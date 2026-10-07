import io
import os
import tempfile
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

tmp = tempfile.mkdtemp(prefix="gamo-smoke-")
os.environ.pop("DATABASE_URL", None)
os.environ["GAMO_DATA_DIR"] = tmp
os.environ["GAMO_ADMIN_PASSWORD"] = "TestGamo2026!"
os.environ["GAMO_HTTPS"] = "0"

import app
from openpyxl import load_workbook

client = app.app.test_client()

def csrf():
    with client.session_transaction() as sess:
        return sess["csrf"]

r = client.get("/login")
assert r.status_code == 200, r.status_code

r = client.post("/login", data={"email": "admin@gamo.sk", "password": "TestGamo2026!"})
assert r.status_code in (302, 303), r.status_code

for path in ("/", "/assets", "/buildings", "/maintenance", "/incidents", "/admin", "/reports"):
    r = client.get(path)
    assert r.status_code == 200, (path, r.status_code)

r = client.get("/api/assets/options")
assert r.status_code == 200 and isinstance(r.get_json(), list)

# Invalid building input must not leave a partially-created record behind.
with client.session_transaction() as sess:
    gamo_org_id = sess["organization_id"]
before_buildings = app.one("select count(*) n from buildings where organization_id=?", (gamo_org_id,))["n"]
r = client.post("/add/building", data={
    "_csrf": csrf(), "code": "BROKEN", "name": "Broken building", "floors_count": "not-a-number"
})
assert r.status_code in (302, 303)
after_buildings = app.one("select count(*) n from buildings where organization_id=?", (gamo_org_id,))["n"]
assert before_buildings == after_buildings

# Full facility hierarchy + operational records.
r = client.post("/add/building", data={
    "_csrf": csrf(), "code": "SMK", "name": "Smoke Building", "address": "Test 1",
    "manager": "QA", "floors_count": "1"
})
assert r.status_code in (302, 303)
building = app.one("select * from buildings where organization_id=? and code=?", (gamo_org_id, "SMK"))
assert building
floor = app.one("select * from floors where building_id=?", (building["id"],))
assert floor

r = client.post("/add/room", data={
    "_csrf": csrf(), "floor_id": str(floor["id"]), "code": "R01", "name": "QA Room",
    "area": "42.5", "tenant": "GAMO", "zone": "QA"
})
assert r.status_code in (302, 303)
room = app.one("select * from rooms where floor_id=? and code=?", (floor["id"], "R01"))
assert room

r = client.post("/add/asset", data={
    "_csrf": csrf(), "asset_id": "QA-000001", "name": "Smoke Asset",
    "building_id": str(building["id"]), "floor_id": str(floor["id"]), "room_id": str(room["id"]),
    "profession": "ELE", "grp": "QA", "type": "TEST", "status": "Prevádzka", "criticality": "B",
    "service_months": "6", "revision_months": "12", "purchase_price": "1250.50"
})
assert r.status_code in (302, 303)
asset = app.one("select * from assets where asset_id=? and organization_id=?", ("QA-000001", gamo_org_id))
assert asset

r = client.post("/add/workorder", data={
    "_csrf": csrf(), "asset_id": str(asset["id"]), "title": "QA servis", "kind": "PM",
    "priority": "Stredná", "status": "Plánované", "due": "2026-01-01",
    "supplier": "QA servis", "technician": "Technik", "cost": "99.90", "description": "Smoke"
})
assert r.status_code in (302, 303)
r = client.post("/add/incident", data={
    "_csrf": csrf(), "asset_id": str(asset["id"]), "title": "QA incident", "severity": "Vysoká",
    "status": "Otvorená", "reported": "2026-10-07", "impact": "Test", "cause": "Smoke", "cost": "10.10"
})
assert r.status_code in (302, 303)

r = client.get("/settings/save")
assert r.status_code in (405, 302)

r = client.post(
    "/settings/save",
    data={"_csrf": csrf(), "section": "smoke-test", "value": "tenant-value"},
    follow_redirects=False,
)
assert r.status_code in (302, 303), r.status_code
r = client.get("/api/setting?section=smoke-test")
assert r.get_json()["value"] == "tenant-value"

r = client.get("/reports/export.xlsx")
assert r.status_code == 200, r.status_code
book = load_workbook(io.BytesIO(r.data), read_only=False, data_only=False)
for sheet in ("Súhrn", "Assety", "Údržba", "Incidenty", "Po termíne"):
    assert sheet in book.sheetnames, book.sheetnames
assert book["Súhrn"]["A1"].value.startswith("GAMO FACILITY REPORT")
asset_values = [cell.value for row in book["Assety"].iter_rows() for cell in row]
assert "QA-000001" in asset_values
assert len(book["Súhrn"]._charts) >= 1

r = client.post(
    "/platform/customer",
    data={
        "_csrf": csrf(),
        "code": "SMOKE",
        "name": "Smoke Customer s.r.o.",
        "admin_name": "Smoke Admin",
        "email": "smoke@example.test",
        "password": "SmokePass2026!",
        "plan": "BASIC",
        "license_status": "Aktívna",
    },
    follow_redirects=False,
)
assert r.status_code in (302, 303), r.status_code
customer = app.one("select * from organizations where code=?", ("SMOKE",))
assert customer and customer["plan"] == "BASIC"

client.get("/logout")
r = client.post("/login", data={"email": "smoke@example.test", "password": "SmokePass2026!"})
assert r.status_code in (302, 303), r.status_code
r = client.get("/assets")
assert r.status_code == 200
assert b"HVAC-000001" not in r.data

r = client.get("/api/buildings/options")
assert r.status_code == 200 and r.get_json() == []

r = client.get("/reports/export.xlsx")
assert r.status_code == 200
empty_book = load_workbook(io.BytesIO(r.data), read_only=False, data_only=False)
assert "Súhrn" in empty_book.sheetnames

r = client.post("/settings/save", data={"section": "missing-csrf", "value": "x"})
assert r.status_code == 400

print("GAMO smoke test OK")
