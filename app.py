from flask import Flask,render_template,request,redirect,url_for,jsonify,flash,session,abort,send_file,has_request_context,make_response
from sqlite3 import IntegrityError
import sqlite3, os, json, secrets, time, io, zipfile, base64, hashlib, re, mimetypes, calendar
try:
 import psycopg
 from psycopg.rows import dict_row
except ImportError:
 psycopg=None; dict_row=None
from functools import wraps
import pyotp, qrcode
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from werkzeug.exceptions import HTTPException
from datetime import date,timedelta,datetime
from urllib.parse import urlsplit
from db_migrations import run_migrations
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
BASE=os.path.dirname(os.path.abspath(__file__))
try:
    with open(os.path.join(BASE,'version.json'),encoding='utf-8-sig') as _vf:
        APP_VERSION=str(json.load(_vf).get('version','unknown'))
except Exception:
    APP_VERSION='unknown'
STATIC_REV=(os.environ.get('RENDER_GIT_COMMIT') or APP_VERSION or 'dev')[:16]
if os.environ.get('GAMO_DESKTOP') == '1':
    DATA_DIR=os.path.join(os.environ.get('LOCALAPPDATA', os.path.expanduser('~')), 'GAMO_FM', 'data')
else:
    DATA_DIR=os.environ.get('GAMO_DATA_DIR', os.path.join(BASE,'data'))
app=Flask(__name__)
app.secret_key=os.environ.get('GAMO_SECRET_KEY') or secrets.token_hex(32)
app.config.update(SESSION_COOKIE_HTTPONLY=True,SESSION_COOKIE_SAMESITE='Strict',SESSION_COOKIE_SECURE=os.environ.get('GAMO_HTTPS','1' if os.environ.get('DATABASE_URL') else '0')=='1',PERMANENT_SESSION_LIFETIME=timedelta(days=30),MAX_CONTENT_LENGTH=16*1024*1024)
DB=os.path.join(DATA_DIR,'gamo.db')
LOGIN_WINDOW=300
LOGIN_MAX_ATTEMPTS=6
_login_attempts={}
ROLE_PERMISSIONS={
 'Administrator':{'view','facility_write','asset_write','maintenance_write','incident_write','documents_write','users_manage','settings_manage','audit_view','platform_manage','reports_view'},
 'Facility Manager':{'view','facility_write','asset_write','maintenance_write','incident_write','documents_write','reports_view'},
 'Technik':{'view','maintenance_write','incident_write'},
 'Servisný technik':{'view','maintenance_write','incident_write'},
 'Viewer':{'view'}
}
PLAN_LIMITS={
 'BASIC':{'users':5,'buildings':2,'assets':500},
 'BUSINESS':{'users':25,'buildings':10,'assets':5000},
 'ENTERPRISE':{'users':None,'buildings':None,'assets':None},
 'INTERNAL':{'users':None,'buildings':None,'assets':None}
}

UPLOAD_ALLOWED_EXTS={'.pdf','.doc','.docx','.xls','.xlsx','.jpg','.jpeg','.png','.txt','.csv','.log','.zip'}
UPLOAD_MAX_BYTES=8*1024*1024
UPLOAD_MIME_BY_EXT={
 '.pdf':{'application/pdf'},'.doc':{'application/msword','application/x-ole-storage'},
 '.docx':{'application/vnd.openxmlformats-officedocument.wordprocessingml.document','application/zip'},
 '.xls':{'application/vnd.ms-excel','application/x-ole-storage'},
 '.xlsx':{'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','application/zip'},
 '.jpg':{'image/jpeg'},'.jpeg':{'image/jpeg'},'.png':{'image/png'},
 '.txt':{'text/plain'},'.csv':{'text/csv','application/csv','text/plain'},
 '.log':{'text/plain'},'.zip':{'application/zip','application/x-zip-compressed'}
}

def read_safe_upload(file,max_bytes=UPLOAD_MAX_BYTES):
 if not file or not getattr(file,'filename',None): return None
 raw_name=os.path.basename(str(file.filename)).strip()
 name=secure_filename(raw_name)[:180]
 if not name:
  raise ValueError('Názov súboru nie je platný.')
 ext=os.path.splitext(name)[1].lower()
 if ext not in UPLOAD_ALLOWED_EXTS:
  raise ValueError('Nepodporovaný typ súboru.')
 mime=(file.mimetype or mimetypes.guess_type(name)[0] or 'application/octet-stream').lower()[:120]
 expected=UPLOAD_MIME_BY_EXT.get(ext,set())
 if mime!='application/octet-stream' and expected and mime not in expected:
  raise ValueError('Typ súboru nezodpovedá jeho prípone.')
 data=file.read(max_bytes+1)
 if len(data)>max_bytes:
  raise ValueError(f'Súbor je väčší ako {max_bytes//(1024*1024)} MB.')
 if ext=='.pdf' and data[:5]!=b'%PDF-':
  raise ValueError('PDF súbor nemá platnú hlavičku.')
 if ext=='.png' and data[:8]!=b'\x89PNG\r\n\x1a\n':
  raise ValueError('PNG súbor nemá platnú hlavičku.')
 if ext in {'.jpg','.jpeg'} and data[:3]!=b'\xff\xd8\xff':
  raise ValueError('JPEG súbor nemá platnú hlavičku.')
 if ext in {'.zip','.docx','.xlsx'} and data[:2]!=b'PK':
  raise ValueError('ZIP/Office súbor nemá platnú hlavičku.')
 return {'name':name,'mime':mime,'size':len(data),'data':data}


def can(permission):
 return permission in ROLE_PERMISSIONS.get(session.get('user_role','Viewer'),{'view'})

def actor_org_id():
 return session.get('organization_id')

def org_id():
 return session.get('support_target_org_id') or actor_org_id()

def support_mode():
 return bool(session.get('support_target_org_id'))

def is_gamo_admin():
 return bool(not support_mode() and actor_org_id() and session.get('organization_code')=='GAMO' and session.get('user_role')=='Administrator')

def owns_building(building_id):
 oid=org_id()
 return bool(oid and one('select id from buildings where id=? and organization_id=?',(building_id,oid)))

def owns_asset(asset_id):
 oid=org_id()
 return bool(oid and one('select a.id from assets a join buildings b on b.id=a.building_id where a.id=? and b.organization_id=?',(asset_id,oid)))

def org_config_values(section,oid=None):
 oid=oid or org_id()
 if not oid: return {}
 try:
  row=one('select v from organization_settings where organization_id=? and k=?',(oid,section))
  raw=json.loads(row['v']) if row and row['v'] else {}
  return raw if isinstance(raw,dict) else {}
 except Exception:
  return {}

def org_runtime_defaults(oid=None):
 oid=oid or org_id()
 tech=org_config_values('Technológie & číselníky',oid)
 asset_cfg=org_config_values('Asset ID generátor',oid)
 service=org_config_values('Servis & SLA',oid)
 notify=org_config_values('Notifikačné centrum',oid)
 roles=org_config_values('Role & bezpečnosť',oid)
 def num(value,default,minimum=0,maximum=999999999):
  try:
   match=re.search(r'-?\d+',str(value or ''))
   return max(minimum,min(maximum,int(match.group(0)) if match else default))
  except Exception:
   return default
 def choice(value,default,allowed):
  value=str(value or default)
  return value if value in allowed else default
 return {
  'asset_id':{
   'mask':str(asset_cfg.get('0') or '{PROF}-000001')[:80],
   'start':num(asset_cfg.get('1'),1,0,999999999),
   'length':num(asset_cfg.get('2'),6,1,12),
   'automatic':str(asset_cfg.get('3') or 'Zapnuté')!='Vypnuté'
  },
  'asset':{
   'profession':str(tech.get('0') or 'HVAC')[:30].upper(),
   'status':choice(tech.get('1'),'Prevádzka',{'Prevádzka','Servis','Mimo prevádzky','Porucha','Vyradené'}),
   'criticality':choice(tech.get('2'),'B',{'A','B','C'}),
   'service_months':num(service.get('0'),6,0,1200),
   'revision_months':num(service.get('1'),12,0,1200)
  },
  'workorder':{
   'priority':choice(service.get('2'),'Stredná',{'Nízka','Stredná','Vysoká','Kritická'})
  },
  'notifications':{
   'due_warning_days':num(notify.get('0'),14,0,365)
  },
  'user':{
   'role':choice(roles.get('0'),'Viewer',{'Administrator','Facility Manager','Technik','Servisný technik','Viewer'})
  }
 }

def next_asset_id(profession,oid=None):
 oid=oid or org_id()
 cfg=org_runtime_defaults(oid)['asset_id']
 prefix=''.join(ch for ch in (profession or 'ASSET').upper() if ch.isalnum())[:10] or 'ASSET'
 mask=(cfg['mask'] or '{PROF}-000001').replace('{PROF}',prefix)
 width=max(1,min(12,int(cfg['length'] or 6)))
 matches=list(re.finditer(r'\d+',mask))
 number_match=matches[-1] if matches else None
 used={str(r['asset_id'] or '').upper() for r in q('select asset_id from assets where organization_id=?',(oid,))}
 n=max(0,int(cfg['start'] or 0))
 for _ in range(1000000):
  if number_match:
   digits=max(width,len(number_match.group(0)))
   candidate=mask[:number_match.start()]+str(n).zfill(digits)+mask[number_match.end():]
  else:
   candidate=f'{mask}{str(n).zfill(width)}'
  candidate=candidate.upper()
  if candidate not in used: return candidate
  n+=1
 raise RuntimeError('Nie je možné nájsť voľné Asset ID podľa zvolenej masky.')

def _as_date(value):
 if not value: return None
 if isinstance(value,datetime): return value.date()
 if isinstance(value,date): return value
 try: return datetime.strptime(str(value)[:10],'%Y-%m-%d').date()
 except Exception: return None

def _add_months(value,months):
 base=_as_date(value)
 if not base or not months: return None
 total=(base.year*12+(base.month-1))+int(months)
 year,month=divmod(total,12); month+=1
 day=min(base.day,calendar.monthrange(year,month)[1])
 return date(year,month,day)

def _last_completed_cycle(asset_id,kind):
 row=one("""select completed_at,due from workorders
  where asset_id=? and kind=? and status='Ukončené'
  order by case when completed_at is null then 1 else 0 end,completed_at desc,id desc limit 1""",(asset_id,kind))
 if not row: return None
 return _as_date(row['completed_at']) or _as_date(row['due'])

def maintenance_plan_rows(limit=None):
 warning_days=org_runtime_defaults()['notifications']['due_warning_days']
 today=date.today()
 assets=q("""select a.id,a.asset_id,a.name,a.status,a.criticality,a.service_months,a.revision_months,a.installed,
  b.code building,b.name building_name,r.code room
  from assets a join buildings b on b.id=a.building_id left join rooms r on r.id=a.room_id
  where a.organization_id=? and a.status<>'Vyradené' order by b.name,a.asset_id""",(org_id(),))
 rows=[]
 for a in assets:
  for kind,interval,label in (
   ('PM',int(a['service_months'] or 0),'Preventívna údržba'),
   ('REV',int(a['revision_months'] or 0),'Pravidelná revízia'),
  ):
   if interval<=0: continue
   last_done=_last_completed_cycle(a['id'],kind)
   anchor=last_done or _as_date(a['installed'])
   due=_add_months(anchor,interval) if anchor else None
   active=one("""select id,status,due,source from workorders
    where asset_id=? and kind=? and status not in ('Ukončené','Zrušené')
    order by case when due is null or due='' then 1 else 0 end,due asc,id desc limit 1""",(a['id'],kind))
   days=(due-today).days if due else None
   if not due: state='missing'
   elif days<0: state='overdue'
   elif days<=warning_days: state='soon'
   else: state='future'
   rows.append({
    'asset_db_id':a['id'],'asset_id':a['asset_id'],'asset':a['name'],'building':a['building'],
    'building_name':a['building_name'],'room':a['room'],'criticality':a['criticality'],'kind':kind,
    'label':label,'interval_months':interval,'anchor':anchor.isoformat() if anchor else None,
    'last_done':last_done.isoformat() if last_done else None,'due':due.isoformat() if due else None,
    'days':days,'state':state,'active_workorder_id':active['id'] if active else None,
    'active_workorder_status':active['status'] if active else None,'active_workorder_due':active['due'] if active else None
   })
 order={'overdue':0,'soon':1,'missing':2,'future':3}
 rows.sort(key=lambda r:(order.get(r['state'],9),r['due'] or '9999-12-31',r['asset_id'],r['kind']))
 return rows[:limit] if limit else rows

def generate_maintenance_plan():
 if support_mode(): abort(403)
 if not can('maintenance_write'): abort(403)
 warning_days=org_runtime_defaults()['notifications']['due_warning_days']
 horizon=date.today()+timedelta(days=warning_days)
 created=0; skipped=0
 for row in maintenance_plan_rows():
  due=_as_date(row['due'])
  if not due or due>horizon or row['active_workorder_id']:
   skipped+=1; continue
  generated_key=f"AUTO:{org_id()}:{row['asset_db_id']}:{row['kind']}:{row['due']}"
  if one('select id from workorders where generated_key=?',(generated_key,)):
   skipped+=1; continue
  priority='Vysoká' if row['criticality']=='A' else ('Nízka' if row['criticality']=='C' else org_runtime_defaults()['workorder']['priority'])
  title=f"{row['label']} · {row['asset_id']}"
  description=f"Automaticky vytvorené z intervalu {row['interval_months']} mes. Základ termínu: {row['anchor'] or '—'}."
  try:
   wid=x("""insert into workorders(asset_id,title,kind,priority,status,due,supplier,technician,cost,description,source,generated_key)
    values(?,?,?,?,?,?,?,?,?,?,?,?)""",(row['asset_db_id'],title,row['kind'],priority,'Plánované',row['due'],'','',0,description,'AUTO',generated_key))
   asset_event(row['asset_db_id'],'WORKORDER_AUTO','Automaticky vytvorený pracovný príkaz',f"{row['kind']} · termín {row['due']}")
   created+=1
  except IntegrityError:
   skipped+=1
 audit('MAINTENANCE_PLAN_SYNC',f'Vytvorených {created} · preskočených {skipped} · horizont {warning_days} dní')
 return created,skipped

def ticket_staff():
 return session.get('user_role') in {'Administrator','Facility Manager','Technik','Servisný technik'}

def platform_ticket_mode():
 return is_gamo_admin()

def next_ticket_no():
 rows=q("select ticket_no from tickets where organization_id=? and ticket_no like ?",(org_id(),'TKT-%'))
 nums=[int(str(r['ticket_no'])[4:]) for r in rows if str(r['ticket_no'] or '')[4:].isdigit()]
 return f"TKT-{(max(nums) if nums else 0)+1:06d}"

def ticket_record(ticket_id):
 if platform_ticket_mode():
  return one_system('select * from tickets where id=?',(ticket_id,))
 t=one('select * from tickets where id=? and organization_id=?',(ticket_id,org_id()))
 if not t: return None
 return t if ticket_staff() or t['created_by']==session.get('user_id') else None

def ticket_unread_count():
 uid=session.get('user_id'); oid=org_id()
 if not uid or not oid: return 0
 if platform_ticket_mode():
  return int(one_system("""select count(*) n from tickets t join organizations o on o.id=t.organization_id
   where o.code<>'GAMO' and exists(
    select 1 from ticket_messages m left join users su on su.id=m.sender_user_id
    where m.ticket_id=t.id and m.id>coalesce(t.platform_last_read_message_id,0)
      and (m.sender_user_id is null or su.organization_id=t.organization_id))""")['n'])
 if ticket_staff():
  return int(one("""select count(*) n from tickets t
   where t.organization_id=? and exists(
    select 1 from ticket_messages m where m.ticket_id=t.id and m.id>coalesce(t.staff_last_read_message_id,0)
      and (m.sender_user_id is null or m.sender_user_id<>?))""",(oid,uid))['n'])
 return int(one("""select count(*) n from tickets t where t.organization_id=? and t.created_by=? and exists(
   select 1 from ticket_messages m where m.ticket_id=t.id and m.id>coalesce(t.customer_last_read_message_id,0)
     and (m.sender_user_id is null or m.sender_user_id<>t.created_by))""",(oid,uid))['n'])

def ticket_inbox_version():
 updated_expr="coalesce(max(t.updated)::text,'')" if USING_POSTGRES else "coalesce(max(t.updated),'')"
 if platform_ticket_mode():
  r=one_system(f"""select {updated_expr} updated,coalesce(sum((select count(*) from ticket_messages m where m.ticket_id=t.id)),0) messages
   from tickets t join organizations o on o.id=t.organization_id where o.code<>'GAMO'""")
 elif ticket_staff():
  r=one(f"""select {updated_expr} updated,coalesce(sum((select count(*) from ticket_messages m where m.ticket_id=t.id)),0) messages
   from tickets t where t.organization_id=?""",(org_id(),))
 else:
  r=one(f"""select {updated_expr} updated,coalesce(sum((select count(*) from ticket_messages m where m.ticket_id=t.id)),0) messages
   from tickets t where t.organization_id=? and t.created_by=?""",(org_id(),session.get('user_id')))
 return f"{r['updated']}|{r['messages']}" if r else '0|0'

def asset_parent_allowed(asset_id,parent_id):
 if not parent_id: return True
 try: current=int(parent_id); target=int(asset_id)
 except (TypeError,ValueError): return False
 seen=set()
 for _ in range(100):
  if current==target or current in seen: return False
  seen.add(current)
  row=one('select parent_id from assets where id=? and organization_id=?',(current,org_id()))
  if not row: return False
  if not row['parent_id']: return True
  current=int(row['parent_id'])
 return False

def owns_floor(floor_id):
 return bool(org_id() and one('select f.id from floors f join buildings b on b.id=f.building_id where f.id=? and b.organization_id=?',(floor_id,org_id())))

def owns_room(room_id):
 return bool(org_id() and one('select r.id from rooms r join floors f on f.id=r.floor_id join buildings b on b.id=f.building_id where r.id=? and b.organization_id=?',(room_id,org_id())))

def owns_workorder(workorder_id):
 return bool(org_id() and one('select w.id from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id where w.id=? and b.organization_id=?',(workorder_id,org_id())))

def owns_incident(incident_id):
 return bool(org_id() and one('select i.id from incidents i join assets a on a.id=i.asset_id join buildings b on b.id=a.building_id where i.id=? and b.organization_id=?',(incident_id,org_id())))

def owns_user(user_id):
 return bool(org_id() and one('select id from users where id=? and organization_id=?',(user_id,org_id())))

def support_access_active(org):
 if not org or not org['support_access_enabled']: return False
 until=org['support_access_until']
 if not until: return True
 try:
  dt=datetime.fromisoformat(str(until).replace('Z','+00:00'))
  now=datetime.now(dt.tzinfo) if dt.tzinfo else datetime.now()
  return dt>=now
 except Exception:
  return False

def customer_access(target_org_id,action,reason=''):
 try:
  writer=x if target_org_id==org_id() else x_system
  writer('insert into customer_access_log(target_organization_id,actor_user_id,actor_name,action,reason,ip) values(?,?,?,?,?,?)',(target_org_id,session.get('user_id'),session.get('user_name','Systém'),action,(reason or '')[:500],request.headers.get('X-Forwarded-For',request.remote_addr or '')))
 except Exception:
  pass

def safe_next_url(value):
 value=(value or '/').strip()
 return value if value.startswith('/') and not value.startswith('//') else '/'

def safe_referrer_url(default='/'):
 ref=(request.referrer or '').strip()
 if not ref: return default
 try:
  parsed=urlsplit(ref)
  if parsed.netloc and parsed.netloc!=request.host: return default
  path=parsed.path or '/'
  if parsed.query: path+='?'+parsed.query
  if parsed.fragment: path+='#'+parsed.fragment
  return safe_next_url(path)
 except Exception:
  return default

def auth_event(organization_id,user_id,event,success,detail=''):
 try:
  x_system('insert into auth_events(organization_id,user_id,event,success,detail,ip,user_agent) values(?,?,?,?,?,?,?)',(
   organization_id,user_id,event,True if (USING_POSTGRES and success) else (1 if success else (False if USING_POSTGRES else 0)),
   (detail or '')[:300],request.headers.get('X-Forwarded-For',request.remote_addr or ''),
   (request.headers.get('User-Agent') or '')[:300]
  ))
 except Exception:
  pass

REMEMBER_COOKIE='gamo_remember_device'
REMEMBER_DAYS=30

def _remember_hash(token):
 return hashlib.sha256((token or '').encode('utf-8')).hexdigest()

def _remember_expired(value):
 if not value: return True
 try:
  if isinstance(value,datetime):
   expiry=value
  else:
   expiry=datetime.fromisoformat(str(value).replace('Z','+00:00'))
  now=datetime.now(expiry.tzinfo) if getattr(expiry,'tzinfo',None) else datetime.utcnow()
  return expiry<=now
 except Exception:
  return True

def revoke_user_devices(user_id):
 try:
  x_system('update remembered_devices set revoked_at=CURRENT_TIMESTAMP where user_id=? and revoked_at is null',(user_id,))
 except Exception:
  pass

def revoke_remember_token(raw_token):
 if not raw_token: return
 try:
  x_system('update remembered_devices set revoked_at=CURRENT_TIMESTAMP where token_hash=? and revoked_at is null',(_remember_hash(raw_token),))
 except Exception:
  pass

def issue_remembered_device(user):
 raw=secrets.token_urlsafe(48)
 expires=datetime.utcnow()+timedelta(days=REMEMBER_DAYS)
 ua=(request.headers.get('User-Agent') or '')[:500]
 label='GAMO Desktop klient' if ('GAMO-Desktop/' in ua or 'pywebview' in ua.lower()) else 'Webové zariadenie'
 ip=(request.headers.get('X-Forwarded-For',request.remote_addr or '') or '').split(',')[0].strip()[:120]
 x_system('insert into remembered_devices(organization_id,user_id,token_hash,label,user_agent,ip_created,expires_at) values(?,?,?,?,?,?,?)',
  (user['organization_id'],user['id'],_remember_hash(raw),label,ua,ip,expires.isoformat(timespec='seconds')+'Z'))
 # Keep the table bounded per account and remove dead tokens.
 try:
  if USING_POSTGRES:
   x_system("""delete from remembered_devices where user_id=? and
    (revoked_at is not null or expires_at<CURRENT_TIMESTAMP or id not in
      (select id from remembered_devices where user_id=? and revoked_at is null and expires_at>=CURRENT_TIMESTAMP order by last_used desc,id desc limit 8))""",(user['id'],user['id']))
  else:
   x_system("""delete from remembered_devices where user_id=? and
    (revoked_at is not null or expires_at<CURRENT_TIMESTAMP or id not in
      (select id from remembered_devices where user_id=? and revoked_at is null and expires_at>=CURRENT_TIMESTAMP order by last_used desc,id desc limit 8))""",(user['id'],user['id']))
 except Exception:
  pass
 return raw

def login_response(target,user,remember=False):
 response=make_response(redirect(safe_next_url(target)))
 previous=request.cookies.get(REMEMBER_COOKIE)
 if remember:
  if previous: revoke_remember_token(previous)
  raw=issue_remembered_device(user)
  response.set_cookie(REMEMBER_COOKIE,raw,max_age=REMEMBER_DAYS*86400,httponly=True,
   secure=bool(app.config.get('SESSION_COOKIE_SECURE')),samesite='Strict',path='/')
 else:
  if previous: revoke_remember_token(previous)
  response.delete_cookie(REMEMBER_COOKIE,path='/',samesite='Strict')
 return response

def restore_remembered_device():
 raw=request.cookies.get(REMEMBER_COOKIE)
 if not raw: return False
 device=one_system('select * from remembered_devices where token_hash=? and revoked_at is null',(_remember_hash(raw),))
 if not device or _remember_expired(device['expires_at']):
  revoke_remember_token(raw); return False
 user=one_system('select * from users where id=?',(device['user_id'],))
 org=one_system('select * from organizations where id=?',(device['organization_id'],))
 license_ok=bool(user and org and user['status']=='Aktívny' and user['organization_id']==org['id'] and
  org['status']=='Aktívny' and org['license_status']=='Aktívna' and
  (not org['license_until'] or str(org['license_until'])[:10]>=date.today().isoformat()))
 if not license_ok:
  revoke_remember_token(raw); return False
 establish_user_session(user,org,True,request.path)
 try: x_system('update remembered_devices set last_used=CURRENT_TIMESTAMP where id=?',(device['id'],))
 except Exception: pass
 auth_event(org['id'],user['id'],'REMEMBER_DEVICE_RESTORE',True,'Relácia obnovená z dôveryhodného zariadenia.')
 return True

def _normalize_mfa_code(code):
 return ''.join(ch for ch in (code or '').upper() if ch.isalnum())

def _recovery_hash(code):
 return hashlib.sha256(_normalize_mfa_code(code).encode('utf-8')).hexdigest()

def verify_mfa_code(user,code,consume_recovery=True):
 normalized=_normalize_mfa_code(code)
 if not normalized: return False
 secret=user['mfa_secret']
 if secret:
  try:
   if pyotp.TOTP(secret).verify(normalized,valid_window=1): return True
  except Exception:
   pass
 try: hashes=json.loads(user['mfa_recovery_codes'] or '[]')
 except Exception: hashes=[]
 digest=_recovery_hash(normalized)
 if digest in hashes:
  if consume_recovery:
   hashes.remove(digest)
   x_system('update users set mfa_recovery_codes=? where id=?',(json.dumps(hashes),user['id']))
  return True
 return False

def establish_user_session(user,org,remember=False,next_url='/'):
 session.clear(); session.permanent=bool(remember)
 session['user_id']=user['id']; session['user_name']=user['name']; session['user_role']=user['role']
 session['organization_id']=user['organization_id']; session['organization_code']=org['code']
 session['must_change_password']=bool(user['must_change_password']); session['csrf']=secrets.token_urlsafe(32)
 x_system("update users set last_login=CURRENT_TIMESTAMP where id=?" if USING_POSTGRES else "update users set last_login=datetime('now') where id=?",(user['id'],))
 auth_event(user['organization_id'],user['id'],'LOGIN_SUCCESS',True,'Prihlásenie dokončené.')
 if session['must_change_password']: return '/account/password'
 if bool(org['mfa_required']) and not bool(user['mfa_enabled']): return '/account/mfa'
 return safe_next_url(next_url)

def plan_limits(oid=None):
 oid=oid or org_id()
 org=one('select plan from organizations where id=?',(oid,)) if oid else None
 return PLAN_LIMITS.get((org['plan'] if org else 'BASIC') or 'BASIC',PLAN_LIMITS['BASIC'])

def resource_count(resource,oid=None):
 oid=oid or org_id()
 if resource=='users': return int(one('select count(*) n from users where organization_id=?',(oid,))['n'])
 if resource=='buildings': return int(one('select count(*) n from buildings where organization_id=?',(oid,))['n'])
 if resource=='assets': return int(one('select count(*) n from assets a join buildings b on b.id=a.building_id where b.organization_id=?',(oid,))['n'])
 return 0

def plan_allows(resource,oid=None):
 limit=plan_limits(oid).get(resource)
 return limit is None or resource_count(resource,oid)<limit

def asset_event(asset_id,event_type,title,detail=''):
 try:
  x('insert into asset_events(asset_id,organization_id,user_id,user_name,event_type,title,detail) values(?,?,?,?,?,?,?)',(asset_id,org_id(),session.get('user_id'),session.get('user_name','Systém'),event_type,title,detail))
 except Exception:
  pass

def audit(action,detail=''):
 try:
  x("insert into audit_log(user_id,user_name,action,detail,ip,organization_id) values(?,?,?,?,?,?)",(session.get('user_id'),session.get('user_name','Systém'),action,detail,request.headers.get('X-Forwarded-For',request.remote_addr or ''),org_id()))
 except Exception:
  pass
DATABASE_URL=os.environ.get('DATABASE_URL','').strip()
USING_POSTGRES=bool(DATABASE_URL)
def _sql(sql):
 return sql.replace('datetime(\'now\')','CURRENT_TIMESTAMP').replace("date('now')",'CURRENT_DATE').replace('?','%s') if USING_POSTGRES else sql
def con(system=False):
 if USING_POSTGRES:
  if not psycopg: raise RuntimeError('DATABASE_URL is set but psycopg is not installed')
  db=psycopg.connect(DATABASE_URL,row_factory=dict_row)
  if system or not has_request_context():
   oid=''; platform_admin='1'
  else:
   oid=str(org_id() or '')
   platform_admin='0'
  db.execute("select set_config('gamo.organization_id',%s,false)",(oid,))
  db.execute("select set_config('gamo.platform_admin',%s,false)",(platform_admin,))
  return db
 c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; c.execute('PRAGMA foreign_keys=ON'); return c
def q(sql,a=()):
 with con() as c:return c.execute(_sql(sql),a).fetchall()
def one(sql,a=()):
 with con() as c:return c.execute(_sql(sql),a).fetchone()
def one_system(sql,a=()):
 with con(system=True) as c:return c.execute(_sql(sql),a).fetchone()
def q_system(sql,a=()):
 with con(system=True) as c:return c.execute(_sql(sql),a).fetchall()
def x(sql,a=()):
 with con() as c:
  statement=_sql(sql)
  lower=statement.lstrip().lower()
  no_id_tables=('insert into settings','insert into organization_settings','insert into platform_meta','insert into schema_migrations')
  if USING_POSTGRES and lower.startswith('insert into') and not lower.startswith(no_id_tables) and ' returning ' not in lower:
   statement+=' RETURNING id'
   r=c.execute(statement,a); row=r.fetchone(); c.commit(); return row['id'] if row else None
  r=c.execute(statement,a); c.commit(); return r.lastrowid if not USING_POSTGRES else r.rowcount
def x_system(sql,a=()):
 with con(system=True) as c:
  statement=_sql(sql)
  lower=statement.lstrip().lower()
  no_id_tables=('insert into settings','insert into organization_settings','insert into platform_meta','insert into schema_migrations')
  if USING_POSTGRES and lower.startswith('insert into') and not lower.startswith(no_id_tables) and ' returning ' not in lower:
   statement+=' RETURNING id'
   r=c.execute(statement,a); row=r.fetchone(); c.commit(); return row['id'] if row else None
  r=c.execute(statement,a); c.commit(); return r.lastrowid if not USING_POSTGRES else r.rowcount
def init_postgres():
 schema=[
  """CREATE TABLE IF NOT EXISTS buildings(id BIGSERIAL PRIMARY KEY,code TEXT UNIQUE,name TEXT,address TEXT,manager TEXT,customer TEXT DEFAULT 'GAMO a.s.',status TEXT DEFAULT 'Aktívna')""",
  """CREATE TABLE IF NOT EXISTS floors(id BIGSERIAL PRIMARY KEY,building_id BIGINT REFERENCES buildings(id) ON DELETE CASCADE,code TEXT,name TEXT)""",
  """CREATE TABLE IF NOT EXISTS rooms(id BIGSERIAL PRIMARY KEY,floor_id BIGINT REFERENCES floors(id) ON DELETE CASCADE,code TEXT,name TEXT,area DOUBLE PRECISION,tenant TEXT,zone TEXT)""",
  """CREATE TABLE IF NOT EXISTS assets(id BIGSERIAL PRIMARY KEY,asset_id TEXT UNIQUE,name TEXT,building_id BIGINT,floor_id BIGINT,room_id BIGINT,profession TEXT,grp TEXT,type TEXT,manufacturer TEXT,model TEXT,serial TEXT,system_id TEXT,parent_id BIGINT,status TEXT,criticality TEXT,service_months INTEGER,revision_months INTEGER,purchase_price DOUBLE PRECISION,installed TEXT,warranty TEXT,ip TEXT,protocol TEXT,notes TEXT)""",
  """CREATE TABLE IF NOT EXISTS workorders(id BIGSERIAL PRIMARY KEY,asset_id BIGINT,title TEXT,kind TEXT,priority TEXT,status TEXT,due TEXT,supplier TEXT,technician TEXT,cost DOUBLE PRECISION,description TEXT)""",
  """CREATE TABLE IF NOT EXISTS incidents(id BIGSERIAL PRIMARY KEY,asset_id BIGINT,title TEXT,severity TEXT,status TEXT,reported TEXT,impact TEXT,cause TEXT,cost DOUBLE PRECISION)""",
  """CREATE TABLE IF NOT EXISTS users(id BIGSERIAL PRIMARY KEY,name TEXT,email TEXT,role TEXT,status TEXT,password_hash TEXT,last_login TEXT)""",
  """CREATE TABLE IF NOT EXISTS settings(k TEXT PRIMARY KEY,v TEXT)""",
  """CREATE TABLE IF NOT EXISTS audit_log(id BIGSERIAL PRIMARY KEY,created TIMESTAMPTZ DEFAULT CURRENT_TIMESTAMP,user_id BIGINT,user_name TEXT,action TEXT,detail TEXT,ip TEXT)""",
  """CREATE TABLE IF NOT EXISTS documents(id BIGSERIAL PRIMARY KEY,building_id BIGINT REFERENCES buildings(id) ON DELETE CASCADE,name TEXT,category TEXT,mime TEXT,size BIGINT,uploaded TIMESTAMPTZ DEFAULT CURRENT_TIMESTAMP,data BYTEA)""",
  """CREATE TABLE IF NOT EXISTS organizations(id BIGSERIAL PRIMARY KEY,code TEXT UNIQUE NOT NULL,name TEXT NOT NULL,status TEXT DEFAULT 'Aktívny',plan TEXT DEFAULT 'INTERNAL',license_status TEXT DEFAULT 'Aktívna',license_until TEXT,branding_name TEXT,brand_color TEXT DEFAULT '#E31B23',brand_tagline TEXT DEFAULT 'FACILITY MANAGEMENT',created TIMESTAMPTZ DEFAULT CURRENT_TIMESTAMP)""",
  """CREATE TABLE IF NOT EXISTS platform_meta(k TEXT PRIMARY KEY,v TEXT)"""
 ]
 with con(system=True) as db:
  for statement in schema: db.execute(statement)
  db.execute("ALTER TABLE buildings ADD COLUMN IF NOT EXISTS customer TEXT DEFAULT 'GAMO a.s.'")
  db.execute("ALTER TABLE buildings ADD COLUMN IF NOT EXISTS organization_id BIGINT")
  db.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS organization_id BIGINT")
  db.execute("ALTER TABLE organizations ADD COLUMN IF NOT EXISTS brand_color TEXT DEFAULT '#E31B23'")
  db.execute("ALTER TABLE organizations ADD COLUMN IF NOT EXISTS brand_tagline TEXT DEFAULT 'FACILITY MANAGEMENT'")
  db.execute("INSERT INTO organizations(code,name,status,plan,license_status,branding_name) VALUES('GAMO','GAMO a.s.','Aktívny','INTERNAL','Aktívna','GAMO a.s.') ON CONFLICT (code) DO NOTHING")
  gamo_org=db.execute("SELECT id FROM organizations WHERE code='GAMO'").fetchone()['id']
  db.execute("UPDATE users SET organization_id=%s WHERE organization_id IS NULL",(gamo_org,))
  db.execute("UPDATE buildings SET organization_id=%s WHERE organization_id IS NULL",(gamo_org,))
  db.commit()
 run_migrations(lambda: con(system=True), True)
 if not one_system('select count(*) n from users')['n']:
  x_system('insert into users(name,email,role,status,password_hash,organization_id) values(?,?,?,?,?,?)',('GAMO Administrator','admin@gamo.sk','Administrator','Aktívny',generate_password_hash(os.environ.get('GAMO_ADMIN_PASSWORD','GamoFM2026!')),gamo_org))
def init():
 os.makedirs(DATA_DIR,exist_ok=True)
 if USING_POSTGRES:
  init_postgres(); return
 with con() as c:
  c.executescript("""CREATE TABLE IF NOT EXISTS buildings(id INTEGER PRIMARY KEY,code TEXT UNIQUE,name TEXT,address TEXT,manager TEXT,customer TEXT DEFAULT 'GAMO a.s.',status TEXT DEFAULT 'Aktívna');CREATE TABLE IF NOT EXISTS floors(id INTEGER PRIMARY KEY,building_id INTEGER REFERENCES buildings(id) ON DELETE CASCADE,code TEXT,name TEXT);CREATE TABLE IF NOT EXISTS rooms(id INTEGER PRIMARY KEY,floor_id INTEGER REFERENCES floors(id) ON DELETE CASCADE,code TEXT,name TEXT,area REAL,tenant TEXT,zone TEXT);CREATE TABLE IF NOT EXISTS assets(id INTEGER PRIMARY KEY,asset_id TEXT UNIQUE,name TEXT,building_id INTEGER,floor_id INTEGER,room_id INTEGER,profession TEXT,grp TEXT,type TEXT,manufacturer TEXT,model TEXT,serial TEXT,system_id TEXT,parent_id INTEGER,status TEXT,criticality TEXT,service_months INTEGER,revision_months INTEGER,purchase_price REAL,installed TEXT,warranty TEXT,ip TEXT,protocol TEXT,notes TEXT);CREATE TABLE IF NOT EXISTS workorders(id INTEGER PRIMARY KEY,asset_id INTEGER,title TEXT,kind TEXT,priority TEXT,status TEXT,due TEXT,supplier TEXT,technician TEXT,cost REAL,description TEXT);CREATE TABLE IF NOT EXISTS incidents(id INTEGER PRIMARY KEY,asset_id INTEGER,title TEXT,severity TEXT,status TEXT,reported TEXT,impact TEXT,cause TEXT,cost REAL);CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,name TEXT,email TEXT,role TEXT,status TEXT);CREATE TABLE IF NOT EXISTS settings(k TEXT PRIMARY KEY,v TEXT);""")
  bcols=[r[1] for r in c.execute("pragma table_info(buildings)").fetchall()]
  if 'customer' not in bcols: c.execute("alter table buildings add column customer TEXT DEFAULT 'GAMO a.s.'")
  cols=[r[1] for r in c.execute("pragma table_info(users)").fetchall()]
  if 'password_hash' not in cols: c.execute("alter table users add column password_hash TEXT")
  if 'last_login' not in cols: c.execute("alter table users add column last_login TEXT")
  c.execute("""CREATE TABLE IF NOT EXISTS audit_log(id INTEGER PRIMARY KEY,created TEXT DEFAULT CURRENT_TIMESTAMP,user_id INTEGER,user_name TEXT,action TEXT,detail TEXT,ip TEXT);""")
  c.execute("""CREATE TABLE IF NOT EXISTS documents(id INTEGER PRIMARY KEY,building_id INTEGER REFERENCES buildings(id) ON DELETE CASCADE,name TEXT,category TEXT,mime TEXT,size INTEGER,uploaded TEXT DEFAULT CURRENT_TIMESTAMP,data BLOB);""")
  c.execute("""CREATE TABLE IF NOT EXISTS organizations(id INTEGER PRIMARY KEY,code TEXT UNIQUE NOT NULL,name TEXT NOT NULL,status TEXT DEFAULT 'Aktívny',plan TEXT DEFAULT 'INTERNAL',license_status TEXT DEFAULT 'Aktívna',license_until TEXT,branding_name TEXT,brand_color TEXT DEFAULT '#E31B23',brand_tagline TEXT DEFAULT 'FACILITY MANAGEMENT',created TEXT DEFAULT CURRENT_TIMESTAMP);""")
  orgcols=[r[1] for r in c.execute("pragma table_info(organizations)").fetchall()]
  if 'brand_color' not in orgcols: c.execute("alter table organizations add column brand_color TEXT DEFAULT '#E31B23'")
  if 'brand_tagline' not in orgcols: c.execute("alter table organizations add column brand_tagline TEXT DEFAULT 'FACILITY MANAGEMENT'")
  c.execute("""CREATE TABLE IF NOT EXISTS platform_meta(k TEXT PRIMARY KEY,v TEXT);""")
  bcols=[r[1] for r in c.execute("pragma table_info(buildings)").fetchall()]
  if 'organization_id' not in bcols: c.execute("alter table buildings add column organization_id INTEGER")
  ucols=[r[1] for r in c.execute("pragma table_info(users)").fetchall()]
  if 'organization_id' not in ucols: c.execute("alter table users add column organization_id INTEGER")
  c.execute("insert or ignore into organizations(code,name,status,plan,license_status,branding_name) values('GAMO','GAMO a.s.','Aktívny','INTERNAL','Aktívna','GAMO a.s.')")
  gamo_org=c.execute("select id from organizations where code='GAMO'").fetchone()[0]
  c.execute("update users set organization_id=? where organization_id is null",(gamo_org,))
  c.execute("update buildings set organization_id=? where organization_id is null",(gamo_org,))
  if not c.execute('select count(*) n from buildings').fetchone()['n']:
   c.execute("insert into buildings(code,name,address,manager,organization_id) values('A','GAMO Centrum – Budova A','Kyjevské námestie 6, Banská Bystrica','Facility Management',?)",(gamo_org,)); c.execute("insert into buildings(code,name,address,manager,organization_id) values('B','GAMO Centrum – Budova B','Banská Bystrica','Facility Management',?)",(gamo_org,))
   a=c.execute("select id from buildings where code='A'").fetchone()[0]; b=c.execute("select id from buildings where code='B'").fetchone()[0]
   f1=c.execute("insert into floors(building_id,code,name) values(?,?,?)",(a,'1.NP','Prízemie')).lastrowid; f2=c.execute("insert into floors(building_id,code,name) values(?,?,?)",(a,'2.NP','Administratíva')).lastrowid; fb=c.execute("insert into floors(building_id,code,name) values(?,?,?)",(b,'1.NP','Technické podlažie')).lastrowid
   r1=c.execute("insert into rooms(floor_id,code,name,area,tenant,zone) values(?,?,?,?,?,?)",(f1,'A005','Kancelária',28.5,'GAMO','HVAC zóna 01')).lastrowid;r2=c.execute("insert into rooms(floor_id,code,name,area,tenant,zone) values(?,?,?,?,?,?)",(f1,'A008','Kancelária',31.2,'GAMO','HVAC zóna 01')).lastrowid;r3=c.execute("insert into rooms(floor_id,code,name,area,tenant,zone) values(?,?,?,?,?,?)",(f2,'A214','Serverovňa',18,'GAMO','IT kritická zóna')).lastrowid;r4=c.execute("insert into rooms(floor_id,code,name,area,tenant,zone) values(?,?,?,?,?,?)",(fb,'B202','Technická miestnosť',42,'','HVAC zóna B')).lastrowid
   p=c.execute("insert into assets(asset_id,name,building_id,floor_id,room_id,profession,grp,type,manufacturer,model,serial,system_id,status,criticality,service_months,revision_months,purchase_price,ip,protocol) values(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",('HVAC-000001','VRV vonkajšia jednotka',a,f1,r1,'HVAC','VRV systém','VRV-OUT','Daikin','RXYQ','SN240001','HVAC-A-VRV-01','Prevádzka','A',6,12,8900,'','Modbus')).lastrowid
   c.execute("insert into assets(asset_id,name,building_id,floor_id,room_id,profession,grp,type,manufacturer,model,system_id,parent_id,status,criticality,purchase_price) values(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",('HVAC-000002','VRV vnútorná jednotka A005',a,f1,r1,'HVAC','VRV systém','VRV-IN','Daikin','FXZQ25','HVAC-A-VRV-01',p,'Prevádzka','B',1450));c.execute("insert into assets(asset_id,name,building_id,floor_id,room_id,profession,grp,type,manufacturer,model,system_id,parent_id,status,criticality,purchase_price) values(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",('HVAC-000003','VRV vnútorná jednotka A008',a,f1,r2,'HVAC','VRV systém','VRV-IN','Daikin','FXZQ25','HVAC-A-VRV-01',p,'Prevádzka','B',1450));c.execute("insert into assets(asset_id,name,building_id,floor_id,room_id,profession,grp,type,manufacturer,model,system_id,status,criticality,purchase_price,ip,protocol) values(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",('SLB-000001','Core switch serverovne',a,f2,r3,'SLB','LAN/WAN','Switch','Cisco','Catalyst','SLB-A-LAN-01','Prevádzka','A',3200,'10.0.1.2','SNMP'))
   c.execute("insert into workorders(asset_id,title,kind,priority,status,due,supplier,technician,cost,description) values(?,?,?,?,?,?,?,?,?,?)",(p,'Preventívny servis VRV','PM','Stredná','Plánované',str(date.today()+timedelta(days=12)),'Servis HVAC s.r.o.','Ján Technik',350,'Kontrola systému a filtrov'))
   c.execute("insert into incidents(asset_id,title,severity,status,reported,impact,cost) values(?,?,?,?,?,?,?)",(p,'Kolísanie tlaku VRV','Porucha','Otvorená',str(date.today()),'A005, A008 – znížený komfort',0))
   c.execute("insert into users(name,email,role,status,password_hash,organization_id) values(?,?,?,?,?,?)",('GAMO Administrator','admin@gamo.sk','Administrator','Aktívny',generate_password_hash(os.environ.get('GAMO_ADMIN_PASSWORD','GamoFM2026!')),gamo_org))
 run_migrations(con, False)
init()
with con() as c:
 c.execute(_sql("update users set password_hash=? where (password_hash is null or password_hash='') and lower(email)=?"),(generate_password_hash(os.environ.get('GAMO_ADMIN_PASSWORD','GamoFM2026!')),'admin@gamo.sk')); c.commit()

@app.after_request
def security_headers(response):
 response.headers['X-Content-Type-Options']='nosniff'
 response.headers['X-Frame-Options']='DENY'
 response.headers['Referrer-Policy']='same-origin'
 response.headers['Permissions-Policy']='camera=(), microphone=(), geolocation=()'
 response.headers['Content-Security-Policy']="default-src 'self'; img-src 'self' data:; style-src 'self' 'unsafe-inline'; script-src 'self' 'unsafe-inline'; connect-src 'self'; font-src 'self' data:; object-src 'none'; base-uri 'self'; frame-ancestors 'none'"
 if app.config.get('SESSION_COOKIE_SECURE'):
  response.headers['Strict-Transport-Security']='max-age=31536000; includeSubDomains'
 if request.endpoint=='static':
  response.headers['Cache-Control']='public, max-age=31536000, immutable'
 elif request.path.startswith('/backup') or request.path.endswith('/backup') or request.path.startswith('/privacy'):
  response.headers['Cache-Control']='no-store, private'
 elif 'Cache-Control' not in response.headers:
  response.headers['Cache-Control']='private, no-cache'
 return response

@app.before_request
def require_login():
 if request.endpoint in ('login','login_mfa','static') or request.path.startswith('/static/'): return
 if not session.get('user_id'):
  if not restore_remembered_device(): return redirect(url_for('login',next=request.path))
 current=one_system('select id,status,role,organization_id,must_change_password,mfa_enabled from users where id=?',(session.get('user_id'),))
 if not current or current['status']!='Aktívny' or current['organization_id']!=actor_org_id():
  session.clear(); return redirect(url_for('login'))
 if current['role']!=session.get('user_role'):
  session['user_role']=current['role']
 session['must_change_password']=bool(current['must_change_password'])
 if session['must_change_password'] and request.endpoint not in {'account_password','logout'}:
  return redirect('/account/password')
 actor_org=one_system('select * from organizations where id=?',(actor_org_id(),)) if actor_org_id() else None
 if actor_org: session['organization_code']=actor_org['code']
 if actor_org and bool(actor_org['mfa_required']) and not bool(current['mfa_enabled']) and request.endpoint not in {'account_mfa','account_password','logout'}:
  return redirect('/account/mfa')
 target=None
 if support_mode():
  target=one_system('select * from organizations where id=?',(session.get('support_target_org_id'),))
  if not actor_org or actor_org['code']!='GAMO' or current['role']!='Administrator' or not target or not support_access_active(target):
   session.pop('support_target_org_id',None); session.pop('support_target_name',None); session.pop('support_target_code',None); target=None
 org=target or actor_org
 license_ok=bool(org and org['status']=='Aktívny' and org['license_status']=='Aktívna' and (not org['license_until'] or str(org['license_until'])[:10]>=date.today().isoformat()))
 if not license_ok:
  session.clear()
  return redirect(url_for('login'))
 if request.method in {'POST','PUT','PATCH','DELETE'}:
  supplied=request.form.get('_csrf') or request.headers.get('X-CSRF-Token','')
  expected=session.get('csrf','')
  if not supplied or not expected or not secrets.compare_digest(str(supplied),str(expected)):
   abort(400,description='Neplatný bezpečnostný token požiadavky.')
 role=session.get('user_role','Viewer')
 if request.path.startswith('/admin') or request.path.startswith('/settings/') or request.endpoint in {'admin','update_user'}:
  if not can('users_manage') and not can('settings_manage'):
   if request.method in {'POST','PUT','PATCH','DELETE'}: abort(403)
   flash('Na túto časť nemáš administrátorské oprávnenie.','error'); return redirect('/')
 if request.method in {'POST','PUT','PATCH','DELETE'}:
  what=(request.view_args or {}).get('what')
  needed={'user':'users_manage','building':'facility_write','floor':'facility_write','room':'facility_write','asset':'asset_write','workorder':'maintenance_write','incident':'incident_write'}.get(what)
  endpoint_permissions={
   'upload_building_document':'documents_write','delete_document':'documents_write','save_setting':'settings_manage',
   'update_user':'users_manage','platform_customer':'platform_manage','platform_customer_update':'platform_manage',
   'platform_customer_branding':'platform_manage'
  }
  needed=endpoint_permissions.get(request.endpoint,needed)
  if needed and not can(needed): abort(403)

def render_app_error(status,title,message):
 return render_template('error.html',status=status,title=title,message=message,app_version=APP_VERSION),status

@app.errorhandler(400)
def error_400(exc):
 return render_app_error(400,'Neplatná požiadavka','Požiadavku sa nepodarilo spracovať. Skontroluj údaje a skús to znova.')

@app.errorhandler(403)
def error_403(exc):
 return render_app_error(403,'Prístup zamietnutý','Na túto časť nemáš oprávnenie alebo je operácia z bezpečnostných dôvodov blokovaná.')

@app.errorhandler(404)
def error_404(exc):
 return render_app_error(404,'Záznam sa nenašiel','Požadovaná stránka alebo záznam neexistuje, prípadne nepatrí tvojej organizácii.')

@app.errorhandler(413)
def error_413(exc):
 return render_app_error(413,'Súbor je príliš veľký','Maximálna veľkosť jedného uploadu je 16 MB.')

@app.errorhandler(500)
def error_500(exc):
 try: app.logger.exception('Unhandled GAMO application error',exc_info=exc)
 except Exception: pass
 return render_app_error(500,'Nastala interná chyba','Dáta zostali zachované. Skús operáciu zopakovať; ak problém pretrváva, vytvor support ticket.')

@app.route('/login',methods=['GET','POST'])
def login():
 if session.get('user_id'): return redirect('/')
 error=None
 if request.method=='POST':
  email=(request.form.get('login_identifier') or request.form.get('email') or '').strip().lower(); password=request.form.get('password') or ''; remember=request.form.get('remember')=='1'
  key=(request.headers.get('X-Forwarded-For',request.remote_addr or '')+'|'+email)
  now=time.time(); attempts=[t for t in _login_attempts.get(key,[]) if now-t<LOGIN_WINDOW]
  if len(attempts)>=LOGIN_MAX_ATTEMPTS:
   return render_template('login.html',error='Príliš veľa neúspešných pokusov. Skús to znova o pár minút.'),429
  u=one_system('select * from users where lower(email)=?',(email,))
  password_ok=bool(u and u['status']=='Aktívny' and u['password_hash'] and check_password_hash(u['password_hash'],password))
  if password_ok:
   org=one_system('select * from organizations where id=?',(u['organization_id'],)) if u['organization_id'] else None
   license_ok=bool(org and org['status']=='Aktívny' and org['license_status']=='Aktívna' and (not org['license_until'] or str(org['license_until'])[:10]>=date.today().isoformat()))
   if not license_ok:
    auth_event(u['organization_id'],u['id'],'LOGIN_BLOCKED',False,'Neaktívna organizácia alebo licencia.')
    error='Licencia organizácie nie je aktívna alebo jej platnosť skončila. Kontaktuj GAMO.'
   elif bool(u['mfa_enabled']):
    _login_attempts.pop(key,None); session.clear(); session.permanent=remember
    session['mfa_pending_user_id']=u['id']; session['mfa_pending_remember']=remember
    session['mfa_pending_next']=safe_next_url(request.args.get('next'))
    session['mfa_csrf']=secrets.token_urlsafe(32); session['mfa_failures']=0
    auth_event(u['organization_id'],u['id'],'PASSWORD_VERIFIED',True,'Čaká sa na druhý faktor.')
    return redirect('/login/mfa')
   else:
    _login_attempts.pop(key,None)
    target=establish_user_session(u,org,remember,request.args.get('next'))
    return login_response(target,u,remember)
  else:
   attempts.append(now); _login_attempts[key]=attempts
   if u: auth_event(u['organization_id'],u['id'],'LOGIN_PASSWORD_FAILURE',False,'Nesprávne heslo.')
   error='Nesprávny e-mail alebo heslo.'
 return render_template('login.html',error=error)

@app.route('/login/mfa',methods=['GET','POST'])
def login_mfa():
 if session.get('user_id'): return redirect('/')
 uid=session.get('mfa_pending_user_id')
 if not uid: return redirect('/login')
 u=one_system('select * from users where id=?',(uid,))
 org=one_system('select * from organizations where id=?',(u['organization_id'],)) if u else None
 if not u or not org or u['status']!='Aktívny' or not bool(u['mfa_enabled']):
  session.clear(); return redirect('/login')
 error=None
 if request.method=='POST':
  supplied=request.form.get('_csrf') or ''
  expected=session.get('mfa_csrf') or ''
  if not supplied or not expected or not secrets.compare_digest(str(supplied),str(expected)):
   abort(400,description='Neplatný bezpečnostný token požiadavky.')
  code=request.form.get('code') or ''
  if verify_mfa_code(u,code,True):
   remember=bool(session.get('mfa_pending_remember')); next_url=session.get('mfa_pending_next') or '/'
   auth_event(u['organization_id'],u['id'],'MFA_SUCCESS',True,'Druhý faktor overený.')
   target=establish_user_session(u,org,remember,next_url)
   return login_response(target,u,remember)
  failures=int(session.get('mfa_failures') or 0)+1; session['mfa_failures']=failures
  auth_event(u['organization_id'],u['id'],'MFA_FAILURE',False,f'Neúspešný MFA pokus #{failures}.')
  if failures>=6:
   session.clear()
   return render_template('login_mfa.html',error='Príliš veľa neúspešných pokusov. Prihlás sa znova.',csrf_token=''),429
  error='Neplatný overovací alebo recovery kód.'
 return render_template('login_mfa.html',error=error,csrf_token=session.get('mfa_csrf',''),user_name=u['name'])

@app.get('/logout')
def logout():
 raw=request.cookies.get(REMEMBER_COOKIE)
 if raw: revoke_remember_token(raw)
 session.clear()
 response=make_response(redirect('/login'))
 response.delete_cookie(REMEMBER_COOKIE,path='/',samesite='Strict')
 return response

@app.route('/account/password',methods=['GET','POST'])
def account_password():
 if support_mode():
  flash('Najprv ukonči GAMO support režim.','error'); return redirect('/')
 u=one('select * from users where id=? and organization_id=?',(session.get('user_id'),actor_org_id()))
 if not u: abort(404)
 if request.method=='POST':
  current_password=request.form.get('current_password') or ''
  new_password=request.form.get('new_password') or ''
  confirm=request.form.get('confirm_password') or ''
  if not check_password_hash(u['password_hash'],current_password):
   flash('Aktuálne heslo nie je správne.','error')
  elif len(new_password)<10:
   flash('Nové heslo musí mať aspoň 10 znakov.','error')
  elif new_password!=confirm:
   flash('Nové heslá sa nezhodujú.','error')
  elif check_password_hash(u['password_hash'],new_password):
   flash('Nové heslo musí byť odlišné od aktuálneho.','error')
  else:
   x('update users set password_hash=?,must_change_password=? where id=?',(generate_password_hash(new_password),False if USING_POSTGRES else 0,u['id']))
   revoke_user_devices(u['id'])
   session['must_change_password']=False
   audit('PASSWORD_CHANGE','Používateľ zmenil svoje heslo.')
   flash('Heslo bolo bezpečne zmenené.','success')
   return redirect('/')
 return render_template('account_password.html',forced=bool(u['must_change_password']),csrf_token=session.get('csrf',''),current_user=u)

@app.route('/account/mfa',methods=['GET','POST'])
def account_mfa():
 if support_mode():
  flash('Nastavenie MFA nie je dostupné v support režime.','error'); return redirect('/')
 u=one('select * from users where id=? and organization_id=?',(session.get('user_id'),actor_org_id()))
 org=one('select * from organizations where id=?',(actor_org_id(),))
 if not u or not org: abort(404)
 new_codes=None
 if request.method=='POST':
  action=request.form.get('action') or 'enable'
  if action=='enable':
   secret=session.get('mfa_setup_secret')
   code=request.form.get('code') or ''
   if not secret or not pyotp.TOTP(secret).verify(_normalize_mfa_code(code),valid_window=1):
    flash('Kód z autentifikátora nie je správny.','error')
   else:
    new_codes=[secrets.token_hex(6).upper() for _ in range(10)]
    hashes=[_recovery_hash(x) for x in new_codes]
    x('update users set mfa_secret=?,mfa_recovery_codes=?,mfa_enabled=?,mfa_enabled_at=CURRENT_TIMESTAMP where id=?' if USING_POSTGRES else "update users set mfa_secret=?,mfa_recovery_codes=?,mfa_enabled=?,mfa_enabled_at=datetime('now') where id=?",(secret,json.dumps(hashes),True if USING_POSTGRES else 1,u['id']))
    revoke_user_devices(u['id'])
    session.pop('mfa_setup_secret',None)
    audit('MFA_ENABLED','Používateľ zapol dvojfaktorové overenie.')
    auth_event(actor_org_id(),u['id'],'MFA_ENABLED',True,'TOTP MFA aktivované.')
    u=one('select * from users where id=?',(u['id'],))
    flash('MFA bolo úspešne zapnuté. Recovery kódy si bezpečne ulož.','success')
  elif action=='regenerate':
   password=request.form.get('current_password') or ''; code=request.form.get('code') or ''
   if not check_password_hash(u['password_hash'],password) or not verify_mfa_code(u,code,True):
    flash('Heslo alebo MFA kód nie je správny.','error')
   else:
    new_codes=[secrets.token_hex(6).upper() for _ in range(10)]
    x('update users set mfa_recovery_codes=? where id=?',(json.dumps([_recovery_hash(x) for x in new_codes]),u['id']))
    audit('MFA_RECOVERY_REGENERATED','Vygenerované nové recovery kódy.')
    flash('Recovery kódy boli nahradené novými.','success')
  elif action=='disable':
   if bool(org['mfa_required']):
    flash('Organizácia vyžaduje MFA. Najprv vypni povinné MFA v Privacy Center.','error')
   else:
    password=request.form.get('current_password') or ''; code=request.form.get('code') or ''
    if not check_password_hash(u['password_hash'],password) or not verify_mfa_code(u,code,True):
     flash('Heslo alebo MFA kód nie je správny.','error')
    else:
     x('update users set mfa_secret=?,mfa_recovery_codes=?,mfa_enabled=?,mfa_enabled_at=? where id=?',(None,None,False if USING_POSTGRES else 0,None,u['id']))
     revoke_user_devices(u['id'])
     audit('MFA_DISABLED','Používateľ vypol dvojfaktorové overenie.')
     auth_event(actor_org_id(),u['id'],'MFA_DISABLED',True,'TOTP MFA vypnuté.')
     u=one('select * from users where id=?',(u['id'],))
     flash('MFA bolo vypnuté.','success')
 if not bool(u['mfa_enabled']):
  secret=session.get('mfa_setup_secret')
  if not secret:
   secret=pyotp.random_base32(); session['mfa_setup_secret']=secret
  uri=pyotp.TOTP(secret).provisioning_uri(name=u['email'],issuer_name=(org['branding_name'] or org['name'] or 'GAMO'))
  img=qrcode.make(uri); qr=io.BytesIO(); img.save(qr,format='PNG')
  qr_data='data:image/png;base64,'+base64.b64encode(qr.getvalue()).decode('ascii')
 else:
  secret=None; qr_data=None
 try: recovery_left=len(json.loads(u['mfa_recovery_codes'] or '[]'))
 except Exception: recovery_left=0
 current_token=request.cookies.get(REMEMBER_COOKIE)
 current_hash=_remember_hash(current_token) if current_token else ''
 devices=q('select id,label,user_agent,ip_created,created,last_used,expires_at,token_hash from remembered_devices where user_id=? and organization_id=? and revoked_at is null order by last_used desc,id desc',(u['id'],actor_org_id()))
 device_rows=[]
 for d in devices:
  item=dict(d); item['is_current']=bool(current_hash and item.get('token_hash')==current_hash); item.pop('token_hash',None); device_rows.append(item)
 return render_template('account_mfa.html',current_user=u,org=org,secret=secret,qr_data=qr_data,
  recovery_left=recovery_left,new_recovery_codes=new_codes,csrf_token=session.get('csrf',''),remembered_devices=device_rows)

@app.post('/account/device/<int:i>/revoke')
def revoke_account_device(i):
 if support_mode(): abort(403)
 d=one('select id,token_hash,label from remembered_devices where id=? and user_id=? and organization_id=? and revoked_at is null',(i,session.get('user_id'),actor_org_id()))
 if not d: abort(404)
 x('update remembered_devices set revoked_at=CURRENT_TIMESTAMP where id=? and user_id=? and organization_id=?',(i,session.get('user_id'),actor_org_id()))
 audit('DEVICE_REVOKE',d['label'] or f'Device #{i}')
 auth_event(actor_org_id(),session.get('user_id'),'REMEMBERED_DEVICE_REVOKED',True,d['label'] or f'Device #{i}')
 response=make_response(redirect('/account/mfa#devices'))
 raw=request.cookies.get(REMEMBER_COOKIE)
 if raw and _remember_hash(raw)==d['token_hash']:
  response.delete_cookie(REMEMBER_COOKIE,path='/',samesite='Strict')
 flash('Zapamätané zariadenie bolo odvolané.','success')
 return response

@app.post('/account/devices/revoke-others')
def revoke_other_account_devices():
 if support_mode(): abort(403)
 raw=request.cookies.get(REMEMBER_COOKIE); current_hash=_remember_hash(raw) if raw else ''
 if current_hash:
  x('update remembered_devices set revoked_at=CURRENT_TIMESTAMP where user_id=? and organization_id=? and revoked_at is null and token_hash<>?',(session.get('user_id'),actor_org_id(),current_hash))
 else:
  x('update remembered_devices set revoked_at=CURRENT_TIMESTAMP where user_id=? and organization_id=? and revoked_at is null',(session.get('user_id'),actor_org_id()))
 audit('DEVICE_REVOKE_OTHERS','Odvolané ostatné zapamätané zariadenia.')
 auth_event(actor_org_id(),session.get('user_id'),'REMEMBERED_DEVICES_REVOKED',True,'Ostatné zariadenia boli odvolané.')
 flash('Ostatné zapamätané zariadenia boli odvolané.','success')
 return redirect('/account/mfa#devices')

@app.context_processor
def ctx():
 org=one('select * from organizations where id=?',(org_id(),)) if org_id() else None
 brand_name=(org['branding_name'] or org['name']) if org else 'GAMO a.s.'
 brand_color=(org['brand_color'] or '#E31B23') if org else '#E31B23'
 brand_tagline=(org['brand_tagline'] or 'FACILITY MANAGEMENT') if org else 'FACILITY MANAGEMENT'
 unread_tickets=0
 try: unread_tickets=ticket_unread_count()
 except Exception: unread_tickets=0
 return dict(today=date.today(),app_version=APP_VERSION,static_rev=STATIC_REV,runtime_defaults=org_runtime_defaults() if org_id() else {},csrf_token=session.get('csrf',''),current_user={'id':session.get('user_id'),'name':session.get('user_name',''),'role':session.get('user_role',''),'organization_id':actor_org_id()},current_org=org,brand_name=brand_name,brand_color=brand_color,brand_tagline=brand_tagline,is_gamo_admin=is_gamo_admin(),support_mode=support_mode(),support_customer_name=session.get('support_target_name',''),ticket_unread_count=unread_tickets,can=can)

@app.route('/onboarding',methods=['GET','POST'])
def onboarding():
 org=one('select * from organizations where id=?',(org_id(),))
 if not org or org['code']=='GAMO': return redirect('/')
 if session.get('user_role')!='Administrator':
  return redirect('/')
 if request.method=='POST':
  code=(request.form.get('code') or '').strip().upper()
  name=(request.form.get('name') or '').strip()
  address=(request.form.get('address') or '').strip()
  manager=(request.form.get('manager') or session.get('user_name','')).strip()
  try: floors=max(1,min(50,int(request.form.get('floors_count') or 1)))
  except ValueError: floors=1
  if not code or not name:
   flash('Kód a názov prvej budovy sú povinné.','error')
  elif not plan_allows('buildings'):
   flash('Licenčný plán už neumožňuje pridať ďalšiu budovu.','error')
  elif one('select id from buildings where organization_id=? and upper(code)=?',(org_id(),code)):
   flash('Budova s týmto kódom už existuje.','error')
  else:
   bid=x('insert into buildings(code,name,address,manager,customer,organization_id) values(?,?,?,?,?,?)',(code,name,address,manager,org['name'],org_id()))
   for n in range(1,floors+1):
    x('insert into floors(building_id,code,name) values(?,?,?)',(bid,f'{n}.NP',f'{n}. nadzemné podlažie'))
   x('update organizations set onboarding_complete=? where id=?',(True if USING_POSTGRES else 1,org_id()))
   audit('ONBOARDING_COMPLETE',f'{name} · {floors} podlaží')
   flash('Firemné prostredie je pripravené. Teraz môžeš doplniť miestnosti a assety.','success')
   return redirect(f'/building/{bid}')
 limits=plan_limits()
 return render_template('onboarding.html',org=org,limits=limits,csrf_token=session.get('csrf',''))

@app.post('/onboarding/skip')
def onboarding_skip():
 org=one('select * from organizations where id=?',(org_id(),))
 if not org or org['code']=='GAMO' or session.get('user_role')!='Administrator': abort(403)
 x('update organizations set onboarding_complete=? where id=?',(True if USING_POSTGRES else 1,org_id()))
 audit('ONBOARDING_SKIP','Onboarding preskočený administrátorom')
 return redirect('/')

@app.route('/')
def dashboard():
 oid=org_id()
 org=one('select * from organizations where id=?',(oid,)) if oid else None
 if org and org['code']!='GAMO' and session.get('user_role')=='Administrator' and not bool(org['onboarding_complete']):
  return redirect('/onboarding')
 s={
  'assets':one('select count(*) n from assets a join buildings b on b.id=a.building_id where b.organization_id=?',(oid,))['n'],
  'buildings':one('select count(*) n from buildings where organization_id=?',(oid,))['n'],
  'rooms':one('select count(*) n from rooms r join floors f on f.id=r.floor_id join buildings b on b.id=f.building_id where b.organization_id=?',(oid,))['n'],
  'open':one("select count(*) n from incidents i join assets a on a.id=i.asset_id join buildings b on b.id=a.building_id where b.organization_id=? and i.status not in ('Ukončená','Vyriešená')",(oid,))['n'],
  'critical':one("select count(*) n from assets a join buildings b on b.id=a.building_id where b.organization_id=? and a.criticality='A'",(oid,))['n'],
  'orders':one("select count(*) n from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id where b.organization_id=? and w.status not in ('Ukončené','Zrušené')",(oid,))['n'],
  'high_incidents':one("select count(*) n from incidents i join assets a on a.id=i.asset_id join buildings b on b.id=a.building_id where b.organization_id=? and i.status not in ('Ukončená','Vyriešená') and i.severity in ('Vysoká','Kritická','Havária')",(oid,))['n'],
  'overdue':one("select count(*) n from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id where b.organization_id=? and w.status not in ('Ukončené','Zrušené') and w.due is not null and w.due!='' and date(w.due)<date('now')",(oid,))['n'],
  'tickets':one("select count(*) n from tickets where organization_id=? and status not in ('Vyriešený','Uzavretý')",(oid,))['n']
 }
 report={
  'maintenance_cost':one('select coalesce(sum(w.cost),0) n from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id where b.organization_id=?',(oid,))['n'],
  'incident_cost':one('select coalesce(sum(i.cost),0) n from incidents i join assets a on a.id=i.asset_id join buildings b on b.id=a.building_id where b.organization_id=?',(oid,))['n'],
  'closed_orders':one("select count(*) n from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id where b.organization_id=? and w.status='Ukončené'",(oid,))['n'],
  'total_orders':one('select count(*) n from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id where b.organization_id=?',(oid,))['n']
 }; report['total_cost']=report['maintenance_cost']+report['incident_cost']
 profession_sql="select coalesce(a.profession,'Iné') label,round(coalesce(sum(w.cost),0)::numeric,2) value from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id where b.organization_id=? group by a.profession order by value desc limit 6" if USING_POSTGRES else "select coalesce(a.profession,'Iné') label,round(coalesce(sum(w.cost),0),2) value from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id where b.organization_id=? group by a.profession order by value desc limit 6"
 return render_template('index.html',page='dashboard',s=s,report=report,profession_costs=q(profession_sql,(oid,)),buildings=q('select * from buildings where organization_id=?',(oid,)),recent=q('select w.*,a.asset_id,a.name asset from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id where b.organization_id=? order by w.id desc limit 6',(oid,)),incidents=q('select i.*,a.asset_id from incidents i join assets a on a.id=i.asset_id join buildings b on b.id=a.building_id where b.organization_id=? order by i.id desc limit 5',(oid,)))
@app.route('/tickets')
def tickets():
 uid=session.get('user_id')
 if platform_ticket_mode():
  base="""select t.*,o.name organization_name,o.code organization_code,cu.name creator_name,au.name assigned_name,
    b.code building_code,a.asset_id asset_code,
    (select count(*) from ticket_messages m where m.ticket_id=t.id) message_count,
    (select count(*) from ticket_messages m left join users su on su.id=m.sender_user_id
      where m.ticket_id=t.id and m.id>coalesce(t.platform_last_read_message_id,0)
      and (m.sender_user_id is null or su.organization_id=t.organization_id)) platform_unread
    from tickets t
    join organizations o on o.id=t.organization_id
    left join users cu on cu.id=t.created_by
    left join users au on au.id=t.assigned_to
    left join buildings b on b.id=t.building_id
    left join assets a on a.id=t.asset_id
    where o.code<>'GAMO'
    order by case t.status when 'Nový' then 0 when 'Otvorený' then 1 when 'Rieši sa' then 2 when 'Čaká na zákazníka' then 3 when 'Vyriešený' then 4 else 5 end,t.updated desc"""
  rows=q_system(base)
  stats={'total':len(rows),'open':sum(1 for r in rows if r['status'] not in {'Vyriešený','Uzavretý'}),'critical':sum(1 for r in rows if r['priority']=='Kritická' and r['status'] not in {'Vyriešený','Uzavretý'}),'waiting':sum(1 for r in rows if r['status']=='Čaká na zákazníka'),'unread':sum(1 for r in rows if int(r['platform_unread'] or 0)>0)}
  return render_template('index.html',page='tickets',tickets=rows,ticket_stats=stats,ticket_is_staff=True,
   ticket_buildings=[],ticket_assets=[],ticket_staff_users=[],ticket_platform_inbox=True)

 oid=org_id(); staff=ticket_staff()
 base="""select t.*,cu.name creator_name,au.name assigned_name,b.code building_code,a.asset_id asset_code,
  (select count(*) from ticket_messages m where m.ticket_id=t.id) message_count,
  (select count(*) from ticket_messages m where m.ticket_id=t.id and m.id>coalesce(t.staff_last_read_message_id,0)
    and (m.sender_user_id is null or m.sender_user_id<>?)) staff_unread,
  (select count(*) from ticket_messages m where m.ticket_id=t.id and m.id>coalesce(t.customer_last_read_message_id,0)
    and (m.sender_user_id is null or m.sender_user_id<>t.created_by)) customer_unread
  from tickets t
  left join users cu on cu.id=t.created_by
  left join users au on au.id=t.assigned_to
  left join buildings b on b.id=t.building_id
  left join assets a on a.id=t.asset_id
  where t.organization_id=?"""
 params=[uid,oid]
 if not staff:
  base+=" and t.created_by=?"; params.append(uid)
 base+=" order by case t.status when 'Nový' then 0 when 'Otvorený' then 1 when 'Rieši sa' then 2 when 'Čaká na zákazníka' then 3 when 'Vyriešený' then 4 else 5 end,t.updated desc"
 rows=[dict(r) for r in q(base,tuple(params))]
 for r in rows:
  r['user_unread']=int((r['staff_unread'] if staff else r['customer_unread']) or 0)
 stats={'total':len(rows),'open':sum(1 for r in rows if r['status'] not in {'Vyriešený','Uzavretý'}),'critical':sum(1 for r in rows if r['priority']=='Kritická' and r['status'] not in {'Vyriešený','Uzavretý'}),'waiting':sum(1 for r in rows if r['status']=='Čaká na zákazníka'),'unread':sum(1 for r in rows if r['user_unread']>0)}
 buildings=q('select id,code,name from buildings where organization_id=? order by name',(oid,))
 assets=q('select id,asset_id,name,building_id from assets where organization_id=? order by asset_id',(oid,))
 prefill_asset_id=(request.args.get('asset_id') or '').strip()
 prefill_building_id=(request.args.get('building_id') or '').strip()
 if prefill_asset_id:
  prefill_asset=one('select id,building_id from assets where id=? and organization_id=?',(prefill_asset_id,oid))
  if prefill_asset: prefill_building_id=str(prefill_asset['building_id'])
  else: prefill_asset_id=''
 if prefill_building_id and not owns_building(prefill_building_id): prefill_building_id=''
 staff_users=q("select id,name,role from users where organization_id=? and status='Aktívny' and role in ('Administrator','Facility Manager','Technik','Servisný technik') order by case role when 'Facility Manager' then 0 when 'Administrator' then 1 else 2 end,name",(oid,))
 return render_template('index.html',page='tickets',tickets=rows,ticket_stats=stats,ticket_is_staff=staff,ticket_buildings=buildings,ticket_assets=assets,ticket_staff_users=staff_users,ticket_platform_inbox=False,ticket_prefill_asset=prefill_asset_id,ticket_prefill_building=prefill_building_id,ticket_auto_open=bool(prefill_asset_id or prefill_building_id))

@app.get('/api/tickets/inbox-state')
def api_ticket_inbox_state():
 return jsonify({'version':ticket_inbox_version(),'unread':ticket_unread_count()})

@app.post('/tickets/create')
def ticket_create():
 oid=org_id(); uid=session.get('user_id'); f=request.form
 subject=(f.get('subject') or '').strip()[:160]; body=(f.get('message') or '').strip()[:5000]
 try: attachment=read_safe_upload(request.files.get('attachment'))
 except ValueError as exc:
  flash(str(exc),'error'); return redirect('/tickets')
 category=(f.get('category') or 'Požiadavka').strip(); priority=(f.get('priority') or 'Stredná').strip()
 building_id=f.get('building_id') or None; asset_id=f.get('asset_id') or None
 if not subject or len(body)<2:
  flash('Ticket potrebuje predmet a úvodnú správu.','error'); return redirect('/tickets')
 if category not in {'Požiadavka','Porucha','Prístup','Budova','Asset','Iné'} or priority not in {'Nízka','Stredná','Vysoká','Kritická'}: abort(400)
 if building_id and not owns_building(building_id): abort(404)
 if asset_id:
  asset=one('select id,building_id from assets where id=? and organization_id=?',(asset_id,oid))
  if not asset: abort(404)
  if building_id and str(asset['building_id'])!=str(building_id):
   flash('Vybraný asset nepatrí do zvolenej budovy.','error'); return redirect('/tickets')
  if not building_id: building_id=asset['building_id']
 assigned=f.get('assigned_to') if ticket_staff() else None
 if assigned and not one("select id from users where id=? and organization_id=? and status='Aktívny' and role in ('Administrator','Facility Manager','Technik','Servisný technik')",(assigned,oid)): abort(404)
 if not assigned:
  manager=one("select id from users where organization_id=? and status='Aktívny' and id<>? and role in ('Facility Manager','Administrator') order by case role when 'Facility Manager' then 0 else 1 end,id limit 1",(oid,uid))
  assigned=manager['id'] if manager else None
 no=next_ticket_no(); now=datetime.utcnow().isoformat(timespec='seconds')+'Z'
 try:
  tid=x('insert into tickets(organization_id,ticket_no,created_by,assigned_to,subject,category,priority,status,building_id,asset_id,customer_last_read_at) values(?,?,?,?,?,?,?,?,?,?,?)',(oid,no,uid,assigned,subject,category,priority,'Nový',building_id,asset_id,now))
 except IntegrityError:
  no=next_ticket_no()
  tid=x('insert into tickets(organization_id,ticket_no,created_by,assigned_to,subject,category,priority,status,building_id,asset_id,customer_last_read_at) values(?,?,?,?,?,?,?,?,?,?,?)',(oid,no,uid,assigned,subject,category,priority,'Nový',building_id,asset_id,now))
 mid=x('insert into ticket_messages(ticket_id,organization_id,sender_user_id,sender_name,body) values(?,?,?,?,?)',(tid,oid,uid,session.get('user_name','Používateľ'),body))
 if attachment:
  x('insert into ticket_attachments(organization_id,ticket_id,message_id,sender_user_id,name,mime,size,data) values(?,?,?,?,?,?,?,?)',(oid,tid,mid,uid,attachment['name'],attachment['mime'],attachment['size'],attachment['data']))
 x('update tickets set customer_last_read_message_id=?,customer_last_read_at=CURRENT_TIMESTAMP where id=? and organization_id=?',(mid,tid,oid))
 audit('TICKET_CREATE',f'{no} · {subject}'); flash(f'Ticket {no} bol vytvorený.','success')
 return redirect(f'/ticket/{tid}')

@app.get('/ticket/<int:i>')
def ticket_detail(i):
 t=ticket_record(i)
 if not t: abort(404)
 platform_view=platform_ticket_mode()
 oid=t['organization_id'] if platform_view else org_id()
 read_one=one_system if platform_view else one
 read_all=q_system if platform_view else q
 write=x_system if platform_view else x
 creator=read_one('select id,name,email,role from users where id=? and organization_id=?',(t['created_by'],oid)) if t['created_by'] else None
 assigned=read_one('select id,name,email,role from users where id=? and organization_id=?',(t['assigned_to'],oid)) if t['assigned_to'] else None
 building=read_one('select id,code,name from buildings where id=? and organization_id=?',(t['building_id'],oid)) if t['building_id'] else None
 asset=read_one('select id,asset_id,name from assets where id=? and organization_id=?',(t['asset_id'],oid)) if t['asset_id'] else None
 if platform_view:
  messages=read_all("""select m.*,u.role sender_role,u.organization_id sender_org_id,
   case when u.organization_id is not null and u.organization_id<>? then 1 else 0 end is_support
   from ticket_messages m left join users u on u.id=m.sender_user_id
   where m.ticket_id=? and m.organization_id=? order by m.id""",(oid,i,oid))
 else:
  messages=read_all("""select m.*,u.role sender_role,u.organization_id sender_org_id,
   case when m.sender_user_id is not null and not exists(
    select 1 from users local_sender where local_sender.id=m.sender_user_id and local_sender.organization_id=?
   ) then 1 else 0 end is_support
   from ticket_messages m left join users u on u.id=m.sender_user_id
   where m.ticket_id=? and m.organization_id=? order by m.id""",(oid,i,oid))
 messages=[dict(m) for m in messages]
 for m in messages:
  m['attachments']=read_all('select id,name,mime,size,uploaded from ticket_attachments where message_id=? and ticket_id=? and organization_id=? order by id',(m['id'],i,oid))
 last_message_id=messages[-1]['id'] if messages else 0
 if platform_view:
  write('update tickets set platform_last_read_message_id=?,platform_last_read_at=CURRENT_TIMESTAMP where id=? and organization_id=?',(last_message_id,i,oid))
  customer_access(oid,'TICKET_SUPPORT_VIEW',f"{t['ticket_no']} · ticket conversation")
 elif t['created_by']==session.get('user_id'):
  write('update tickets set customer_last_read_message_id=?,customer_last_read_at=CURRENT_TIMESTAMP where id=? and organization_id=?',(last_message_id,i,oid))
 elif ticket_staff():
  write('update tickets set staff_last_read_message_id=?,staff_last_read_at=CURRENT_TIMESTAMP where id=? and organization_id=?',(last_message_id,i,oid))
 staff_users=read_all("select id,name,role from users where organization_id=? and status='Aktívny' and role in ('Administrator','Facility Manager','Technik','Servisný technik') order by name",(oid,))
 customer_org=one_system('select id,code,name,support_access_enabled,support_access_until from organizations where id=?',(oid,)) if platform_view else None
 return render_template('index.html',page='ticket',ticket=t,ticket_creator=creator,ticket_assigned=assigned,ticket_building=building,ticket_asset=asset,ticket_messages=messages,ticket_staff_users=staff_users,ticket_is_staff=(True if platform_view else ticket_staff()),ticket_platform_view=platform_view,ticket_customer_org=customer_org)

@app.get('/api/ticket/<int:i>/messages')
def api_ticket_messages(i):
 t=ticket_record(i)
 if not t: abort(404)
 try: after=max(0,int(request.args.get('after') or 0))
 except (TypeError,ValueError): after=0
 platform_view=platform_ticket_mode()
 oid=t['organization_id'] if platform_view else org_id()
 read_all=q_system if platform_view else q
 write=x_system if platform_view else x
 if platform_view:
  rows=read_all("""select m.*,u.role sender_role,u.organization_id sender_org_id,
   case when u.organization_id is not null and u.organization_id<>? then 1 else 0 end is_support
   from ticket_messages m left join users u on u.id=m.sender_user_id
   where m.ticket_id=? and m.organization_id=? and m.id>? order by m.id""",(oid,i,oid,after))
 else:
  rows=read_all("""select m.*,u.role sender_role,u.organization_id sender_org_id,
   case when m.sender_user_id is not null and not exists(
    select 1 from users local_sender where local_sender.id=m.sender_user_id and local_sender.organization_id=?
   ) then 1 else 0 end is_support
   from ticket_messages m left join users u on u.id=m.sender_user_id
   where m.ticket_id=? and m.organization_id=? and m.id>? order by m.id""",(oid,i,oid,after))
 if rows:
  last_id=rows[-1]['id']
  if platform_view:
   write('update tickets set platform_last_read_message_id=?,platform_last_read_at=CURRENT_TIMESTAMP where id=? and organization_id=?',(last_id,i,oid))
  elif t['created_by']==session.get('user_id'):
   write('update tickets set customer_last_read_message_id=?,customer_last_read_at=CURRENT_TIMESTAMP where id=? and organization_id=?',(last_id,i,oid))
  elif ticket_staff():
   write('update tickets set staff_last_read_message_id=?,staff_last_read_at=CURRENT_TIMESTAMP where id=? and organization_id=?',(last_id,i,oid))
 else:
  last_id=after
 fresh=(one_system if platform_view else one)('select status,priority,updated from tickets where id=? and organization_id=?',(i,oid))
 messages=[]
 for r in rows:
  external_support=bool(r['is_support'])
  attachments=read_all('select id,name,mime,size,uploaded from ticket_attachments where message_id=? and ticket_id=? and organization_id=? order by id',(r['id'],i,oid))
  messages.append({
   'id':r['id'],
   'sender_name':r['sender_name'] or 'Systém',
   'sender_role':'GAMO Support' if external_support else (r['sender_role'] or 'Používateľ'),
   'body':r['body'] or '',
   'created':str(r['created'] or ''),
   'mine':bool(r['sender_user_id']==session.get('user_id')),
   'support':external_support,
   'attachments':[{'id':a['id'],'name':a['name'],'mime':a['mime'],'size':a['size'],'url':f"/ticket-attachment/{a['id']}/download"} for a in attachments]
  })
 return jsonify({
  'messages':messages,'last_id':last_id,
  'status':fresh['status'] if fresh else t['status'],
  'priority':fresh['priority'] if fresh else t['priority'],
  'updated':str(fresh['updated'] if fresh else t['updated'])
 })

@app.post('/ticket/<int:i>/message')
def ticket_message(i):
 t=ticket_record(i)
 if not t: abort(404)
 body=(request.form.get('message') or '').strip()[:5000]
 try: attachment=read_safe_upload(request.files.get('attachment'))
 except ValueError as exc:
  if request.headers.get('X-Requested-With')=='GAMO-Live-Chat': return jsonify({'ok':False,'error':str(exc)}),400
  flash(str(exc),'error'); return redirect(f'/ticket/{i}')
 if not body:
  if request.headers.get('X-Requested-With')=='GAMO-Live-Chat': return jsonify({'ok':False,'error':'Správa nemôže byť prázdna.'}),400
  flash('Správa nemôže byť prázdna.','error'); return redirect(f'/ticket/{i}')
 uid=session.get('user_id'); platform_view=platform_ticket_mode()
 oid=t['organization_id'] if platform_view else org_id()
 write=x_system if platform_view else x
 mid=write('insert into ticket_messages(ticket_id,organization_id,sender_user_id,sender_name,body) values(?,?,?,?,?)',(i,oid,uid,session.get('user_name','Používateľ'),body))
 if attachment:
  write('insert into ticket_attachments(organization_id,ticket_id,message_id,sender_user_id,name,mime,size,data) values(?,?,?,?,?,?,?,?)',(oid,i,mid,uid,attachment['name'],attachment['mime'],attachment['size'],attachment['data']))
 if platform_view:
  new_status='Otvorený' if t['status'] in {'Nový','Uzavretý'} else t['status']
  write('update tickets set status=?,updated=CURRENT_TIMESTAMP,platform_last_read_message_id=?,platform_last_read_at=CURRENT_TIMESTAMP where id=? and organization_id=?',(new_status,mid,i,oid))
  customer_access(oid,'TICKET_SUPPORT_REPLY',f"{t['ticket_no']} · GAMO odpoveď")
  audit_for_org(oid,'TICKET_SUPPORT_REPLY',f"{t['ticket_no']} · GAMO odpoveď")
 else:
  is_creator=t['created_by']==uid
  if is_creator:
   new_status='Otvorený' if t['status'] in {'Vyriešený','Uzavretý','Čaká na zákazníka'} else t['status']
   write('update tickets set status=?,updated=CURRENT_TIMESTAMP,customer_last_read_message_id=?,customer_last_read_at=CURRENT_TIMESTAMP where id=? and organization_id=?',(new_status,mid,i,oid))
  else:
   new_status='Otvorený' if t['status']=='Nový' else t['status']
   write('update tickets set status=?,updated=CURRENT_TIMESTAMP,staff_last_read_message_id=?,staff_last_read_at=CURRENT_TIMESTAMP where id=? and organization_id=?',(new_status,mid,i,oid))
  audit('TICKET_MESSAGE',f"{t['ticket_no']} · nová správa")
 if request.headers.get('X-Requested-With')=='GAMO-Live-Chat':
  return jsonify({'ok':True,'message_id':mid,'status':new_status})
 return redirect(f'/ticket/{i}#conversation')

@app.get('/ticket-attachment/<int:i>/download')
def download_ticket_attachment(i):
 platform_view=platform_ticket_mode()
 read_one=one_system if platform_view else one
 a=read_one('select * from ticket_attachments where id=?',(i,)) if platform_view else read_one('select * from ticket_attachments where id=? and organization_id=?',(i,org_id()))
 if not a: abort(404)
 t=ticket_record(a['ticket_id'])
 if not t or int(t['organization_id'])!=int(a['organization_id']): abort(404)
 if platform_view: customer_access(a['organization_id'],'TICKET_ATTACHMENT_DOWNLOAD',a['name'])
 return send_file(io.BytesIO(bytes(a['data'])),mimetype=a['mime'] or 'application/octet-stream',as_attachment=True,download_name=a['name'])

@app.post('/ticket/<int:i>/manage')
def ticket_manage(i):
 platform_view=platform_ticket_mode()
 if not platform_view and not ticket_staff(): abort(403)
 t=ticket_record(i)
 if not t: abort(404)
 oid=t['organization_id'] if platform_view else org_id()
 status=(request.form.get('status') or '').strip(); priority=(request.form.get('priority') or '').strip(); assigned=request.form.get('assigned_to') or None
 if status not in {'Nový','Otvorený','Rieši sa','Čaká na zákazníka','Vyriešený','Uzavretý'} or priority not in {'Nízka','Stredná','Vysoká','Kritická'}: abort(400)
 read_one=one_system if platform_view else one
 write=x_system if platform_view else x
 if assigned and not read_one("select id from users where id=? and organization_id=? and status='Aktívny' and role in ('Administrator','Facility Manager','Technik','Servisný technik')",(assigned,oid)): abort(404)
 closed=datetime.utcnow().isoformat(timespec='seconds')+'Z' if status=='Uzavretý' else None
 write('update tickets set status=?,priority=?,assigned_to=?,closed_at=?,updated=CURRENT_TIMESTAMP where id=? and organization_id=?',(status,priority,assigned,closed,i,oid))
 if platform_view:
  customer_access(oid,'TICKET_SUPPORT_UPDATE',f"{t['ticket_no']} · {status} · {priority}")
  audit_for_org(oid,'TICKET_SUPPORT_UPDATE',f"{t['ticket_no']} · {status} · {priority}")
 else:
  audit('TICKET_UPDATE',f"{t['ticket_no']} · {status} · {priority}")
 flash('Ticket bol aktualizovaný.','success')
 return redirect(f'/ticket/{i}')

@app.route('/reports')
def reports():
 if not can('reports_view'): abort(403)
 oid=org_id()
 stats={
  'assets':one('select count(*) n from assets a join buildings b on b.id=a.building_id where b.organization_id=?',(oid,))['n'],
  'maintenance_cost':one('select coalesce(sum(w.cost),0) n from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id where b.organization_id=?',(oid,))['n'],
  'incident_cost':one('select coalesce(sum(i.cost),0) n from incidents i join assets a on a.id=i.asset_id join buildings b on b.id=a.building_id where b.organization_id=?',(oid,))['n'],
  'open_incidents':one("select count(*) n from incidents i join assets a on a.id=i.asset_id join buildings b on b.id=a.building_id where b.organization_id=? and i.status not in ('Ukončená','Vyriešená')",(oid,))['n'],
  'overdue':one("select count(*) n from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id where b.organization_id=? and w.status not in ('Ukončené','Zrušené') and w.due is not null and w.due!='' and date(w.due)<date('now')",(oid,))['n']
 }
 stats['total_cost']=stats['maintenance_cost']+stats['incident_cost']
 building_rows=q("""select b.id,b.code,b.name,
  (select count(*) from assets a where a.building_id=b.id) assets,
  (select coalesce(sum(w.cost),0) from workorders w join assets a on a.id=w.asset_id where a.building_id=b.id) maintenance_cost,
  (select coalesce(sum(i.cost),0) from incidents i join assets a on a.id=i.asset_id where a.building_id=b.id) incident_cost,
  (select count(*) from incidents i join assets a on a.id=i.asset_id where a.building_id=b.id and i.status not in ('Ukončená','Vyriešená')) open_incidents
  from buildings b where b.organization_id=? order by b.name""",(oid,))
 profession_rows=q("""select coalesce(a.profession,'Iné') profession,count(*) assets,
  coalesce(sum(a.purchase_price),0) asset_value
  from assets a join buildings b on b.id=a.building_id
  where b.organization_id=? group by a.profession order by assets desc""",(oid,))
 return render_template('index.html',page='reports',report_stats=stats,report_buildings=building_rows,report_professions=profession_rows)

@app.get('/reports/export.xlsx')
def reports_export_xlsx():
 if not can('reports_view'): abort(403)
 oid=org_id(); org=one('select * from organizations where id=?',(oid,))
 brand=(org['brand_color'] if org and org['brand_color'] else '#17365D').replace('#','').upper()
 if len(brand)!=6 or any(ch not in '0123456789ABCDEF' for ch in brand): brand='17365D'
 dark='17365D'; blue='246BFD'; light='EAF1FB'; pale='F7F9FC'; green='DFF3E8'; red='FCE8EC'; amber='FFF3D8'; white='FFFFFF'; gray='667085'
 wb=Workbook(); ws=wb.active; ws.title='Súhrn'
 wb.properties.creator='GAMO Facility Platform'
 wb.properties.title=f"Facility report · {org['name'] if org else ''}"
 wb.properties.subject='Assety, údržba, incidenty a manažérske KPI'
 thin=Side(style='thin',color='DDE3EC')
 money_fmt='#,##0.00 [$€-sk-SK]'
 date_fmt='dd.mm.yyyy'

 def excel_safe(value):
  if isinstance(value,str) and value[:1] in {'=','+','-','@'}: return "'" + value
  return value

 def excel_date(value):
  if not value: return None
  if isinstance(value,(date,datetime)): return value
  try: return datetime.strptime(str(value)[:10],'%Y-%m-%d').date()
  except Exception: return excel_safe(value)

 buildings=q("""select b.id,b.code,b.name,b.address,b.manager,b.status,
  (select count(*) from floors f where f.building_id=b.id) floors,
  (select count(*) from rooms r join floors f on f.id=r.floor_id where f.building_id=b.id) rooms,
  (select coalesce(sum(r.area),0) from rooms r join floors f on f.id=r.floor_id where f.building_id=b.id) area,
  (select count(*) from assets a where a.building_id=b.id) assets,
  (select coalesce(sum(w.cost),0) from workorders w join assets a on a.id=w.asset_id where a.building_id=b.id) maintenance_cost,
  (select coalesce(sum(i.cost),0) from incidents i join assets a on a.id=i.asset_id where a.building_id=b.id) incident_cost,
  (select count(*) from incidents i join assets a on a.id=i.asset_id where a.building_id=b.id and i.status not in ('Ukončená','Vyriešená')) open_incidents
  from buildings b where b.organization_id=? order by b.name""",(oid,))
 professions=q("""select coalesce(a.profession,'Iné') profession,count(*) assets,
  coalesce(sum(a.purchase_price),0) asset_value
  from assets a join buildings b on b.id=a.building_id
  where b.organization_id=? group by a.profession order by assets desc""",(oid,))
 assets_rows=q("""select a.asset_id,a.asset_tag,a.name,b.name building,f.code floor,r.code room,a.profession,a.grp,a.type,a.manufacturer,a.model,a.serial,a.system_id,a.status,a.criticality,a.purchase_price,a.installed,a.warranty
  from assets a join buildings b on b.id=a.building_id left join floors f on f.id=a.floor_id left join rooms r on r.id=a.room_id
  where b.organization_id=? order by a.asset_id""",(oid,))
 wo_rows=q("""select a.asset_id,a.name asset,b.name building,w.title,w.kind,w.priority,w.status,w.due,w.supplier,w.technician,w.cost,w.description
  from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id
  where b.organization_id=? order by w.id desc""",(oid,))
 inc_rows=q("""select a.asset_id,a.name asset,b.name building,i.title,i.severity,i.status,i.reported,i.impact,i.cause,i.cost
  from incidents i join assets a on a.id=i.asset_id join buildings b on b.id=a.building_id
  where b.organization_id=? order by i.id desc""",(oid,))
 overdue_rows=[r for r in wo_rows if r['status'] not in {'Ukončené','Zrušené'} and r['due'] and str(r['due'])[:10]<date.today().isoformat()]

 maintenance_cost=float(sum(float(r['cost'] or 0) for r in wo_rows))
 incident_cost=float(sum(float(r['cost'] or 0) for r in inc_rows))
 total_cost=maintenance_cost+incident_cost
 open_incidents=sum(1 for r in inc_rows if r['status'] not in {'Ukončená','Vyriešená'})

 # Executive summary
 ws.sheet_view.showGridLines=False
 ws.freeze_panes='A12'
 ws.merge_cells('A1:H2')
 ws['A1']=f"GAMO FACILITY REPORT · {org['name'] if org else ''}"
 ws['A1'].font=Font(color=white,bold=True,size=20)
 ws['A1'].fill=PatternFill('solid',fgColor=brand or dark)
 ws['A1'].alignment=Alignment(vertical='center')
 for row in ws['A1:H2']:
  for cell in row: cell.fill=PatternFill('solid',fgColor=brand or dark)
 ws.row_dimensions[1].height=25; ws.row_dimensions[2].height=16
 ws.merge_cells('A3:H3'); ws['A3']=f"Generované {datetime.now().strftime('%d.%m.%Y %H:%M')} · {org['plan'] if org else ''} · organizácia {org['code'] if org else ''}"
 ws['A3'].font=Font(color=gray,size=9); ws['A3'].alignment=Alignment(vertical='center')

 cards=[
  ('A5','B7','CELKOVÉ NÁKLADY',total_cost,money_fmt,brand or dark),
  ('C5','D7','ÚDRŽBA',maintenance_cost,money_fmt,blue),
  ('E5','F7','OTVORENÉ INCIDENTY',open_incidents,'0','C0394E' if open_incidents else '16845A'),
  ('G5','H7','PO TERMÍNE',len(overdue_rows),'0','C98412' if overdue_rows else '16845A')
 ]
 for start,end,label,value,numfmt,color in cards:
  ws.merge_cells(f'{start}:{end}')
  cell=ws[start]; cell.value=value; cell.number_format=numfmt
  cell.font=Font(size=17,bold=True,color=dark)
  cell.fill=PatternFill('solid',fgColor=pale)
  cell.alignment=Alignment(horizontal='center',vertical='center')
  cell.border=Border(left=Side(style='medium',color=color),top=thin,right=thin,bottom=thin)
  label_cell=ws.cell(cell.row-1,cell.column); label_cell.value=label; label_cell.font=Font(size=8,bold=True,color=gray)

 ws['A9']='PORTFÓLIO'; ws['A9'].font=Font(size=9,bold=True,color=gray)
 portfolio=[('Budovy',len(buildings)),('Assety',len(assets_rows)),('Servisné záznamy',len(wo_rows)),('Incidenty',len(inc_rows))]
 for idx,(label,value) in enumerate(portfolio):
  col=1+idx*2
  ws.cell(9,col,label).font=Font(size=8,bold=True,color=gray)
  ws.cell(9,col+1,value).font=Font(size=11,bold=True,color=dark)

 # Building cost table
 headers=['Kód','Budova','Assety','Údržba','Incidenty','Otvorené','Spolu']
 start_row=11
 for col,hdr in enumerate(headers,1):
  cell=ws.cell(start_row,col,hdr); cell.fill=PatternFill('solid',fgColor=dark); cell.font=Font(color=white,bold=True,size=9); cell.alignment=Alignment(vertical='center')
 for r_idx,b in enumerate(buildings,start_row+1):
  total=float(b['maintenance_cost'] or 0)+float(b['incident_cost'] or 0)
  vals=[excel_safe(b['code']),excel_safe(b['name']),b['assets'],float(b['maintenance_cost'] or 0),float(b['incident_cost'] or 0),b['open_incidents'],total]
  for col,val in enumerate(vals,1):
   cell=ws.cell(r_idx,col,val); cell.border=Border(bottom=thin); cell.alignment=Alignment(vertical='top',wrap_text=True)
  for col in (4,5,7): ws.cell(r_idx,col).number_format=money_fmt
  if b['open_incidents']: ws.cell(r_idx,6).fill=PatternFill('solid',fgColor=red)
 end_row=max(start_row+1,start_row+len(buildings))
 if buildings:
  tab=Table(displayName='BuildingCostTable',ref=f'A{start_row}:G{end_row}')
  tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showFirstColumn=False,showLastColumn=False)
  ws.add_table(tab)

 # Charts use a dedicated hidden data sheet so chart anchors stay visible.
 chart_data=wb.create_sheet('_Grafy'); chart_data.sheet_state='hidden'
 chart_data.append(['Budova','Údržba','Incidenty','','Profesia','Assety'])
 chart_buildings=sorted(buildings,key=lambda b:float(b['maintenance_cost'] or 0)+float(b['incident_cost'] or 0),reverse=True)[:12]
 chart_professions=list(professions)[:10]
 for idx,b in enumerate(chart_buildings,2):
  chart_data.cell(idx,1,excel_safe(b['code'])); chart_data.cell(idx,2,float(b['maintenance_cost'] or 0)); chart_data.cell(idx,3,float(b['incident_cost'] or 0))
 for idx,p in enumerate(chart_professions,2):
  chart_data.cell(idx,5,excel_safe(p['profession'])); chart_data.cell(idx,6,p['assets'])
 if chart_buildings:
  chart=BarChart(); chart.type='col'; chart.style=10; chart.title='Top náklady podľa budovy'; chart.y_axis.title='EUR'; chart.height=7.5; chart.width=13
  chart.add_data(Reference(chart_data,min_col=2,max_col=3,min_row=1,max_row=1+len(chart_buildings)),titles_from_data=True)
  chart.set_categories(Reference(chart_data,min_col=1,min_row=2,max_row=1+len(chart_buildings)))
  chart.legend.position='b'; ws.add_chart(chart,'J4')
 if chart_professions:
  pie=PieChart(); pie.title='Assety podľa profesie'; pie.height=7.5; pie.width=10
  pie.add_data(Reference(chart_data,min_col=6,min_row=1,max_row=1+len(chart_professions)),titles_from_data=True)
  pie.set_categories(Reference(chart_data,min_col=5,min_row=2,max_row=1+len(chart_professions)))
  pie.legend.position='r'; ws.add_chart(pie,'J19')
 for col,width in {'A':13,'B':31,'C':11,'D':16,'E':16,'F':12,'G':16,'H':4}.items(): ws.column_dimensions[col].width=width
 ws.auto_filter.ref=f'A{start_row}:G{end_row}' if buildings else None
 ws.print_title_rows='1:11'; ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.sheet_properties.pageSetUpPr.fitToPage=True

 def styled_sheet(name,title,headers,rows,formats=None,status_col=None,critical_col=None,table_name=None):
  sh=wb.create_sheet(name); sh.sheet_view.showGridLines=False
  sh.merge_cells(start_row=1,start_column=1,end_row=2,end_column=max(1,len(headers)))
  sh.cell(1,1,title); sh.cell(1,1).font=Font(color=white,bold=True,size=16); sh.cell(1,1).fill=PatternFill('solid',fgColor=brand or dark); sh.cell(1,1).alignment=Alignment(vertical='center')
  for row in sh.iter_rows(min_row=1,max_row=2,min_col=1,max_col=len(headers)):
   for cell in row: cell.fill=PatternFill('solid',fgColor=brand or dark)
  sh.cell(3,1,f"Organizácia: {org['name'] if org else ''} · export {datetime.now().strftime('%d.%m.%Y %H:%M')}")
  sh.cell(3,1).font=Font(color=gray,size=9)
  header_row=5
  for col,hdr in enumerate(headers,1):
   cell=sh.cell(header_row,col,hdr); cell.fill=PatternFill('solid',fgColor=dark); cell.font=Font(color=white,bold=True,size=9); cell.alignment=Alignment(vertical='center',wrap_text=True)
  for ridx,row in enumerate(rows,header_row+1):
   for cidx,val in enumerate(row,1):
    cell=sh.cell(ridx,cidx,excel_safe(val)); cell.border=Border(bottom=thin); cell.alignment=Alignment(vertical='top',wrap_text=True)
   if ridx%2==0:
    for cidx in range(1,len(headers)+1): sh.cell(ridx,cidx).fill=PatternFill('solid',fgColor='FAFBFD')
  if formats:
   for col_idx,fmt in formats.items():
    for ridx in range(header_row+1,header_row+1+len(rows)): sh.cell(ridx,col_idx).number_format=fmt
  if status_col:
   for ridx in range(header_row+1,header_row+1+len(rows)):
    val=str(sh.cell(ridx,status_col).value or '')
    if val in {'Ukončené','Ukončená','Vyriešená','Prevádzka'}: sh.cell(ridx,status_col).fill=PatternFill('solid',fgColor=green)
    elif val in {'Kritická','Havária','Porucha','Mimo prevádzky'}: sh.cell(ridx,status_col).fill=PatternFill('solid',fgColor=red)
    elif val in {'Prebieha','Rieši sa','Pozastavené','Čaká na diel'}: sh.cell(ridx,status_col).fill=PatternFill('solid',fgColor=amber)
  if critical_col:
   for ridx in range(header_row+1,header_row+1+len(rows)):
    val=str(sh.cell(ridx,critical_col).value or '')
    if val=='A': sh.cell(ridx,critical_col).fill=PatternFill('solid',fgColor=red)
    elif val=='B': sh.cell(ridx,critical_col).fill=PatternFill('solid',fgColor=amber)
    elif val=='C': sh.cell(ridx,critical_col).fill=PatternFill('solid',fgColor=green)
  end=max(header_row+1,header_row+len(rows))
  if rows and table_name:
   table=Table(displayName=table_name,ref=f"A{header_row}:{get_column_letter(len(headers))}{end}")
   table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showFirstColumn=False,showLastColumn=False)
   sh.add_table(table)
  sh.freeze_panes=f'A{header_row+1}'; sh.auto_filter.ref=f"A{header_row}:{get_column_letter(len(headers))}{end}" if rows else None
  for cidx,hdr in enumerate(headers,1):
   values=[str(hdr)]+[str(row[cidx-1] or '') for row in rows[:250]]
   width=min(38,max(11,max(len(v) for v in values)+2))
   sh.column_dimensions[get_column_letter(cidx)].width=width
  sh.row_dimensions[1].height=24; sh.page_setup.orientation='landscape'; sh.page_setup.fitToWidth=1; sh.sheet_properties.pageSetUpPr.fitToPage=True
  return sh

 building_data=[[b['code'],b['name'],b['address'],b['manager'],b['status'],b['floors'],b['rooms'],float(b['area'] or 0),b['assets'],float(b['maintenance_cost'] or 0),float(b['incident_cost'] or 0),b['open_incidents'],float(b['maintenance_cost'] or 0)+float(b['incident_cost'] or 0)] for b in buildings]
 styled_sheet('Budovy','BUDOVY & PORTFÓLIO',['Kód','Budova','Adresa','Správca','Stav','Podlažia','Miestnosti','Plocha m²','Assety','Údržba','Incidenty','Otvorené incidenty','Náklady spolu'],building_data,{8:'#,##0.00',10:money_fmt,11:money_fmt,13:money_fmt},5,None,'BuildingsTable')

 asset_data=[[r['asset_id'],r['asset_tag'],r['name'],r['building'],r['floor'],r['room'],r['profession'],r['grp'],r['type'],r['manufacturer'],r['model'],r['serial'],r['system_id'],r['status'],r['criticality'],excel_date(r['installed']),excel_date(r['warranty']),r['purchase_price']] for r in assets_rows]
 styled_sheet('Assety','ASSET REGISTER',['Asset ID','Asset Tag / QR','Názov','Budova','Podlažie','Miestnosť','Profesia','Skupina','Typ','Výrobca','Model','Sériové číslo','System ID','Stav','Kritickosť','Inštalácia','Záruka do','Cena'],asset_data,{16:date_fmt,17:date_fmt,18:money_fmt},14,15,'AssetsTable')

 work_data=[[r['asset_id'],r['asset'],r['building'],r['title'],r['kind'],r['priority'],r['status'],excel_date(r['due']),r['supplier'],r['technician'],r['cost'],r['description']] for r in wo_rows]
 work=styled_sheet('Údržba','ÚDRŽBA & REVÍZIE',['Asset ID','Asset','Budova','Pracovný príkaz','Typ','Priorita','Stav','Termín','Dodávateľ','Technik','Náklad','Popis'],work_data,{8:date_fmt,11:money_fmt},7,None,'MaintenanceTable')
 for ridx in range(6,6+len(work_data)):
  due=work.cell(ridx,8).value
  if due and str(due)[:10]<date.today().isoformat() and str(work.cell(ridx,7).value)!='Ukončené':
   work.cell(ridx,8).fill=PatternFill('solid',fgColor=red); work.cell(ridx,8).font=Font(color='A61B2B',bold=True)

 incident_data=[[r['asset_id'],r['asset'],r['building'],r['title'],r['severity'],r['status'],excel_date(r['reported']),r['impact'],r['cause'],r['cost']] for r in inc_rows]
 incidents_sh=styled_sheet('Incidenty','PORUCHY & HAVÁRIE',['Asset ID','Asset','Budova','Incident','Závažnosť','Stav','Nahlásené','Dopad','Príčina','Náklad'],incident_data,{7:date_fmt,10:money_fmt},6,None,'IncidentsTable')
 for ridx in range(6,6+len(incident_data)):
  sev=str(incidents_sh.cell(ridx,5).value or '')
  incidents_sh.cell(ridx,5).fill=PatternFill('solid',fgColor=red if sev in {'Kritická','Havária','Vysoká'} else amber)

 overdue_data=[[r['asset_id'],r['asset'],r['building'],r['title'],r['priority'],excel_date(r['due']),r['technician'],r['supplier'],r['status']] for r in overdue_rows]
 overdue_sh=styled_sheet('Po termíne','ÚLOHY PO TERMÍNE',['Asset ID','Asset','Budova','Úloha','Priorita','Termín','Technik','Dodávateľ','Stav'],overdue_data,{6:date_fmt},9,None,'OverdueTable')
 for ridx in range(6,6+len(overdue_data)):
  overdue_sh.cell(ridx,6).fill=PatternFill('solid',fgColor=red); overdue_sh.cell(ridx,6).font=Font(color='A61B2B',bold=True)

 stream=io.BytesIO(); wb.save(stream); stream.seek(0)
 audit('REPORT_EXPORT',f'XLSX management report · {len(assets_rows)} assets · {len(wo_rows)} workorders · {len(inc_rows)} incidents')
 code=(org['code'] if org else 'ORG')
 return send_file(stream,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',as_attachment=True,download_name=f'GAMO_report_{code}_{date.today().isoformat()}.xlsx')

@app.route('/buildings')
def buildings(): return render_template('index.html',page='buildings',buildings=q('select b.*,(select count(*) from floors where building_id=b.id) floors,(select count(*) from assets where building_id=b.id) assets from buildings b where b.organization_id=?',(org_id(),)))
@app.route('/building/<int:i>')
def building(i):
 b=one('select * from buildings where id=? and organization_id=?',(i,org_id()))
 if not b: abort(404)
 return render_template('index.html',page='building',b=b,floors=q('select * from floors where building_id=?',(i,)),rooms=q("""select r.id,r.floor_id,r.code,r.name,coalesce(r.area,0) area,r.tenant,r.zone,f.code floor,
  (select count(*) from assets a where a.room_id=r.id) asset_count,
  (select count(*) from incidents x join assets a on a.id=x.asset_id where a.room_id=r.id and x.status not in ('Ukončená','Vyriešená')) incident_count
  from rooms r join floors f on f.id=r.floor_id where f.building_id=? order by f.id,r.code""",(i,)),assets=q('select a.*,r.code room from assets a left join rooms r on r.id=a.room_id where a.building_id=? order by a.asset_id',(i,)),documents=q('select id,name,category,mime,size,uploaded from documents where building_id=? order by id desc',(i,)))
@app.post('/building/<int:i>/document')
def upload_building_document(i):
 if not can('documents_write'): abort(403)
 if not owns_building(i): abort(404)
 category=(request.form.get('category') or 'Technická').strip()[:80]
 try:
  upload=read_safe_upload(request.files.get('document'))
  if not upload: raise ValueError('Vyber dokument na nahratie.')
 except ValueError as exc:
  flash(str(exc),'error'); return redirect(f'/building/{i}#documents')
 x('insert into documents(building_id,name,category,mime,size,data) values(?,?,?,?,?,?)',(i,upload['name'],category,upload['mime'],upload['size'],upload['data']))
 audit('DOCUMENT_UPLOAD',upload['name']); flash('Dokument bol nahratý.','success'); return redirect(f'/building/{i}#documents')

@app.get('/document/<int:i>/download')
def download_document(i):
 d=one('select d.* from documents d join buildings b on b.id=d.building_id where d.id=? and b.organization_id=?',(i,org_id()))
 if not d: abort(404)
 return send_file(io.BytesIO(d['data']),mimetype=d['mime'] or 'application/octet-stream',as_attachment=True,download_name=d['name'])

@app.post('/document/<int:i>/delete')
def delete_document(i):
 if not can('documents_write'): abort(403)
 d=one('select d.building_id,d.name from documents d join buildings b on b.id=d.building_id where d.id=? and b.organization_id=?',(i,org_id()))
 if not d: abort(404)
 x('delete from documents where id=?',(i,)); audit('DOCUMENT_DELETE',d['name']); flash('Dokument bol odstránený.','success')
 return redirect(f"/building/{d['building_id']}#documents")

@app.route('/assets')
def assets(): return render_template('index.html',page='assets',assets=q('select a.*,b.code building,r.code room from assets a join buildings b on b.id=a.building_id left join rooms r on r.id=a.room_id where b.organization_id=? order by a.asset_id',(org_id(),)))
@app.route('/asset/<int:i>')
def asset(i):
 a=one('select a.*,b.name building,f.code floor,r.code room,r.name room_name,r.area from assets a join buildings b on b.id=a.building_id left join floors f on f.id=a.floor_id left join rooms r on r.id=a.room_id where a.id=? and b.organization_id=?',(i,org_id()))
 if not a: abort(404)
 children=q("""select a.*,r.code room,r.name room_name,r.area from assets a
  join buildings b on b.id=a.building_id left join rooms r on r.id=a.room_id
  where a.parent_id=? and b.organization_id=? order by a.asset_id""",(i,org_id()))
 parent=one("""select a.id,a.asset_id,a.name,a.status from assets a join buildings b on b.id=a.building_id
  where a.id=? and b.organization_id=?""",(a['parent_id'],org_id())) if a['parent_id'] else None
 impact_rooms=len({x['room'] for x in children if x['room']}); impact_area=sum(float(x['area'] or 0) for x in children if x['room'])
 orders=q("""select w.* from workorders w join assets aa on aa.id=w.asset_id join buildings b on b.id=aa.building_id
  where w.asset_id=? and b.organization_id=? order by w.id desc""",(i,org_id()))
 incidents=q("""select x.* from incidents x join assets aa on aa.id=x.asset_id join buildings b on b.id=aa.building_id
  where x.asset_id=? and b.organization_id=? order by x.id desc""",(i,org_id()))
 events=q('select * from asset_events where asset_id=? and organization_id=? order by id desc limit 100',(i,org_id()))
 active_incidents=sum(1 for row in incidents if row['status'] not in {'Ukončená','Vyriešená'})
 asset_documents=q('select id,name,category,mime,size,uploaded from asset_documents where asset_id=? and organization_id=? order by id desc',(i,org_id()))
 asset_plan=[r for r in maintenance_plan_rows() if r['asset_db_id']==i]
 return render_template('index.html',page='asset',a=a,parent=parent,children=children,impact_rooms=impact_rooms,impact_area=impact_area,orders=orders,incidents=incidents,events=events,active_incidents=active_incidents,asset_documents=asset_documents,asset_plan=asset_plan)
@app.post('/asset/<int:i>/document')
def upload_asset_document(i):
 if not can('documents_write'): abort(403)
 a=one('select id,asset_id from assets where id=? and organization_id=?',(i,org_id()))
 if not a: abort(404)
 try:
  upload=read_safe_upload(request.files.get('document'))
  if not upload: raise ValueError('Vyber dokument na nahratie.')
 except ValueError as exc:
  flash(str(exc),'error'); return redirect(f'/asset/{i}#docs')
 category=(request.form.get('category') or 'Technická').strip()[:80]
 did=x('insert into asset_documents(organization_id,asset_id,name,category,mime,size,data) values(?,?,?,?,?,?,?)',
  (org_id(),i,upload['name'],category,upload['mime'],upload['size'],upload['data']))
 asset_event(i,'DOCUMENT_UPLOAD','Dokument assetu nahraný',f"{upload['name']} · {category}")
 audit('ASSET_DOCUMENT_UPLOAD',f"{a['asset_id']} · {upload['name']}")
 flash('Dokument assetu bol nahraný.','success')
 return redirect(f'/asset/{i}#docs')

@app.get('/asset-document/<int:i>/download')
def download_asset_document(i):
 d=one('select d.* from asset_documents d join assets a on a.id=d.asset_id where d.id=? and d.organization_id=? and a.organization_id=?',(i,org_id(),org_id()))
 if not d: abort(404)
 return send_file(io.BytesIO(bytes(d['data'])),mimetype=d['mime'] or 'application/octet-stream',as_attachment=True,download_name=d['name'])

@app.post('/asset-document/<int:i>/delete')
def delete_asset_document(i):
 if not can('documents_write'): abort(403)
 d=one('select d.id,d.asset_id,d.name,a.asset_id asset_code from asset_documents d join assets a on a.id=d.asset_id where d.id=? and d.organization_id=? and a.organization_id=?',(i,org_id(),org_id()))
 if not d: abort(404)
 x('delete from asset_documents where id=? and organization_id=?',(i,org_id()))
 asset_event(d['asset_id'],'DOCUMENT_DELETE','Dokument assetu odstránený',d['name'])
 audit('ASSET_DOCUMENT_DELETE',f"{d['asset_code']} · {d['name']}")
 flash('Dokument assetu bol odstránený.','success')
 return redirect(f"/asset/{d['asset_id']}#docs")

@app.route('/maintenance')
def maintenance():
 orders=q('select w.*,w.asset_id asset_db_id,a.asset_id asset_code,a.name asset,b.code building from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id where b.organization_id=? order by w.id desc',(org_id(),))
 today=date.today().isoformat()
 stats={'total':len(orders),'active':sum(1 for r in orders if r['status'] not in {'Ukončené','Zrušené'}),'overdue':sum(1 for r in orders if r['status'] not in {'Ukončené','Zrušené'} and r['due'] and str(r['due'])[:10]<today),'critical':sum(1 for r in orders if r['priority']=='Kritická' and r['status'] not in {'Ukončené','Zrušené'}),'completed':sum(1 for r in orders if r['status']=='Ukončené')}
 plan=maintenance_plan_rows()
 plan_stats={'rules':len(plan),'overdue':sum(1 for r in plan if r['state']=='overdue'),'soon':sum(1 for r in plan if r['state']=='soon'),'missing':sum(1 for r in plan if r['state']=='missing'),'covered':sum(1 for r in plan if r['active_workorder_id'])}
 return render_template('index.html',page='maintenance',orders=orders,maintenance_stats=stats,today_iso=today,maintenance_plan=plan,maintenance_plan_stats=plan_stats)

@app.post('/maintenance/generate')
def maintenance_generate():
 created,skipped=generate_maintenance_plan()
 if created: flash(f'Plán údržby vytvoril {created} pracovných príkazov.','success')
 else: flash('Plán je synchronizovaný. Nebolo potrebné vytvoriť nový pracovný príkaz.','success')
 return redirect('/maintenance#planner')

@app.route('/incidents')
def incidents():
 rows=q('select i.*,i.asset_id asset_db_id,a.asset_id asset_code,a.name asset,b.code building from incidents i join assets a on a.id=i.asset_id join buildings b on b.id=a.building_id where b.organization_id=? order by i.id desc',(org_id(),))
 stats={'total':len(rows),'open':sum(1 for r in rows if r['status'] not in {'Ukončená','Vyriešená'}),'critical':sum(1 for r in rows if r['severity'] in {'Kritická','Havária'} and r['status'] not in {'Ukončená','Vyriešená'}),'resolved':sum(1 for r in rows if r['status'] in {'Ukončená','Vyriešená'}),'cost':sum(float(r['cost'] or 0) for r in rows)}
 return render_template('index.html',page='incidents',incidents=rows,incident_stats=stats)
def platform_customer_portfolio():
 if not is_gamo_admin(): return []
 rows=q_system("""select o.*,
  (select count(*) from users u where u.organization_id=o.id) users_count,
  (select count(*) from users u where u.organization_id=o.id and u.status='Aktívny') active_users_count,
  (select count(*) from buildings b where b.organization_id=o.id) buildings_count,
  (select count(*) from assets a where a.organization_id=o.id) assets_count,
  (select count(*) from incidents x join assets a on a.id=x.asset_id where a.organization_id=o.id and x.status not in ('Ukončená','Vyriešená')) open_incidents_count,
  (select count(*) from workorders w join assets a on a.id=w.asset_id where a.organization_id=o.id and w.status not in ('Ukončené','Zrušené')) open_orders_count,
  (select count(*) from tickets t where t.organization_id=o.id and t.status not in ('Vyriešený','Uzavretý')) open_tickets_count,
  (select count(*) from tickets t where t.organization_id=o.id and exists(
    select 1 from ticket_messages m left join users su on su.id=m.sender_user_id
    where m.ticket_id=t.id and m.id>coalesce(t.platform_last_read_message_id,0)
      and (m.sender_user_id is null or su.organization_id=t.organization_id)
  )) unread_tickets_count,
  (select u.name from users u where u.organization_id=o.id and u.role='Administrator' order by u.id limit 1) admin_name,
  (select u.email from users u where u.organization_id=o.id and u.role='Administrator' order by u.id limit 1) admin_email,
  (select u.last_login from users u where u.organization_id=o.id and u.role='Administrator' order by u.id limit 1) admin_last_login
  from organizations o where o.code<>'GAMO' order by o.name""")
 out=[]; today_value=date.today()
 for raw in rows:
  item=dict(raw); security=security_snapshot(item['id'],True)
  item['security_score']=security['score']; item['mfa_coverage']=security['mfa_coverage']; item['backup_ok']=security['backup_ok']; item['support_active']=security['support_active']; item['security_warnings']=security['warnings']
  item['license_days']=None; item['license_expired']=False; item['license_expiring']=False
  if item.get('license_until'):
   try:
    days=(datetime.strptime(str(item['license_until'])[:10],'%Y-%m-%d').date()-today_value).days
    item['license_days']=days; item['license_expired']=days<0; item['license_expiring']=0<=days<=30
   except Exception: pass
  active=bool(item.get('status')=='Aktívny' and item.get('license_status')=='Aktívna' and not item['license_expired'])
  item['commercial_active']=active
  attention=[]
  if not active: attention.append('Licencia / organizácia')
  elif item['license_expiring']: attention.append('Licencia čoskoro expiruje')
  if item['security_score']<70: attention.append('Security')
  if not item['backup_ok']: attention.append('Backup')
  if item['support_active']: attention.append('Support otvorený')
  if int(item.get('unread_tickets_count') or 0)>0: attention.append('Nové tickety')
  item['attention']=attention; item['attention_count']=len(attention); out.append(item)
 return out

@app.route('/admin')
def admin():
 org=one('select * from organizations where id=?',(org_id(),))
 organizations=[]; platform_attention=[]; recent_platform_tickets=[]; recent_platform_access=[]
 customer_stats={'total':0,'active':0,'paused':0,'users':0,'buildings':0,'assets':0,'open_incidents':0,'open_orders':0,'open_tickets':0,'unread_tickets':0,'expiring':0,'support_open':0,'security_attention':0}
 if is_gamo_admin():
  organizations=platform_customer_portfolio()
  customer_stats={
   'total':len(organizations),'active':sum(1 for o in organizations if o['commercial_active']),'paused':sum(1 for o in organizations if not o['commercial_active']),
   'users':sum(int(o.get('users_count') or 0) for o in organizations),'buildings':sum(int(o.get('buildings_count') or 0) for o in organizations),'assets':sum(int(o.get('assets_count') or 0) for o in organizations),
   'open_incidents':sum(int(o.get('open_incidents_count') or 0) for o in organizations),'open_orders':sum(int(o.get('open_orders_count') or 0) for o in organizations),
   'open_tickets':sum(int(o.get('open_tickets_count') or 0) for o in organizations),'unread_tickets':sum(int(o.get('unread_tickets_count') or 0) for o in organizations),
   'expiring':sum(1 for o in organizations if o['license_expiring']),'support_open':sum(1 for o in organizations if o['support_active']),'security_attention':sum(1 for o in organizations if o['security_score']<70)
  }
  for o in organizations:
   for issue in o['attention'][:3]: platform_attention.append({'organization_id':o['id'],'code':o['code'],'name':o['name'],'issue':issue,'score':o['security_score']})
  platform_attention=platform_attention[:12]
  recent_platform_tickets=[dict(r) for r in q_system("""select t.id,t.ticket_no,t.subject,t.priority,t.status,t.updated,o.code organization_code,o.name organization_name,
   (select count(*) from ticket_messages m left join users su on su.id=m.sender_user_id
    where m.ticket_id=t.id and m.id>coalesce(t.platform_last_read_message_id,0)
      and (m.sender_user_id is null or su.organization_id=t.organization_id)) platform_unread
   from tickets t join organizations o on o.id=t.organization_id
   where o.code<>'GAMO'
   order by case when t.status='Nový' then 0 when t.status='Otvorený' then 1 else 2 end,t.updated desc limit 7""")]
  recent_platform_access=[dict(r) for r in q_system("""select l.created,l.action,l.actor_name,l.reason,o.code organization_code,o.name organization_name
   from customer_access_log l join organizations o on o.id=l.target_organization_id
   where o.code<>'GAMO' and l.action<>'GAMO_METADATA_VIEW'
   order by l.id desc limit 7""")]
 audit_rows=q('select * from audit_log where organization_id=? order by id desc limit 20',(org_id(),))
 return render_template('index.html',page='admin',
  users=q('select * from users where organization_id=? order by name',(org_id(),)),
  buildings=q('select * from buildings where organization_id=? order by name',(org_id(),)),
  audit_rows=audit_rows,organizations=organizations,customer_stats=customer_stats,current_admin_org=org,
  platform_attention=platform_attention,recent_platform_tickets=recent_platform_tickets,recent_platform_access=recent_platform_access)

@app.get('/platform/customers/export.xlsx')
def platform_customers_export():
 if not is_gamo_admin(): abort(403)
 organizations=platform_customer_portfolio()
 wb=Workbook(); ws=wb.active; ws.title='Zákazníci'
 brand='17365D'; white='FFFFFF'; gray='667085'; green='E4F5EC'; red='FBE9ED'; amber='FFF4D9'
 ws.sheet_view.showGridLines=False
 ws.merge_cells('A1:P2'); ws['A1']='GAMO SUPER ADMIN · CUSTOMER PORTFOLIO'
 ws['A1'].font=Font(color=white,bold=True,size=18); ws['A1'].fill=PatternFill('solid',fgColor=brand); ws['A1'].alignment=Alignment(vertical='center')
 for row in ws['A1:P2']:
  for cell in row: cell.fill=PatternFill('solid',fgColor=brand)
 ws.merge_cells('A3:P3'); ws['A3']=f"Generované {datetime.now().strftime('%d.%m.%Y %H:%M')} · iba platformové metadáta"; ws['A3'].font=Font(color=gray,size=9)
 headers=['Kód','Zákazník','Plán','Organizácia','Licencia','Platnosť do','Dní','Používatelia','Budovy','Assety','Otvorené incidenty','Otvorené úlohy','Tickety','Unread','Security Score','MFA %']
 row0=5
 for col,h in enumerate(headers,1):
  cell=ws.cell(row0,col,h); cell.fill=PatternFill('solid',fgColor=brand); cell.font=Font(color=white,bold=True,size=9)
 for ridx,o in enumerate(organizations,row0+1):
  values=[o['code'],o['name'],o['plan'],o['status'],o['license_status'],o['license_until'] or '',o['license_days'] if o['license_days'] is not None else '',o['users_count'],o['buildings_count'],o['assets_count'],o['open_incidents_count'],o['open_orders_count'],o['open_tickets_count'],o['unread_tickets_count'],o['security_score'],o['mfa_coverage']]
  for cidx,val in enumerate(values,1): ws.cell(ridx,cidx,val).alignment=Alignment(vertical='top')
  for cidx,val in enumerate(values,1): ws.cell(ridx,cidx).value=val
  ws.cell(ridx,15).fill=PatternFill('solid',fgColor=green if o['security_score']>=80 else (amber if o['security_score']>=60 else red))
  if not o['commercial_active']: ws.cell(ridx,5).fill=PatternFill('solid',fgColor=red)
  elif o['license_expiring']: ws.cell(ridx,6).fill=PatternFill('solid',fgColor=amber)
 if organizations:
  table=Table(displayName='GamoCustomerPortfolio',ref=f'A{row0}:P{row0+len(organizations)}'); table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showFirstColumn=False,showLastColumn=False); ws.add_table(table)
 ws.freeze_panes='A6'
 widths=[13,30,14,15,15,15,9,13,10,10,18,16,12,10,15,10]
 for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
 ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.fitToWidth=1
 summary=wb.create_sheet('Súhrn'); summary.sheet_view.showGridLines=False
 summary['A1']='GAMO PLATFORM OVERVIEW'; summary['A1'].font=Font(size=18,bold=True,color=white); summary['A1'].fill=PatternFill('solid',fgColor=brand); summary.merge_cells('A1:D2')
 for row in summary['A1:D2']:
  for cell in row: cell.fill=PatternFill('solid',fgColor=brand)
 metrics=[('Zákazníci',len(organizations)),('Aktívni',sum(1 for o in organizations if o['commercial_active'])),('Licencie ≤30 dní',sum(1 for o in organizations if o['license_expiring'])),('Security <70',sum(1 for o in organizations if o['security_score']<70)),('Unread tickety',sum(int(o['unread_tickets_count'] or 0) for o in organizations)),('Aktívny support',sum(1 for o in organizations if o['support_active']))]
 for idx,(label,value) in enumerate(metrics,4): summary.cell(idx,1,label).font=Font(bold=True,color=gray); summary.cell(idx,2,value).font=Font(bold=True,size=14)
 summary.column_dimensions['A'].width=24; summary.column_dimensions['B'].width=16
 stream=io.BytesIO(); wb.save(stream); stream.seek(0)
 audit('PLATFORM_CUSTOMERS_EXPORT',f'{len(organizations)} zákazníkov')
 return send_file(stream,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',as_attachment=True,download_name=f'GAMO_customer_portfolio_{date.today().isoformat()}.xlsx')


def audit_for_org(target_org_id,action,detail=''):
 try:
  x_system("insert into audit_log(user_id,user_name,action,detail,ip,organization_id) values(?,?,?,?,?,?)",(
   session.get('user_id'),session.get('user_name','Systém'),action,(detail or '')[:1000],
   request.headers.get('X-Forwarded-For',request.remote_addr or ''),target_org_id
  ))
 except Exception:
  pass

def security_snapshot(oid,privileged=False):
 read_one=one_system if privileged else one
 org=read_one('select * from organizations where id=?',(oid,))
 if not org: return {'score':0,'warnings':['Organizácia neexistuje.']}
 active=int(read_one("select count(*) n from users where organization_id=? and status='Aktívny'",(oid,))['n'])
 mfa=int(read_one("select count(*) n from users where organization_id=? and status='Aktívny' and mfa_enabled=?",(oid,True if USING_POSTGRES else 1))['n'])
 admins=int(read_one("select count(*) n from users where organization_id=? and role='Administrator' and status='Aktívny'",(oid,))['n'])
 admin_mfa=int(read_one("select count(*) n from users where organization_id=? and role='Administrator' and status='Aktívny' and mfa_enabled=?",(oid,True if USING_POSTGRES else 1))['n'])
 coverage=round((mfa/active*100),1) if active else 100.0
 warnings=[]; score=0
 if USING_POSTGRES: score+=25
 else: warnings.append('Desktop/SQLite režim nemá databázové RLS.')
 if bool(org['mfa_required']): score+=20
 else: warnings.append('MFA nie je povinné pre organizáciu.')
 if admins and admin_mfa==admins: score+=20
 else: warnings.append(f'{max(0,admins-admin_mfa)} aktívnych administrátorov nemá MFA.')
 score+=round(15*(coverage/100))
 if not support_access_active(org): score+=10
 else: warnings.append('GAMO support prístup je momentálne aktívny.')
 backup_ok=(org['backup_last_verified_status']=='OK' and bool(org['backup_last_verified_at']))
 if backup_ok: score+=10
 else: warnings.append('Integrita posledného zákazníckeho backupu nebola overená.')
 return {
  'score':min(100,int(score)),'active_users':active,'mfa_users':mfa,'mfa_coverage':coverage,
  'admins':admins,'admin_mfa':admin_mfa,'mfa_required':bool(org['mfa_required']),
  'support_active':support_access_active(org),'backup_ok':backup_ok,
  'backup_verified_at':org['backup_last_verified_at'],'warnings':warnings
 }

def apply_retention_if_due(oid,force=False):
 org=one('select * from organizations where id=?',(oid,))
 if not org: return 0
 last=org['retention_last_run']
 if last and not force:
  try:
   dt=datetime.fromisoformat(str(last).replace('Z','+00:00'))
   now=datetime.now(dt.tzinfo) if dt.tzinfo else datetime.now()
   if now-dt<timedelta(hours=24): return 0
  except Exception:
   pass
 days=max(365,min(3650,int(org['retention_days'] or 3650)))
 cutoff=(datetime.utcnow()-timedelta(days=days)).isoformat(timespec='seconds')+'Z'
 deleted=0
 with con() as db:
  cur=db.execute(_sql('delete from auth_events where organization_id=? and created<?'),(oid,cutoff)); deleted+=max(0,cur.rowcount or 0)
  cur=db.execute(_sql('delete from customer_access_log where target_organization_id=? and created<?'),(oid,cutoff)); deleted+=max(0,cur.rowcount or 0)
  db.execute(_sql('update organizations set retention_last_run=CURRENT_TIMESTAMP where id=?'),(oid,))
  db.commit()
 audit('RETENTION_RUN',f'{days} dní · odstránených {deleted} bezpečnostných logov')
 return deleted

def create_organization_backup_archive(oid,privileged=False):
 read_one=one_system if privileged else one
 read_all=q_system if privileged else q
 org=read_one('select * from organizations where id=?',(oid,))
 if not org: abort(404)
 data={}
 data['buildings']=[dict(r) for r in read_all('select * from buildings where organization_id=? order by id',(oid,))]
 data['floors']=[dict(r) for r in read_all('select f.* from floors f join buildings b on b.id=f.building_id where b.organization_id=? order by f.id',(oid,))]
 data['rooms']=[dict(r) for r in read_all('select r.* from rooms r join floors f on f.id=r.floor_id join buildings b on b.id=f.building_id where b.organization_id=? order by r.id',(oid,))]
 data['assets']=[dict(r) for r in read_all('select a.* from assets a where a.organization_id=? order by a.id',(oid,))]
 data['workorders']=[dict(r) for r in read_all('select w.* from workorders w join assets a on a.id=w.asset_id where a.organization_id=? order by w.id',(oid,))]
 data['incidents']=[dict(r) for r in read_all('select x.* from incidents x join assets a on a.id=x.asset_id where a.organization_id=? order by x.id',(oid,))]
 data['users']=[dict(r) for r in read_all('select id,name,email,role,status,last_login,organization_id,mfa_enabled,mfa_enabled_at,erased_at from users where organization_id=? order by id',(oid,))]
 data['settings']=[dict(r) for r in read_all('select k,v from organization_settings where organization_id=? order by k',(oid,))]
 data['asset_events']=[dict(r) for r in read_all('select * from asset_events where organization_id=? order by id',(oid,))]
 data['access_log']=[dict(r) for r in read_all('select * from customer_access_log where target_organization_id=? order by id',(oid,))]
 data['auth_events']=[dict(r) for r in read_all('select * from auth_events where organization_id=? order by id',(oid,))]
 data['privacy_requests']=[dict(r) for r in read_all('select * from privacy_requests where organization_id=? order by id',(oid,))]
 data['tickets']=[dict(r) for r in read_all('select * from tickets where organization_id=? order by id',(oid,))]
 data['ticket_messages']=[dict(r) for r in read_all('select * from ticket_messages where organization_id=? order by id',(oid,))]
 docs=read_all('select id,building_id,name,category,mime,size,uploaded,data from documents where building_id in (select id from buildings where organization_id=?) order by id',(oid,))
 asset_docs=read_all('select id,asset_id,name,category,mime,size,uploaded,data from asset_documents where organization_id=? order by id',(oid,))
 ticket_files=read_all('select id,ticket_id,message_id,sender_user_id,name,mime,size,uploaded,data from ticket_attachments where organization_id=? order by id',(oid,))
 files={}
 for name,rows in data.items():
  files[f'data/{name}.json']=json.dumps(rows,ensure_ascii=False,indent=2,default=str).encode('utf-8')
 doc_meta=[]
 for d in docs:
  item={k:d[k] for k in ['id','building_id','name','category','mime','size','uploaded']}
  safe=os.path.basename(d['name'] or f"document_{d['id']}")
  path=f"documents/{d['id']}_{safe}"
  item['archive_path']=path; doc_meta.append(item)
  files[path]=bytes(d['data']) if d['data'] is not None else b''
 files['data/documents.json']=json.dumps(doc_meta,ensure_ascii=False,indent=2,default=str).encode('utf-8')
 asset_doc_meta=[]
 for d in asset_docs:
  item={k:d[k] for k in ['id','asset_id','name','category','mime','size','uploaded']}
  safe=os.path.basename(d['name'] or f"asset_document_{d['id']}")
  path=f"asset_documents/{d['id']}_{safe}"
  item['archive_path']=path; asset_doc_meta.append(item)
  files[path]=bytes(d['data']) if d['data'] is not None else b''
 files['data/asset_documents.json']=json.dumps(asset_doc_meta,ensure_ascii=False,indent=2,default=str).encode('utf-8')
 ticket_file_meta=[]
 for d in ticket_files:
  item={k:d[k] for k in ['id','ticket_id','message_id','sender_user_id','name','mime','size','uploaded']}
  safe=os.path.basename(d['name'] or f"ticket_attachment_{d['id']}")
  path=f"ticket_attachments/{d['id']}_{safe}"
  item['archive_path']=path; ticket_file_meta.append(item)
  files[path]=bytes(d['data']) if d['data'] is not None else b''
 files['data/ticket_attachments.json']=json.dumps(ticket_file_meta,ensure_ascii=False,indent=2,default=str).encode('utf-8')
 checksums={path:hashlib.sha256(blob).hexdigest() for path,blob in files.items()}
 manifest={
  'format':'GAMO_ORGANIZATION_BACKUP_V3','generated_at':datetime.utcnow().isoformat(timespec='seconds')+'Z',
  'organization':{k:org[k] for k in ['id','code','name','status','plan','license_status','license_until','branding_name','brand_color','brand_tagline','privacy_contact','data_region','retention_days','mfa_required','created'] if k in org.keys()},
  'privacy':{'password_hashes_included':False,'mfa_secrets_included':False,'support_access_active':support_access_active(org)},
  'counts':{name:len(rows) for name,rows in data.items()} | {'documents':len(doc_meta),'asset_documents':len(asset_doc_meta),'ticket_attachments':len(ticket_file_meta)},
  'checksums_sha256':checksums,
  'note':'Password hashes, TOTP secrets and recovery codes are intentionally excluded.'
 }
 stream=io.BytesIO()
 with zipfile.ZipFile(stream,'w',zipfile.ZIP_DEFLATED) as z:
  z.writestr('manifest.json',json.dumps(manifest,ensure_ascii=False,indent=2,default=str))
  for path,blob in files.items(): z.writestr(path,blob)
 filename=f"GAMO_backup_{org['code']}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
 return stream.getvalue(),filename,manifest

def verify_organization_backup_bytes(blob):
 try:
  with zipfile.ZipFile(io.BytesIO(blob),'r') as z:
   manifest=json.loads(z.read('manifest.json').decode('utf-8'))
   if manifest.get('format')!='GAMO_ORGANIZATION_BACKUP_V3': return False,'Neplatný formát backupu.'
   checks=manifest.get('checksums_sha256') or {}
   required={'data/buildings.json','data/assets.json','data/users.json','data/documents.json','data/asset_documents.json','data/ticket_attachments.json'}
   if not required.issubset(set(checks)): return False,'Backup nemá všetky povinné dátové súbory.'
   for path,expected in checks.items():
    if path not in z.namelist(): return False,f'Chýba {path}.'
    if hashlib.sha256(z.read(path)).hexdigest()!=expected: return False,f'Checksum nesedí: {path}.'
  return True,'Všetky súbory a SHA-256 checksumy sú v poriadku.'
 except Exception as exc:
  return False,f'Backup sa nepodarilo overiť: {exc}'

@app.get('/privacy')
def privacy_center():
 if support_mode(): abort(403)
 org=one('select * from organizations where id=?',(org_id(),))
 if not org or session.get('user_role')!='Administrator': abort(403)
 try: apply_retention_if_due(org_id(),False)
 except Exception: pass
 org=one('select * from organizations where id=?',(org_id(),))
 access_rows=q('select * from customer_access_log where target_organization_id=? order by id desc limit 50',(org_id(),))
 auth_rows=q('select * from auth_events where organization_id=? order by id desc limit 30',(org_id(),))
 privacy_users=q('select id,name,email,role,status,last_login,mfa_enabled,mfa_enabled_at,erased_at from users where organization_id=? order by case when role=\'Administrator\' then 0 else 1 end,name',(org_id(),))
 privacy_requests=q('select p.*,u.name subject_name from privacy_requests p left join users u on u.id=p.subject_user_id where p.organization_id=? order by p.id desc limit 30',(org_id(),))
 return render_template('index.html',page='privacy',privacy_org=org,support_active=support_access_active(org),
  access_rows=access_rows,auth_rows=auth_rows,privacy_users=privacy_users,privacy_requests=privacy_requests,
  security=security_snapshot(org_id()),rls_enabled=USING_POSTGRES)

@app.post('/privacy/settings')
def privacy_settings():
 if support_mode(): abort(403)
 org=one('select * from organizations where id=?',(org_id(),))
 if not org or session.get('user_role')!='Administrator': abort(403)
 contact=(request.form.get('privacy_contact') or '').strip()[:160]
 try: retention=max(365,min(3650,int(request.form.get('retention_days') or 3650)))
 except (TypeError,ValueError): retention=3650
 mfa_required=request.form.get('mfa_required')=='1'
 x('update organizations set privacy_contact=?,retention_days=?,mfa_required=? where id=?',(contact,retention,True if (USING_POSTGRES and mfa_required) else (1 if mfa_required else (False if USING_POSTGRES else 0)),org_id()))
 audit('PRIVACY_SETTINGS',f'retention={retention} · mfa_required={mfa_required}')
 current=one('select mfa_enabled from users where id=?',(session.get('user_id'),))
 if mfa_required and current and not bool(current['mfa_enabled']):
  flash('Povinné MFA bolo zapnuté. Najprv zabezpeč svoj administrátorský účet.','success')
  return redirect('/account/mfa')
 flash('Nastavenia súkromia a bezpečnostnej politiky boli uložené.','success')
 return redirect('/privacy')

@app.post('/privacy/retention/run')
def privacy_retention_run():
 if support_mode() or session.get('user_role')!='Administrator': abort(403)
 deleted=apply_retention_if_due(org_id(),True)
 flash(f'Retenčná politika bola vykonaná. Odstránených bezpečnostných logov: {deleted}.','success')
 return redirect('/privacy')

@app.post('/privacy/backup/verify')
def privacy_backup_verify():
 if support_mode() or session.get('user_role')!='Administrator': abort(403)
 blob,_,_=create_organization_backup_archive(org_id(),False)
 ok,detail=verify_organization_backup_bytes(blob)
 status='OK' if ok else 'FAILED'
 x('update organizations set backup_last_verified_at=CURRENT_TIMESTAMP,backup_last_verified_status=? where id=?',(status,org_id()))
 audit('BACKUP_VERIFY',f'{status} · {detail}')
 flash(('Backup integrity: '+detail), 'success' if ok else 'error')
 return redirect('/privacy')

@app.get('/privacy/user/<int:i>/export.json')
def privacy_user_export(i):
 if support_mode() or session.get('user_role')!='Administrator': abort(403)
 u=one('select id,name,email,role,status,last_login,mfa_enabled,mfa_enabled_at,erased_at,organization_id from users where id=? and organization_id=?',(i,org_id()))
 if not u: abort(404)
 payload={
  'format':'GAMO_PERSONAL_DATA_EXPORT_V1','generated_at':datetime.utcnow().isoformat(timespec='seconds')+'Z',
  'account':dict(u),
  'audit_events':[dict(r) for r in q('select created,action,detail,ip from audit_log where organization_id=? and user_id=? order by id',(org_id(),i))],
  'asset_events':[dict(r) for r in q('select created,event_type,title,detail from asset_events where organization_id=? and user_id=? order by id',(org_id(),i))],
  'support_events':[dict(r) for r in q('select created,action,reason,ip from customer_access_log where target_organization_id=? and actor_user_id=? order by id',(org_id(),i))],
  'auth_events':[dict(r) for r in q('select created,event,success,detail,ip,user_agent from auth_events where organization_id=? and user_id=? order by id',(org_id(),i))]
 }
 x('insert into privacy_requests(organization_id,subject_user_id,request_type,status,requested_by,note,completed_at) values(?,?,?,?,?,?,CURRENT_TIMESTAMP)',(org_id(),i,'EXPORT','COMPLETED',session.get('user_id'),'Osobné údaje exportované administrátorom.'))
 audit('GDPR_USER_EXPORT',f'user_id={i}')
 response=app.response_class(json.dumps(payload,ensure_ascii=False,indent=2,default=str),mimetype='application/json')
 response.headers['Content-Disposition']=f'attachment; filename=GAMO_personal_data_{i}_{date.today().isoformat()}.json'
 response.headers['Cache-Control']='no-store'
 return response

@app.post('/privacy/user/<int:i>/anonymize')
def privacy_user_anonymize(i):
 if support_mode() or session.get('user_role')!='Administrator': abort(403)
 u=one('select * from users where id=? and organization_id=?',(i,org_id()))
 if not u: abort(404)
 if i==session.get('user_id'):
  flash('Vlastný aktívny účet nie je možné anonymizovať.','error'); return redirect('/privacy')
 if (request.form.get('confirm_text') or '').strip().upper()!='ANONYMIZE':
  flash('Pre anonymizáciu napíš presne ANONYMIZE.','error'); return redirect('/privacy')
 if u['role']=='Administrator' and u['status']=='Aktívny' and one("select count(*) n from users where organization_id=? and role='Administrator' and status='Aktívny'",(org_id(),))['n']<=1:
  flash('Posledného aktívneho administrátora nie je možné anonymizovať.','error'); return redirect('/privacy')
 anon_name=f'Anonymizovaný používateľ #{i}'; anon_email=f'deleted+{org_id()}-{i}@anonymized.invalid'
 with con() as db:
  db.execute(_sql("""update users set name=?,email=?,role='Viewer',status='Neaktívny',
   password_hash=?,last_login=NULL,must_change_password=?,mfa_secret=NULL,mfa_recovery_codes=NULL,
   mfa_enabled=?,mfa_enabled_at=NULL,erased_at=CURRENT_TIMESTAMP where id=? and organization_id=?"""),
   (anon_name,anon_email,generate_password_hash(secrets.token_urlsafe(32)),False if USING_POSTGRES else 0,False if USING_POSTGRES else 0,i,org_id()))
  db.execute(_sql("update audit_log set user_name='Anonymizovaný používateľ' where organization_id=? and user_id=?"),(org_id(),i))
  db.execute(_sql("update asset_events set user_name='Anonymizovaný používateľ' where organization_id=? and user_id=?"),(org_id(),i))
  db.execute(_sql("update customer_access_log set actor_name='Anonymizovaný používateľ' where target_organization_id=? and actor_user_id=?"),(org_id(),i))
  db.execute(_sql('insert into privacy_requests(organization_id,subject_user_id,request_type,status,requested_by,note,completed_at) values(?,?,?,?,?,?,CURRENT_TIMESTAMP)'),(org_id(),i,'ANONYMIZE','COMPLETED',session.get('user_id'),'Účet a priame identifikátory boli anonymizované; prevádzkové záznamy ostali zachované.'))
  db.commit()
 audit('GDPR_USER_ANONYMIZE',f'user_id={i}')
 flash('Používateľ bol anonymizovaný. Prevádzková história ostala zachovaná bez priamych identifikátorov.','success')
 return redirect('/privacy')

@app.post('/privacy/support-access')
def privacy_support_access():
 if support_mode(): abort(403)
 org=one('select * from organizations where id=?',(org_id(),))
 if not org or org['code']=='GAMO' or session.get('user_role')!='Administrator': abort(403)
 action=request.form.get('action')
 reason=(request.form.get('reason') or '').strip()[:500]
 if action!='revoke' and len(reason)<5:
  flash('Pri povolení support prístupu uveď dôvod alebo číslo ticketu.','error'); return redirect('/privacy')
 if action=='revoke':
  x('update organizations set support_access_enabled=?,support_access_until=? where id=?',(False if USING_POSTGRES else 0,None,org_id()))
  customer_access(org_id(),'SUPPORT_ACCESS_REVOKED',reason or 'Prístup GAMO bol odvolaný zákazníkom.')
  audit('SUPPORT_ACCESS_REVOKED',reason)
  flash('Prístup GAMO k zákazníckym dátam bol okamžite odvolaný.','success')
 else:
  try: hours=int(request.form.get('hours') or 24)
  except (TypeError,ValueError): hours=24
  if hours not in {1,24,168}: hours=24
  until=(datetime.utcnow()+timedelta(hours=hours)).isoformat(timespec='seconds')+'Z'
  x('update organizations set support_access_enabled=?,support_access_until=? where id=?',(True if USING_POSTGRES else 1,until,org_id()))
  customer_access(org_id(),'SUPPORT_ACCESS_GRANTED',reason or f'Dočasný prístup na {hours} h.')
  audit('SUPPORT_ACCESS_GRANTED',f'{hours}h · {reason}')
  flash(f'GAMO support má dočasný prístup na {hours} hodín.','success')
 return redirect('/privacy')

@app.get('/backup/my')
def backup_my_organization():
 if support_mode(): abort(403)
 if session.get('user_role')!='Administrator': abort(403)
 customer_access(org_id(),'CUSTOMER_DATA_EXPORT','Export vytvorený administrátorom organizácie.')
 return build_organization_backup(org_id())

@app.get('/privacy/access-log.json')
def privacy_access_log_export():
 if support_mode(): abort(403)
 if session.get('user_role')!='Administrator': abort(403)
 org=one('select code,name from organizations where id=?',(org_id(),))
 if not org: abort(404)
 rows=[dict(r) for r in q('select created,actor_name,action,reason,ip from customer_access_log where target_organization_id=? order by id desc',(org_id(),))]
 payload={
  'organization':{'code':org['code'],'name':org['name']},
  'generated_at':datetime.utcnow().isoformat(timespec='seconds')+'Z',
  'events':rows
 }
 audit('PRIVACY_ACCESS_LOG_EXPORT',f"{len(rows)} udalostí")
 response=app.response_class(json.dumps(payload,ensure_ascii=False,indent=2,default=str),mimetype='application/json')
 response.headers['Content-Disposition']=f'attachment; filename=GAMO_access_log_{org["code"]}_{date.today().isoformat()}.json'
 return response

@app.post('/platform/customer/<int:i>/support-enter')
def platform_customer_support_enter(i):
 if not is_gamo_admin(): abort(403)
 customer=one_system('select * from organizations where id=? and code<>?',(i,'GAMO'))
 if not customer: abort(404)
 if not support_access_active(customer):
  flash('Zákazník nemá aktívny dočasný support prístup.','error')
  return redirect(f'/platform/customer/{i}')
 session['support_target_org_id']=customer['id']
 session['support_target_name']=customer['name']
 session['support_target_code']=customer['code']
 customer_access(i,'GAMO_SUPPORT_SESSION_ENTER','GAMO vstúpilo do auditovaného support režimu zákazníka.')
 audit('SUPPORT_SESSION_ENTER',f"{customer['code']} · {customer['name']}")
 return redirect('/')

@app.post('/support/exit')
def support_exit():
 target_id=session.get('support_target_org_id')
 if not target_id: return redirect('/')
 customer=one_system('select id,code,name from organizations where id=?',(target_id,))
 if customer:
  customer_access(target_id,'GAMO_SUPPORT_SESSION_EXIT','GAMO ukončilo auditovaný support režim.')
  audit('SUPPORT_SESSION_EXIT',f"{customer['code']} · {customer['name']}")
 session.pop('support_target_org_id',None); session.pop('support_target_name',None); session.pop('support_target_code',None)
 return redirect(f'/platform/customer/{target_id}' if customer else '/admin')

@app.get('/platform/customer/<int:i>/backup')
def platform_customer_backup(i):
 if not is_gamo_admin(): abort(403)
 customer=one_system('select * from organizations where id=? and code<>?',(i,'GAMO'))
 if not customer: abort(404)
 if not support_access_active(customer):
  flash('Zákazník nemá aktívny súhlas na prístup GAMO k dátam. Backup nie je dostupný.','error')
  return redirect(f'/platform/customer/{i}')
 customer_access(i,'GAMO_BACKUP_EXPORT','GAMO administrátor stiahol zákaznícky export počas aktívneho support prístupu.')
 return build_organization_backup(i,privileged=True)

def build_organization_backup(oid,privileged=False):
 blob,filename,manifest=create_organization_backup_archive(oid,privileged)
 if privileged: audit_for_org(oid,'BACKUP_EXPORT',f"{manifest['organization']['code']} · GAMO support export")
 else: audit('BACKUP_EXPORT',f"{manifest['organization']['code']} · organization export")
 response=send_file(io.BytesIO(blob),mimetype='application/zip',as_attachment=True,download_name=filename)
 response.headers['Cache-Control']='no-store'
 return response

@app.get('/platform/customer/<int:i>')
def platform_customer_detail(i):
 if not is_gamo_admin(): abort(403)
 customer=one_system('select * from organizations where id=? and code<>?',(i,'GAMO'))
 if not customer: abort(404)
 support_active=support_access_active(customer)
 customer_access(i,'GAMO_METADATA_VIEW','GAMO otvorilo licenčné a agregované metadáta zákazníka.')
 customer_stats={
  'users':one_system('select count(*) n from users where organization_id=?',(i,))['n'],
  'active_users':one_system("select count(*) n from users where organization_id=? and status='Aktívny'",(i,))['n'],
  'buildings':one_system('select count(*) n from buildings where organization_id=?',(i,))['n'],
  'assets':one_system('select count(*) n from assets where organization_id=?',(i,))['n'],
  'open_incidents':one_system("select count(*) n from incidents x join assets a on a.id=x.asset_id where a.organization_id=? and x.status not in ('Ukončená','Vyriešená')",(i,))['n'],
  'open_orders':one_system("select count(*) n from workorders w join assets a on a.id=w.asset_id where a.organization_id=? and w.status not in ('Ukončené','Zrušené')",(i,))['n'],
  'maintenance_cost':one_system('select coalesce(sum(w.cost),0) n from workorders w join assets a on a.id=w.asset_id where a.organization_id=?',(i,))['n'],
  'incident_cost':one_system('select coalesce(sum(x.cost),0) n from incidents x join assets a on a.id=x.asset_id where a.organization_id=?',(i,))['n']
 }
 customer_stats['total_cost']=customer_stats['maintenance_cost']+customer_stats['incident_cost']
 license_days=None
 if customer['license_until']:
  try: license_days=(datetime.strptime(str(customer['license_until'])[:10],'%Y-%m-%d').date()-date.today()).days
  except Exception: license_days=None
 customer_users=[]; customer_buildings=[]; recent_incidents=[]; recent_orders=[]; customer_audit=[]
 if support_active:
  customer_access(i,'GAMO_SUPPORT_DATA_VIEW','Customer 360 sensitive data viewed during active support access.')
  customer_users=q_system('select id,name,email,role,status,last_login,mfa_enabled,mfa_enabled_at,erased_at from users where organization_id=? order by case when role=\'Administrator\' then 0 else 1 end,name',(i,))
  customer_buildings=q_system("""select b.*,
   (select count(*) from floors f where f.building_id=b.id) floors_count,
   (select count(*) from assets a where a.building_id=b.id) assets_count,
   (select count(*) from incidents x join assets a on a.id=x.asset_id where a.building_id=b.id and x.status not in ('Ukončená','Vyriešená')) incidents_count
   from buildings b where b.organization_id=? order by b.name""",(i,))
  recent_incidents=q_system("""select x.*,a.asset_id,a.name asset,b.name building from incidents x
   join assets a on a.id=x.asset_id join buildings b on b.id=a.building_id
   where a.organization_id=? order by x.id desc limit 6""",(i,))
  recent_orders=q_system("""select w.*,a.asset_id,a.name asset,b.name building from workorders w
   join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id
   where a.organization_id=? order by w.id desc limit 6""",(i,))
  customer_audit=q_system('select * from audit_log where organization_id=? order by id desc limit 10',(i,))
 last_backup=one_system("select created from audit_log where organization_id=? and action='BACKUP_EXPORT' order by id desc limit 1",(i,))
 last_login=one_system("select max(last_login) last_login from users where organization_id=?",(i,))
 access_rows=q_system('select * from customer_access_log where target_organization_id=? order by id desc limit 20',(i,))
 customer_security=security_snapshot(i,True)
 limits=PLAN_LIMITS.get((customer['plan'] or 'BASIC'),PLAN_LIMITS['BASIC'])
 return render_template('index.html',page='customer',customer=customer,customer_stats=customer_stats,
  customer_users=customer_users,customer_buildings=customer_buildings,recent_incidents=recent_incidents,
  recent_orders=recent_orders,license_days=license_days,customer_limits=limits,customer_audit=customer_audit,
  customer_last_backup=(last_backup['created'] if last_backup else None),customer_last_login=(last_login['last_login'] if last_login else None),
  support_access=support_active,customer_access_rows=access_rows,customer_security=customer_security,rls_enabled=USING_POSTGRES)

@app.post('/platform/customer/<int:i>/reset-admin-password')
def platform_customer_reset_admin_password(i):
 if not is_gamo_admin(): abort(403)
 customer=one_system('select * from organizations where id=? and code<>?',(i,'GAMO'))
 if not customer: abort(404)
 if not support_access_active(customer):
  flash('Reset hesla vyžaduje aktívny support prístup udelený zákazníkom.','error')
  return redirect(f'/platform/customer/{i}#customerUsers')
 password=request.form.get('password') or ''
 if len(password)<10:
  flash('Dočasné heslo musí mať aspoň 10 znakov.','error'); return redirect(f'/platform/customer/{i}#customerUsers')
 admin_user=one_system("select id,name,email from users where organization_id=? and role='Administrator' order by id limit 1",(i,))
 if not admin_user:
  flash('Zákazník nemá administrátorský účet.','error'); return redirect(f'/platform/customer/{i}#customerUsers')
 x_system('update users set password_hash=?,status=?,must_change_password=? where id=?',(generate_password_hash(password),'Aktívny',True if USING_POSTGRES else 1,admin_user['id']))
 revoke_user_devices(admin_user['id'])
 customer_access(i,'GAMO_ADMIN_PASSWORD_RESET',f"Reset hesla administrátora {admin_user['email']}.")
 audit('CUSTOMER_ADMIN_PASSWORD_RESET',f"{customer['code']} · {admin_user['email']}")
 flash('Dočasné heslo zákazníckeho administrátora bolo zmenené.','success')
 return redirect(f'/platform/customer/{i}#customerUsers')

@app.post('/platform/customer/<int:i>/reset-admin-mfa')
def platform_customer_reset_admin_mfa(i):
 if not is_gamo_admin(): abort(403)
 customer=one_system('select * from organizations where id=? and code<>?',(i,'GAMO'))
 if not customer: abort(404)
 if not support_access_active(customer):
  flash('Reset MFA vyžaduje aktívny support prístup udelený zákazníkom.','error')
  return redirect(f'/platform/customer/{i}#customerUsers')
 admin_user=one_system("select id,email from users where organization_id=? and role='Administrator' order by id limit 1",(i,))
 if not admin_user: abort(404)
 x_system('update users set mfa_secret=?,mfa_recovery_codes=?,mfa_enabled=?,mfa_enabled_at=? where id=?',(None,None,False if USING_POSTGRES else 0,None,admin_user['id']))
 revoke_user_devices(admin_user['id'])
 customer_access(i,'GAMO_ADMIN_MFA_RESET',f"Reset MFA administrátora {admin_user['email']}.")
 audit_for_org(i,'CUSTOMER_ADMIN_MFA_RESET',f"Reset MFA {admin_user['email']}")
 flash('MFA administrátora bolo resetované. Ak organizácia vyžaduje MFA, pri ďalšom prihlásení ho musí nastaviť znova.','success')
 return redirect(f'/platform/customer/{i}#customerUsers')

@app.post('/platform/customer/<int:i>/branding')
def platform_customer_branding(i):
 if not is_gamo_admin(): abort(403)
 customer=one_system('select * from organizations where id=? and code<>?',(i,'GAMO'))
 if not customer: abort(404)
 name=(request.form.get('branding_name') or customer['name']).strip()[:80]
 tagline=(request.form.get('brand_tagline') or 'FACILITY MANAGEMENT').strip()[:80]
 color=(request.form.get('brand_color') or '#E31B23').strip().upper()
 if len(color)!=7 or not color.startswith('#') or any(ch not in '0123456789ABCDEF' for ch in color[1:]):
  flash('Farba brandingu musí byť vo formáte #RRGGBB.','error'); return redirect(f'/platform/customer/{i}#branding')
 x_system('update organizations set branding_name=?,brand_color=?,brand_tagline=? where id=?',(name,color,tagline,i))
 audit('CUSTOMER_BRANDING',f"{customer['code']} · {name} · {color}")
 flash('Branding zákazníka bol uložený.','success')
 return redirect(f'/platform/customer/{i}#branding')

@app.post('/platform/customer')
def platform_customer():
 if not is_gamo_admin(): abort(403)
 f=request.form
 code=(f.get('code') or '').strip().upper()
 name=(f.get('name') or '').strip()
 admin_name=(f.get('admin_name') or '').strip()
 email=(f.get('email') or '').strip().lower()
 password=f.get('password') or ''
 plan=(f.get('plan') or 'BUSINESS').upper()
 license_status=f.get('license_status') or 'Aktívna'
 license_until=(f.get('license_until') or '').strip() or None
 if not code or not name or not admin_name or not email or len(password)<10:
  flash('Vyplň povinné údaje. Dočasné heslo musí mať aspoň 10 znakov.','error'); return redirect('/admin#customersAdmin')
 if plan not in {'BASIC','BUSINESS','ENTERPRISE'} or license_status not in {'Aktívna','Pozastavená'}:
  flash('Neplatný licenčný plán alebo stav.','error'); return redirect('/admin#customersAdmin')
 if one_system('select id from organizations where upper(code)=?',(code,)) or one_system('select id from users where lower(email)=?',(email,)):
  flash('Kód zákazníka alebo e-mail administrátora už existuje.','error'); return redirect('/admin#customersAdmin')
 try:
  with con(system=True) as db:
   if USING_POSTGRES:
    row=db.execute(_sql('insert into organizations(code,name,status,plan,license_status,license_until,branding_name,privacy_contact,mfa_required) values(?,?,?,?,?,?,?,?,?) returning id'),(code,name,'Aktívny',plan,license_status,license_until,name,email,True)).fetchone()
    oid=row['id']
   else:
    cur=db.execute(_sql('insert into organizations(code,name,status,plan,license_status,license_until,branding_name,privacy_contact,mfa_required) values(?,?,?,?,?,?,?,?,?)'),(code,name,'Aktívny',plan,license_status,license_until,name,email,1)); oid=cur.lastrowid
   db.execute(_sql('insert into users(name,email,role,status,password_hash,organization_id,must_change_password) values(?,?,?,?,?,?,?)'),(admin_name,email,'Administrator','Aktívny',generate_password_hash(password),oid,True if USING_POSTGRES else 1))
   db.commit()
  audit('CUSTOMER_CREATE',f'{code} · {name} · {plan}')
  flash('Zákazník bol vytvorený. Má vlastnú organizáciu a administrátorský účet.','success')
 except Exception:
  flash('Zákazníka sa nepodarilo vytvoriť. Neboli uložené neúplné dáta.','error')
 return redirect('/admin#customersAdmin')

@app.post('/platform/customer/<int:i>/update')
def platform_customer_update(i):
 if not is_gamo_admin(): abort(403)
 customer=one_system('select * from organizations where id=?',(i,))
 if not customer or customer['code']=='GAMO': abort(404)
 f=request.form
 plan=(f.get('plan') or customer['plan']).upper()
 license_status=f.get('license_status') or customer['license_status']
 org_status=f.get('status') or customer['status']
 license_until=(f.get('license_until') or '').strip() or None
 if plan not in {'BASIC','BUSINESS','ENTERPRISE'} or license_status not in {'Aktívna','Pozastavená'} or org_status not in {'Aktívny','Neaktívny'}:
  abort(400)
 x_system('update organizations set plan=?,license_status=?,license_until=?,status=? where id=?',(plan,license_status,license_until,org_status,i))
 audit('CUSTOMER_UPDATE',f"{customer['code']} · {plan} · {license_status}")
 flash('Nastavenia zákazníka boli uložené.','success')
 return redirect(f"/platform/customer/{i}" if request.form.get('return_to')=='detail' else '/admin#customersAdmin')

@app.route('/add/<what>',methods=['GET','POST'])
def add(what):
 if request.method=='GET':
  return redirect({'incident':'/incidents','workorder':'/maintenance','asset':'/assets','user':'/admin','building':'/buildings'}.get(what,'/'))
 required={'building':'facility_write','floor':'facility_write','room':'facility_write','asset':'asset_write','workorder':'maintenance_write','incident':'incident_write','user':'users_manage'}
 if what not in required: abort(404)
 if not can(required[what]): abort(403)
 f=request.form
 try:
  if what=='building':
   if not plan_allows('buildings'):
    flash('Licenčný limit počtu budov bol dosiahnutý. GAMO môže upraviť licenčný plán.','error'); return redirect('/buildings')
   code=f.get('code','').strip().upper(); name=f.get('name','').strip()
   try: floor_count=max(0,min(50,int(f.get('floors_count') or 0)))
   except (TypeError,ValueError): raise ValueError()
   if not code or not name: flash('Kód a názov budovy sú povinné.','error')
   elif one('select id from buildings where upper(code)=? and organization_id=?',(code,org_id())): flash(f'Budova s kódom {code} už existuje.','error')
   else:
    owner=one('select name from organizations where id=?',(org_id(),))
    customer_name=(owner['name'] if owner else 'GAMO a.s.')
    bid=x('insert into buildings(code,name,address,manager,customer,organization_id) values(?,?,?,?,?,?)',(code,name,f.get('address','').strip(),f.get('manager','').strip(),customer_name,org_id()))
    for n in range(1,floor_count+1): x('insert into floors(building_id,code,name) values(?,?,?)',(bid,f'{n}.NP',f'{n}. nadzemné podlažie'))
    audit('BUILDING_CREATE',f'{name} · {floor_count} podlaží'); flash('Budova a jej základná 3D štruktúra boli vytvorené.','success')
  elif what=='floor':
   building_id=f.get('building_id'); code=(f.get('code') or '').strip().upper(); name=(f.get('name') or '').strip()
   if not owns_building(building_id): abort(404)
   if not code or not name or one('select id from floors where building_id=? and upper(code)=?',(building_id,code)): raise ValueError()
   x('insert into floors(building_id,code,name) values(?,?,?)',(building_id,code,name)); audit('FLOOR_CREATE',f'{code} · {name}'); flash('Podlažie bolo pridané.','success')
  elif what=='room':
   floor_id=f.get('floor_id'); code=(f.get('code') or '').strip().upper(); name=(f.get('name') or '').strip()
   floor=one('select f.id from floors f join buildings b on b.id=f.building_id where f.id=? and b.organization_id=?',(floor_id,org_id()))
   if not floor: abort(404)
   if not code or not name or one('select id from rooms where floor_id=? and upper(code)=?',(floor_id,code)): raise ValueError()
   try: area=max(0,float(f.get('area') or 0))
   except (TypeError,ValueError): raise ValueError()
   x('insert into rooms(floor_id,code,name,area,tenant,zone) values(?,?,?,?,?,?)',(floor_id,code,name,area,f.get('tenant','').strip(),f.get('zone','').strip())); audit('ROOM_CREATE',f'{code} · {name}'); flash('Miestnosť bola pridaná.','success')
  elif what=='asset':
   if not plan_allows('assets'):
    flash('Licenčný limit počtu assetov bol dosiahnutý. GAMO môže upraviť licenčný plán.','error'); return redirect('/assets')
   defaults=org_runtime_defaults()
   name=(f.get('name') or '').strip(); building_id=f.get('building_id'); floor_id=f.get('floor_id'); room_id=f.get('room_id'); parent_id=f.get('parent_id') or None
   profession=(f.get('profession') or defaults['asset']['profession']).strip().upper(); grp=(f.get('grp') or '').strip(); asset_type=(f.get('type') or '').strip()
   if not name: raise ValueError('Zadaj názov zariadenia.')
   if not building_id or not floor_id or not room_id: raise ValueError('Vyber budovu, podlažie aj miestnosť.')
   if not profession or not grp or not asset_type: raise ValueError('Profesia, skupina a typ zariadenia sú povinné.')
   if not owns_building(building_id): abort(404)
   if not one('select f.id from floors f join buildings b on b.id=f.building_id where f.id=? and b.id=? and b.organization_id=?',(floor_id,building_id,org_id())): raise ValueError('Vybrané podlažie nepatrí do zvolenej budovy.')
   if not one('select r.id from rooms r join floors fl on fl.id=r.floor_id join buildings b on b.id=fl.building_id where r.id=? and r.floor_id=? and b.id=? and b.organization_id=?',(room_id,floor_id,building_id,org_id())): raise ValueError('Vybraná miestnosť nepatrí do zvoleného podlažia.')
   if parent_id and not owns_asset(parent_id): raise ValueError('Parent asset nepatrí do tvojej organizácie.')
   status=(f.get('status') or defaults['asset']['status']).strip(); criticality=(f.get('criticality') or defaults['asset']['criticality']).strip()
   if status not in {'Prevádzka','Mimo prevádzky','Servis','Porucha','Vyradené'}: raise ValueError('Neplatný stav assetu.')
   if criticality not in {'A','B','C'}: raise ValueError('Neplatná kritickosť assetu.')
   try:
    service=max(0,int(f.get('service_months') or defaults['asset']['service_months'])); revision=max(0,int(f.get('revision_months') or defaults['asset']['revision_months'])); price=max(0,float(f.get('purchase_price') or 0))
   except (TypeError,ValueError): raise ValueError('Servisný interval, revízia a cena musia byť platné čísla.')
   installed=(f.get('installed') or '').strip() or None; warranty=(f.get('warranty') or '').strip() or None
   for raw,label in ((installed,'Dátum inštalácie'),(warranty,'Záruka')):
    if raw:
     try: datetime.strptime(raw[:10],'%Y-%m-%d')
     except ValueError: raise ValueError(f'{label} nemá platný dátum.')
   asset_tag=(f.get('asset_tag') or '').strip()[:120] or None
   if asset_tag and one('select id from assets where organization_id=? and lower(asset_tag)=lower(?)',(org_id(),asset_tag)):
    raise ValueError(f'Asset Tag {asset_tag} už v tvojej organizácii existuje.')
   manual_aid=(f.get('asset_id') or '').strip().upper()
   if not manual_aid and not defaults['asset_id']['automatic']: raise ValueError('Automatické Asset ID je vypnuté. Zadaj Asset ID ručne.')
   aid=manual_aid or next_asset_id(profession)
   if one('select id from assets where organization_id=? and upper(asset_id)=?',(org_id(),aid)):
    if manual_aid: raise ValueError(f'Asset ID {aid} už v tvojej organizácii existuje. Zmeň ho alebo nechaj pole prázdne pre automatické ID.')
    aid=next_asset_id(profession)
   values=(aid,asset_tag,name,building_id,floor_id,room_id,profession,grp,asset_type,(f.get('manufacturer') or '').strip() or None,(f.get('model') or '').strip() or None,(f.get('serial') or '').strip() or None,(f.get('system_id') or '').strip() or None,parent_id,status,criticality,service,revision,price,installed,warranty,(f.get('ip') or '').strip() or None,(f.get('protocol') or '').strip() or None,(f.get('notes') or '').strip() or None,org_id())
   new_asset=x('insert into assets(asset_id,asset_tag,name,building_id,floor_id,room_id,profession,grp,type,manufacturer,model,serial,system_id,parent_id,status,criticality,service_months,revision_months,purchase_price,installed,warranty,ip,protocol,notes,organization_id) values(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',values)
   asset_event(new_asset,'ASSET_CREATE','Asset vytvorený',f"{aid} · {name}"); audit('ASSET_CREATE',f'{aid} · {name}')
   flash(f'Asset {aid} bol vytvorený.','success')
  elif what=='workorder':
   allowed_priority={'Nízka','Stredná','Vysoká','Kritická'}; allowed_status={'Plánované','Pridelené','Prebieha','Pozastavené','Ukončené','Zrušené'}; allowed_kind={'PM','REV','OPR','VYM'}
   title=(f.get('title') or '').strip(); asset_id=f.get('asset_id'); due=(f.get('due') or '').strip()
   priority=(f.get('priority') or org_runtime_defaults()['workorder']['priority']).strip()
   if not title or priority not in allowed_priority or f.get('status') not in allowed_status or f.get('kind') not in allowed_kind or not owns_asset(asset_id): raise ValueError('Skontroluj asset, typ, prioritu, stav a názov pracovného príkazu.')
   if due:
    try: datetime.strptime(due[:10],'%Y-%m-%d')
    except ValueError: raise ValueError('Termín pracovného príkazu nemá platný dátum.')
   try: cost=max(0,float(f.get('cost') or 0))
   except (TypeError,ValueError): raise ValueError('Náklad pracovného príkazu musí byť platné číslo.')
   x('insert into workorders(asset_id,title,kind,priority,status,due,supplier,technician,cost,description) values(?,?,?,?,?,?,?,?,?,?)',(asset_id,title,f.get('kind'),priority,f.get('status'),due,(f.get('supplier') or '').strip(),(f.get('technician') or '').strip(),cost,(f.get('description') or '').strip()))
   asset_event(asset_id,'WORKORDER_CREATE','Nový pracovný príkaz',f"{f.get('kind','')} · {title}")
   audit('WORKORDER_CREATE',title); flash('Pracovný príkaz bol vytvorený.','success')
  elif what=='incident':
   allowed_severity={'Nízka','Stredná','Vysoká','Kritická','Havária'}; allowed_status={'Otvorená','Pridelená','Rieši sa','Čaká na diel','Vyriešená','Ukončená'}
   title=(f.get('title') or '').strip(); asset_id=f.get('asset_id'); reported=(f.get('reported') or '').strip()
   if not title or f.get('severity') not in allowed_severity or f.get('status') not in allowed_status or not owns_asset(asset_id): raise ValueError('Skontroluj asset, závažnosť, stav a názov incidentu.')
   if reported:
    try: datetime.strptime(reported[:10],'%Y-%m-%d')
    except ValueError: raise ValueError('Dátum nahlásenia incidentu nie je platný.')
   try: cost=max(0,float(f.get('cost') or 0))
   except (TypeError,ValueError): raise ValueError('Náklad incidentu musí byť platné číslo.')
   x('insert into incidents(asset_id,title,severity,status,reported,impact,cause,cost) values(?,?,?,?,?,?,?,?)',(asset_id,title,f.get('severity'),f.get('status'),reported,(f.get('impact') or '').strip(),(f.get('cause') or '').strip(),cost))
   asset_event(asset_id,'INCIDENT_CREATE','Incident zaevidovaný',f"{f.get('severity','')} · {title}")
   audit('INCIDENT_CREATE',title); flash('Incident bol zaevidovaný.','success')
  elif what=='user':
   if not plan_allows('users'):
    flash('Licenčný limit používateľov bol dosiahnutý. GAMO môže upraviť licenčný plán.','error'); return redirect('/admin#usersAdmin')
   name=(f.get('name') or '').strip(); email=(f.get('email') or '').strip().lower(); pwd=f.get('password') or ''
   role=f.get('role') or org_runtime_defaults()['user']['role']; status=f.get('status') or 'Aktívny'
   if not name or not email or len(pwd)<10 or role not in {'Administrator','Facility Manager','Technik','Servisný technik','Viewer'} or status not in {'Aktívny','Neaktívny'}: raise ValueError()
   if one('select id from users where lower(email)=?',(email,)): raise IntegrityError()
   x('insert into users(name,email,role,status,password_hash,organization_id,must_change_password) values(?,?,?,?,?,?,?)',(name,email,role,status,generate_password_hash(pwd),org_id(),True if USING_POSTGRES else 1))
   audit('USER_CREATE',f'{name} · {role}'); flash('Používateľ bol vytvorený.','success')
 except HTTPException:
  raise
 except IntegrityError:
  flash('Záznam sa nepodarilo uložiť pre konflikt v databáze. Skontroluj unikátne kódy a identifikátory.','error')
 except ValueError as exc:
  flash(str(exc) or 'Záznam sa nepodarilo uložiť. Skontroluj zadané hodnoty.','error')
 except Exception:
  flash('Pri ukladaní nastala chyba. Dáta neboli poškodené.','error')
 target={'incident':'/incidents','workorder':'/maintenance','asset':'/assets','user':'/admin','building':'/buildings'}.get(what)
 return redirect(target or safe_referrer_url('/'))
@app.post('/user/<int:i>/update')
def update_user(i):
 if not can('users_manage'): abort(403)
 f=request.form
 u=one('select * from users where id=? and organization_id=?',(i,org_id()))
 if not u:
  abort(404)
 name=(f.get('name') or '').strip(); email=(f.get('email') or '').strip().lower()
 role=f.get('role') or 'Viewer'; status=f.get('status') or 'Aktívny'; pwd=f.get('password') or ''
 allowed_roles={'Administrator','Facility Manager','Technik','Servisný technik','Viewer'}
 if not name or not email or role not in allowed_roles or status not in {'Aktívny','Neaktívny'}:
  flash('Skontroluj údaje používateľa.','error'); return redirect('/admin#usersAdmin')
 if i==session.get('user_id') and status!='Aktívny':
  flash('Aktuálne prihlásený účet nie je možné deaktivovať.','error'); return redirect('/admin#usersAdmin')
 try:
  duplicate=one('select id from users where lower(email)=? and id<>?',(email,i))
  if duplicate: raise IntegrityError()
  if pwd:
   if len(pwd)<10:
    flash('Dočasné heslo musí mať aspoň 10 znakov.','error'); return redirect('/admin#usersAdmin')
   x('update users set name=?,email=?,role=?,status=?,password_hash=?,must_change_password=? where id=?',(name,email,role,status,generate_password_hash(pwd),True if USING_POSTGRES else 1,i))
  else:
   x('update users set name=?,email=?,role=?,status=? where id=?',(name,email,role,status,i))
  if pwd or status!='Aktívny':
   revoke_user_devices(i)
  if i==session.get('user_id'):
   session['user_name']=name; session['user_role']=role
  audit('USER_UPDATE',f'{name} · {role} · {status}')
  flash('Používateľ bol úspešne upravený.','success')
 except IntegrityError:
  flash('Tento e-mail už používa iný účet.','error')
 return redirect('/admin#usersAdmin')

@app.post('/edit/<what>/<int:i>')
def edit_record(what,i):
 required={'building':'facility_write','floor':'facility_write','room':'facility_write','asset':'asset_write','workorder':'maintenance_write','incident':'incident_write'}
 if what not in required: abort(404)
 if not can(required[what]): abort(403)
 f=request.form
 try:
  if what=='building':
   if not owns_building(i): abort(404)
   code=(f.get('code') or '').strip().upper(); name=(f.get('name') or '').strip()
   if not code or not name: raise ValueError()
   if one('select id from buildings where organization_id=? and upper(code)=? and id<>?',(org_id(),code,i)): raise IntegrityError()
   x('update buildings set code=?,name=?,address=?,manager=? where id=? and organization_id=?',(code,name,(f.get('address') or '').strip(),(f.get('manager') or '').strip(),i,org_id()))
   audit('BUILDING_UPDATE',f'{code} · {name}'); flash('Budova bola upravená.','success')
   return redirect(f'/building/{i}')
  if what=='floor':
   if not owns_floor(i): abort(404)
   row=one('select building_id from floors where id=?',(i,)); code=(f.get('code') or '').strip().upper(); name=(f.get('name') or '').strip()
   if not code or not name or not row: raise ValueError()
   if one('select id from floors where building_id=? and upper(code)=? and id<>?',(row['building_id'],code,i)): raise IntegrityError()
   x('update floors set code=?,name=? where id=?',(code,name,i)); audit('FLOOR_UPDATE',f'{code} · {name}'); flash('Podlažie bolo upravené.','success')
   return redirect(safe_referrer_url('/buildings'))
  if what=='room':
   if not owns_room(i): abort(404)
   row=one('select r.floor_id,f.building_id from rooms r join floors f on f.id=r.floor_id where r.id=?',(i,))
   code=(f.get('code') or '').strip().upper(); name=(f.get('name') or '').strip()
   if not code or not name or not row: raise ValueError()
   if one('select id from rooms where floor_id=? and upper(code)=? and id<>?',(row['floor_id'],code,i)): raise IntegrityError()
   area=max(0,float(f.get('area') or 0))
   x('update rooms set code=?,name=?,area=?,tenant=?,zone=? where id=?',(code,name,area,(f.get('tenant') or '').strip(),(f.get('zone') or '').strip(),i))
   audit('ROOM_UPDATE',f'{code} · {name}'); flash('Miestnosť bola upravená.','success')
   return redirect(safe_referrer_url('/buildings'))
  if what=='asset':
   if not owns_asset(i): abort(404)
   aid=(f.get('asset_id') or '').strip().upper(); name=(f.get('name') or '').strip()
   building_id=f.get('building_id'); floor_id=f.get('floor_id'); room_id=f.get('room_id'); parent_id=f.get('parent_id') or None
   profession=(f.get('profession') or '').strip(); grp=(f.get('grp') or '').strip(); asset_type=(f.get('type') or '').strip()
   status=f.get('status'); criticality=f.get('criticality')
   if not all([aid,name,building_id,floor_id,room_id,profession,grp,asset_type]): raise ValueError()
   if status not in {'Prevádzka','Mimo prevádzky','Servis','Porucha','Vyradené'} or criticality not in {'A','B','C'}: raise ValueError()
   if not owns_building(building_id): abort(404)
   if not one('select f.id from floors f join buildings b on b.id=f.building_id where f.id=? and b.id=? and b.organization_id=?',(floor_id,building_id,org_id())): raise ValueError()
   if not one('select r.id from rooms r join floors fl on fl.id=r.floor_id join buildings b on b.id=fl.building_id where r.id=? and r.floor_id=? and b.id=? and b.organization_id=?',(room_id,floor_id,building_id,org_id())): raise ValueError()
   if parent_id and (not owns_asset(parent_id) or not asset_parent_allowed(i,parent_id)): raise ValueError()
   if one('select id from assets where organization_id=? and upper(asset_id)=? and id<>?',(org_id(),aid,i)): raise IntegrityError()
   asset_tag=(f.get('asset_tag') or '').strip()[:120] or None
   if asset_tag and one('select id from assets where organization_id=? and lower(asset_tag)=lower(?) and id<>?',(org_id(),asset_tag,i)): raise IntegrityError()
   service=max(0,int(f.get('service_months') or 0)); revision=max(0,int(f.get('revision_months') or 0)); price=max(0,float(f.get('purchase_price') or 0))
   installed=(f.get('installed') or '').strip() or None; warranty=(f.get('warranty') or '').strip() or None
   for raw in (installed,warranty):
    if raw: datetime.strptime(raw[:10],'%Y-%m-%d')
   old=one('select asset_id,status from assets where id=?',(i,))
   x('''update assets set asset_id=?,asset_tag=?,name=?,building_id=?,floor_id=?,room_id=?,profession=?,grp=?,type=?,manufacturer=?,model=?,serial=?,system_id=?,parent_id=?,status=?,criticality=?,service_months=?,revision_months=?,purchase_price=?,installed=?,warranty=?,ip=?,protocol=?,notes=? where id=? and organization_id=?''',
    (aid,asset_tag,name,building_id,floor_id,room_id,profession,grp,asset_type,(f.get('manufacturer') or '').strip(),(f.get('model') or '').strip(),(f.get('serial') or '').strip(),(f.get('system_id') or '').strip(),parent_id,status,criticality,service,revision,price,installed,warranty,(f.get('ip') or '').strip(),(f.get('protocol') or '').strip(),(f.get('notes') or '').strip(),i,org_id()))
   detail=f"{old['asset_id']} → {aid} · {old['status']} → {status}" if old else f'{aid} · {status}'
   asset_event(i,'ASSET_UPDATE','Asset upravený',detail); audit('ASSET_UPDATE',detail); flash('Asset bol upravený.','success')
   return redirect(f'/asset/{i}')
  if what=='workorder':
   if not owns_workorder(i): abort(404)
   allowed_priority={'Nízka','Stredná','Vysoká','Kritická'}; allowed_status={'Plánované','Pridelené','Prebieha','Pozastavené','Ukončené','Zrušené'}; allowed_kind={'PM','REV','OPR','VYM'}
   asset_id=f.get('asset_id'); title=(f.get('title') or '').strip()
   if not title or not owns_asset(asset_id) or f.get('priority') not in allowed_priority or f.get('status') not in allowed_status or f.get('kind') not in allowed_kind: raise ValueError()
   due=(f.get('due') or '').strip()
   if due: datetime.strptime(due[:10],'%Y-%m-%d')
   cost=max(0,float(f.get('cost') or 0))
   old=one('select status from workorders where id=?',(i,))
   x('update workorders set asset_id=?,title=?,kind=?,priority=?,status=?,due=?,supplier=?,technician=?,cost=?,description=? where id=?',(asset_id,title,f.get('kind'),f.get('priority'),f.get('status'),due,(f.get('supplier') or '').strip(),(f.get('technician') or '').strip(),cost,(f.get('description') or '').strip(),i))
   if f.get('status')=='Ukončené': x('update workorders set completed_at=CURRENT_TIMESTAMP where id=?',(i,))
   else: x('update workorders set completed_at=? where id=?',(None,i))
   asset_event(asset_id,'WORKORDER_UPDATE','Pracovný príkaz upravený',f"{title} · {(old['status'] if old else '—')} → {f.get('status')}")
   audit('WORKORDER_UPDATE',f'{i} · {title}'); flash('Pracovný príkaz bol upravený.','success')
   return redirect(safe_referrer_url('/maintenance'))
  if what=='incident':
   if not owns_incident(i): abort(404)
   allowed_severity={'Nízka','Stredná','Vysoká','Kritická','Havária'}; allowed_status={'Otvorená','Pridelená','Rieši sa','Čaká na diel','Vyriešená','Ukončená'}
   asset_id=f.get('asset_id'); title=(f.get('title') or '').strip()
   if not title or not owns_asset(asset_id) or f.get('severity') not in allowed_severity or f.get('status') not in allowed_status: raise ValueError()
   reported=(f.get('reported') or '').strip()
   if reported: datetime.strptime(reported[:10],'%Y-%m-%d')
   cost=max(0,float(f.get('cost') or 0)); old=one('select status from incidents where id=?',(i,))
   x('update incidents set asset_id=?,title=?,severity=?,status=?,reported=?,impact=?,cause=?,cost=? where id=?',(asset_id,title,f.get('severity'),f.get('status'),reported,(f.get('impact') or '').strip(),(f.get('cause') or '').strip(),cost,i))
   asset_event(asset_id,'INCIDENT_UPDATE','Incident upravený',f"{title} · {(old['status'] if old else '—')} → {f.get('status')}")
   audit('INCIDENT_UPDATE',f'{i} · {title}'); flash('Incident bol upravený.','success')
   return redirect(safe_referrer_url('/incidents'))
  abort(404)
 except HTTPException:
  raise
 except IntegrityError:
  flash('Záznam s rovnakým kódom alebo identifikátorom už existuje.','error')
 except (ValueError,TypeError):
  flash('Záznam sa nepodarilo upraviť. Skontroluj povinné polia a zadané hodnoty.','error')
 except Exception:
  flash('Pri úprave nastala chyba. Pôvodné dáta zostali zachované.','error')
 return redirect(safe_referrer_url('/'))

@app.post('/delete/<what>/<int:i>')
def delete(what,i):
 ownership={
  'building':owns_building,'floor':owns_floor,'room':owns_room,'asset':owns_asset,
  'workorder':owns_workorder,'incident':owns_incident,'user':owns_user
 }
 required={'building':'facility_write','floor':'facility_write','room':'facility_write','asset':'asset_write','workorder':'maintenance_write','incident':'incident_write','user':'users_manage'}
 if what not in ownership or what not in required: abort(404)
 if not can(required[what]): abort(403)
 if not ownership[what](i): abort(404)
 if what=='building':
  if one('select id from assets where building_id=? limit 1',(i,)):
   flash('Budovu nie je možné odstrániť, kým obsahuje assety.','error'); return redirect(safe_referrer_url('/buildings'))
  x('delete from buildings where id=?',(i,)); audit('BUILDING_DELETE',str(i))
 elif what=='floor':
  if one('select id from assets where floor_id=? limit 1',(i,)):
   flash('Podlažie nie je možné odstrániť, kým obsahuje assety.','error'); return redirect(safe_referrer_url('/buildings'))
  x('delete from floors where id=?',(i,)); audit('FLOOR_DELETE',str(i))
 elif what=='room':
  if one('select id from assets where room_id=? limit 1',(i,)):
   flash('Miestnosť nie je možné odstrániť, kým obsahuje assety.','error'); return redirect(safe_referrer_url('/buildings'))
  x('delete from rooms where id=?',(i,)); audit('ROOM_DELETE',str(i))
 elif what=='asset':
  if one('select id from assets where parent_id=? limit 1',(i,)) or one('select id from workorders where asset_id=? limit 1',(i,)) or one('select id from incidents where asset_id=? limit 1',(i,)):
   flash('Asset nie je možné odstrániť, kým má podriadené assety, servisnú históriu alebo incidenty.','error'); return redirect(safe_referrer_url('/assets'))
  x('delete from assets where id=?',(i,)); audit('ASSET_DELETE',str(i))
 elif what=='workorder':
  row=one('select asset_id,title from workorders where id=?',(i,))
  x('delete from workorders where id=?',(i,))
  if row: asset_event(row['asset_id'],'WORKORDER_DELETE','Pracovný príkaz odstránený',row['title'] or '')
  audit('WORKORDER_DELETE',str(i))
 elif what=='incident':
  row=one('select asset_id,title from incidents where id=?',(i,))
  x('delete from incidents where id=?',(i,))
  if row: asset_event(row['asset_id'],'INCIDENT_DELETE','Incident odstránený',row['title'] or '')
  audit('INCIDENT_DELETE',str(i))
 elif what=='user':
  if i==session.get('user_id'):
   flash('Aktuálne prihlásený účet nie je možné odstrániť.','error'); return redirect('/admin#usersAdmin')
  u=one('select role,name from users where id=?',(i,))
  if u and u['role']=='Administrator' and one("select count(*) n from users where organization_id=? and role='Administrator' and status='Aktívny'",(org_id(),))['n']<=1:
   flash('Posledného aktívneho administrátora organizácie nie je možné odstrániť.','error'); return redirect('/admin#usersAdmin')
  x('delete from users where id=?',(i,)); audit('USER_DELETE',u['name'] if u else str(i))
 flash('Záznam bol bezpečne odstránený.','success')
 return redirect(safe_referrer_url('/'))

@app.post('/status/<what>/<int:i>')
def status(what,i):
 allowed={
  'asset':({'Prevádzka','Mimo prevádzky','Servis','Porucha','Vyradené'},owns_asset,'asset_write'),
  'workorder':({'Plánované','Pridelené','Prebieha','Pozastavené','Ukončené','Zrušené'},owns_workorder,'maintenance_write'),
  'incident':({'Otvorená','Pridelená','Rieši sa','Čaká na diel','Vyriešená','Ukončená'},owns_incident,'incident_write')
 }
 if what not in allowed: abort(404)
 statuses,owner_check,permission=allowed[what]
 if not can(permission): abort(403)
 if not owner_check(i): abort(404)
 new_status=request.form.get('status')
 if new_status not in statuses: abort(400)
 if what=='asset':
  x('update assets set status=? where id=?',(new_status,i)); asset_event(i,'STATUS_CHANGE','Zmena stavu assetu',new_status)
 elif what=='workorder':
  row=one('select asset_id,title,status from workorders where id=?',(i,))
  if new_status=='Ukončené':
   x('update workorders set status=?,completed_at=CURRENT_TIMESTAMP where id=?',(new_status,i))
  else:
   x('update workorders set status=?,completed_at=? where id=?',(new_status,None,i))
  if row: asset_event(row['asset_id'],'WORKORDER_STATUS','Zmena stavu pracovného príkazu',f"{row['title']} · {row['status']} → {new_status}")
 else:
  row=one('select asset_id,title from incidents where id=?',(i,)); x('update incidents set status=? where id=?',(new_status,i))
  if row: asset_event(row['asset_id'],'INCIDENT_STATUS','Zmena stavu incidentu',f"{row['title']} → {new_status}")
 audit('STATUS_CHANGE',f'{what}:{i} → {new_status}')
 return redirect(safe_referrer_url('/'))

@app.get('/api/buildings/options')
def api_building_options():
 rows=q('select id,code,name,address from buildings where organization_id=? order by name',(org_id(),))
 return jsonify([dict(r) for r in rows])

@app.get('/api/floors/options')
def api_floor_options():
 rows=q("""select f.id,f.code,f.name,b.id building_id,b.code building_code,b.name building_name
  from floors f join buildings b on b.id=f.building_id
  where b.organization_id=? order by b.name,f.id""",(org_id(),))
 return jsonify([dict(r) for r in rows])

@app.route('/api/floors/<int:b>')
def api_floors(b):
 if not owns_building(b): abort(404)
 return jsonify([dict(r) for r in q('select * from floors where building_id=? order by id',(b,))])

@app.route('/api/rooms/<int:f>')
def api_rooms(f):
 if not owns_floor(f): abort(404)
 return jsonify([dict(r) for r in q('select * from rooms where floor_id=? order by id',(f,))])

@app.get('/api/assets/options')
def api_asset_options():
 rows=q("""select a.id,a.asset_id,a.name,a.profession,b.code building,r.code room
  from assets a join buildings b on b.id=a.building_id
  left join rooms r on r.id=a.room_id
  where b.organization_id=? order by a.asset_id""",(org_id(),))
 return jsonify([dict(r) for r in rows])

@app.get('/api/assets/next-id')
def api_asset_next_id():
 profession=(request.args.get('profession') or 'ASSET').strip()
 return jsonify({'asset_id':next_asset_id(profession)})

@app.get('/api/search')
def api_search():
 term=(request.args.get('q') or '').strip()
 if len(term)<2:return jsonify([])
 needle=term.lower(); like=f'%{needle}%'; out=[]; oid=org_id()
 for r in q("""select a.id,a.asset_id,a.name,a.profession,b.code building,r.code room
  from assets a join buildings b on b.id=a.building_id left join rooms r on r.id=a.room_id
  where b.organization_id=? and (
   lower(coalesce(a.asset_id,'')) like ? or lower(coalesce(a.name,'')) like ?
   or lower(coalesce(a.manufacturer,'')) like ? or lower(coalesce(a.model,'')) like ?
   or lower(coalesce(a.serial,'')) like ? or lower(coalesce(a.system_id,'')) like ?
  ) order by a.asset_id limit 8""",(oid,like,like,like,like,like,like)):
  location=' / '.join(x for x in [r['building'],r['room']] if x)
  out.append({'kind':'Asset','title':f"{r['asset_id']} · {r['name']}",'subtitle':f"{r['profession'] or ''} · {location}".strip(' ·'),'url':f"/asset/{r['id']}"})
 for r in q("""select id,code,name,address from buildings
  where organization_id=? and (lower(coalesce(code,'')) like ? or lower(coalesce(name,'')) like ? or lower(coalesce(address,'')) like ?)
  order by name limit 5""",(oid,like,like,like)):
  out.append({'kind':'Budova','title':f"{r['code']} · {r['name']}",'subtitle':r['address'] or '','url':f"/building/{r['id']}"})
 for r in q("""select r.id,r.code,r.name,f.code floor_code,b.id building_id,b.code building_code,b.name building_name
  from rooms r join floors f on f.id=r.floor_id join buildings b on b.id=f.building_id
  where b.organization_id=? and (lower(coalesce(r.code,'')) like ? or lower(coalesce(r.name,'')) like ? or lower(coalesce(r.tenant,'')) like ? or lower(coalesce(r.zone,'')) like ?)
  order by b.name,f.id,r.code limit 6""",(oid,like,like,like,like)):
  out.append({'kind':'Miestnosť','title':f"{r['code']} · {r['name']}",'subtitle':f"{r['building_code']} · {r['building_name']} / {r['floor_code']}",'url':f"/building/{r['building_id']}#spaces"})
 if can('maintenance_write') or can('reports_view'):
  for r in q("""select w.id,w.title,w.status,w.priority,a.id asset_id,a.asset_id asset_code
   from workorders w join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id
   where b.organization_id=? and (lower(coalesce(w.title,'')) like ? or lower(coalesce(w.technician,'')) like ? or lower(coalesce(w.supplier,'')) like ?)
   order by w.id desc limit 5""",(oid,like,like,like)):
   out.append({'kind':'Údržba','title':r['title'],'subtitle':f"{r['asset_code']} · {r['priority']} · {r['status']}",'url':f"/asset/{r['asset_id']}#service"})
 if can('incident_write') or can('reports_view'):
  for r in q("""select i.id,i.title,i.status,i.severity,a.id asset_id,a.asset_id asset_code
   from incidents i join assets a on a.id=i.asset_id join buildings b on b.id=a.building_id
   where b.organization_id=? and (lower(coalesce(i.title,'')) like ? or lower(coalesce(i.impact,'')) like ? or lower(coalesce(i.cause,'')) like ?)
   order by i.id desc limit 5""",(oid,like,like,like)):
   out.append({'kind':'Incident','title':r['title'],'subtitle':f"{r['asset_code']} · {r['severity']} · {r['status']}",'url':f"/asset/{r['asset_id']}#faults"})
 if platform_ticket_mode():
  ticket_rows=q_system("""select t.id,t.ticket_no,t.subject,t.status,t.priority,o.name organization_name
   from tickets t join organizations o on o.id=t.organization_id
   where o.code<>'GAMO' and (lower(coalesce(t.ticket_no,'')) like ? or lower(coalesce(t.subject,'')) like ?)
   order by t.updated desc limit 6""",(like,like))
 elif ticket_staff():
  ticket_rows=q("select id,ticket_no,subject,status,priority from tickets where organization_id=? and (lower(coalesce(ticket_no,'')) like ? or lower(coalesce(subject,'')) like ?) order by updated desc limit 6",(oid,like,like))
 else:
  ticket_rows=q("select id,ticket_no,subject,status,priority from tickets where organization_id=? and created_by=? and (lower(coalesce(ticket_no,'')) like ? or lower(coalesce(subject,'')) like ?) order by updated desc limit 6",(oid,session.get('user_id'),like,like))
 for r in ticket_rows:
  subtitle=f"{r['priority']} · {r['status']}"
  if platform_ticket_mode(): subtitle=f"{r['organization_name']} · {subtitle}"
  out.append({'kind':'Ticket','title':f"{r['ticket_no']} · {r['subject']}",'subtitle':subtitle,'url':f"/ticket/{r['id']}"})
 return jsonify(out[:20])

@app.get('/api/health')
def api_health():
 started=time.time(); db_ok=False; db_ms=None; counts={}; oid=org_id()
 try:
  t=time.time(); row=one('select 1 as ok'); db_ms=round((time.time()-t)*1000,1); db_ok=bool(row and row['ok']==1)
  if db_ok and oid:
   counts={
    'buildings':one('select count(*) n from buildings where organization_id=?',(oid,))['n'],
    'assets':one('select count(*) n from assets a join buildings b on b.id=a.building_id where b.organization_id=?',(oid,))['n']
   }
 except Exception:
  db_ok=False
 ok=db_ok
 return jsonify({
  'status':'online' if ok else 'degraded','database':'online' if db_ok else 'offline','api':'online',
  'database_ms':db_ms,'response_ms':round((time.time()-started)*1000,1),'version':APP_VERSION,'counts':counts,
  'database_engine':'PostgreSQL' if USING_POSTGRES else 'SQLite','checked_at':datetime.now().isoformat(timespec='seconds')
 }), (200 if ok else 503)

@app.get('/api/notifications')
def api_notifications():
 out=[]; oid=org_id(); today_iso=date.today().isoformat()
 for r in q("""select i.id,i.title,i.status,i.reported,i.severity,a.id aid,a.asset_id from incidents i
  join assets a on a.id=i.asset_id join buildings b on b.id=a.building_id
  where b.organization_id=? and i.status not in ('Ukončená','Vyriešená') order by i.id desc limit 6""",(oid,)):
  critical=r['severity'] in {'Vysoká','Kritická','Havária'}
  out.append({'key':f"incident:{r['id']}",'title':r['title'],'subtitle':f"{r['asset_id']} · {r['severity']}",'status':r['status'],'level':'red' if critical else 'orange','url':f"/asset/{r['aid']}",'created_at':r['reported'] or ''})
 warning_days=org_runtime_defaults()['notifications']['due_warning_days']
 for r in q("""select w.id,w.title,w.status,w.due,a.id aid,a.asset_id from workorders w
  join assets a on a.id=w.asset_id join buildings b on b.id=a.building_id
  where b.organization_id=? and w.status not in ('Ukončené','Zrušené')
  order by case when w.due is null or w.due='' then 1 else 0 end,w.due asc,w.id desc limit 12""",(oid,)):
  overdue=bool(r['due'] and str(r['due'])[:10]<today_iso)
  due_soon=False
  if r['due'] and not overdue:
   try: due_soon=(datetime.strptime(str(r['due'])[:10],'%Y-%m-%d').date()-date.today()).days<=warning_days
   except Exception: due_soon=False
  state='Po termíne' if overdue else ('Blíži sa termín' if due_soon else r['status'])
  out.append({'key':f"workorder:{r['id']}",'title':r['title'],'subtitle':f"{r['asset_id']} · termín {r['due'] or '—'}",'status':state,'level':'red' if overdue else ('orange' if due_soon else 'blue'),'url':f"/asset/{r['aid']}",'created_at':r['due'] or ''})
 for p in [r for r in maintenance_plan_rows() if r['state'] in {'overdue','soon'} and not r['active_workorder_id']][:4]:
  label='Po termíne' if p['state']=='overdue' else 'Blíži sa termín'
  out.append({'key':f"planner:{p['asset_db_id']}:{p['kind']}:{p['due']}",'title':f"{p['label']} · {p['asset_id']}",'subtitle':f"{p['building']} · termín {p['due']}",'status':label,'level':'red' if p['state']=='overdue' else 'orange','url':'/maintenance#planner','created_at':p['due'] or ''})
 uid=session.get('user_id')
 if platform_ticket_mode():
  ticket_rows=q_system("""select t.id,t.ticket_no,t.subject,t.status,t.priority,t.updated,o.name organization_name
   from tickets t join organizations o on o.id=t.organization_id
   where o.code<>'GAMO' and exists(
    select 1 from ticket_messages m left join users su on su.id=m.sender_user_id
    where m.ticket_id=t.id and m.id>coalesce(t.platform_last_read_message_id,0)
    and (m.sender_user_id is null or su.organization_id=t.organization_id))
   order by t.updated desc limit 8""")
 elif ticket_staff():
  ticket_rows=q("""select t.id,t.ticket_no,t.subject,t.status,t.priority,t.updated from tickets t
   where t.organization_id=? and t.created_by<>? and exists(
    select 1 from ticket_messages m
    where m.ticket_id=t.id and m.id>coalesce(t.staff_last_read_message_id,0)
    and (
      m.sender_user_id=t.created_by
      or (
        m.sender_user_id is not null
        and not exists(
          select 1 from users local_sender
          where local_sender.id=m.sender_user_id and local_sender.organization_id=t.organization_id
        )
      )
    ))
   order by t.updated desc limit 5""",(oid,uid))
 else:
  ticket_rows=q("""select t.id,t.ticket_no,t.subject,t.status,t.priority,t.updated from tickets t
   where t.organization_id=? and t.created_by=? and exists(
    select 1 from ticket_messages m where m.ticket_id=t.id and (m.sender_user_id is null or m.sender_user_id<>t.created_by)
    and m.id>coalesce(t.customer_last_read_message_id,0))
   order by t.updated desc limit 5""",(oid,uid))
 for r in ticket_rows:
  subtitle=(f"{r['organization_name']} · nová správa od zákazníka" if platform_ticket_mode() else 'Nová správa v tickete')
  out.insert(0,{'key':f"ticket-unread:{r['id']}:{r['updated']}",'title':f"{r['ticket_no']} · {r['subject']}",'subtitle':subtitle,'status':r['status'],'level':'red' if r['priority']=='Kritická' else 'blue','url':f"/ticket/{r['id']}",'created_at':str(r['updated'] or '')})
 if session.get('user_role')=='Administrator':
  org=one('select license_until,plan from organizations where id=?',(oid,))
  if org and org['license_until']:
   try:
    days=(datetime.strptime(str(org['license_until'])[:10],'%Y-%m-%d').date()-date.today()).days
    if days<=30:
     out.insert(0,{'key':f'license:{oid}:{org["license_until"]}','title':'Platnosť licencie GAMO','subtitle':('Licencia exspirovala' if days<0 else f'Zostáva {days} dní'),'status':org['plan'],'level':'red' if days<7 else 'orange','url':'/admin','created_at':str(org['license_until'])})
   except Exception:
    pass
 return jsonify(out[:12])

@app.get('/api/setting')
def api_setting():
 if not can('settings_manage'): abort(403)
 section=(request.args.get('section') or '').strip()
 r=one('select v from organization_settings where organization_id=? and k=?',(org_id(),section))
 return jsonify({'value':r['v'] if r else ''})

@app.post('/settings/save')
def save_setting():
 if not can('settings_manage'): abort(403)
 section=(request.form.get('section') or '').strip(); value=(request.form.get('value') or '').strip()
 if not section:
  flash('Chýba názov konfiguračnej sekcie.','error')
 else:
  x('insert into organization_settings(organization_id,k,v) values(?,?,?) on conflict(organization_id,k) do update set v=excluded.v',(org_id(),section,value))
  audit('SETTING_UPDATE',section); flash(f'Konfigurácia „{section}“ bola uložená pre tvoju organizáciu.','success')
 return redirect('/admin')

if __name__=='__main__': app.run(debug=False,port=5050)
